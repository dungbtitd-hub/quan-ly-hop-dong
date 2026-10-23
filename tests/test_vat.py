"""Tự tính bộ giá trị VAT theo 4 hướng: ở bộ tính, file Excel mẫu, bảng kê đầu vào, biểu mẫu, dữ liệu cũ."""
import pytest
from openpyxl import load_workbook

from qlhd.core import in_invoice_excel, vat
from qlhd.core.contracts import ContractService, ValidationError, backfill_vat
from qlhd.core.excel_import import import_rows, read_workbook, write_template


@pytest.mark.parametrize("args,expect", [
    ((100_000_000, None, None, "10%"), (100_000_000, 10_000_000, 110_000_000)),   # trước thuế + %
    ((None, None, 102_275, "8%"), (94_699, 7_576, 102_275)),                     # sau thuế + %
    ((94_699, 7_576, None, ""), (94_699, 7_576, 102_275)),                       # trước thuế + VAT
    ((None, 7_576, 102_275, ""), (94_699, 7_576, 102_275)),                      # sau thuế + VAT
    ((94_699, None, 102_275, ""), (94_699, 7_576, 102_275)),                     # trước + sau thuế
    ((94_699, None, None, "8%"), (94_699, 7_576, 102_275)),                      # 7.575,92 làm tròn lên
    ((5_000_000, None, None, "KCT"), (5_000_000, 0, 5_000_000)),
    ((-1_000_000, None, None, "10%"), (-1_000_000, -100_000, -1_100_000)),        # hóa đơn điều chỉnh giảm
])
def test_four_directions(args, expect):
    res = vat.complete(*args)
    assert not res.error and (res.before, res.vat, res.after) == expect


def test_rate_parsing_and_checks():
    assert [vat.parse_rate(x) for x in (0.1, 0.08, 10, "10", "8 %", "5%", "kct", "0,1", None)] == \
        ["10%", "8%", "10%", "10%", "8%", "5%", "KCT", "10%", ""]
    with pytest.raises(ValueError):
        vat.parse_rate("mười")
    assert vat.complete(100, 10, None).rate == "10%"                            # suy ra thuế suất
    assert vat.complete(100, 10, 120).error                                     # mâu thuẫn
    assert vat.complete(1_000_000, 50_000, None, "10%").notes                   # VAT lệch thuế suất
    assert vat.complete(None, 5, None).error                                    # chỉ có tiền VAT


def test_excel_template_auto_vat(tmp_path, company):
    db = company["gd"].db
    tpl = write_template(tmp_path / "mau.xlsx")
    wb = load_workbook(tpl)
    ws = wb["Hợp đồng"]
    head = [c.value.rstrip(" *") for c in ws[1]]
    base = [c.value for c in ws[2]]
    ws.delete_rows(2, ws.max_row)
    col = {k: head.index(k) for k in head}
    cases = [  # số HĐ, trước thuế, thuế suất, VAT, sau thuế
        ("V1", 100_000_000, "10%", None, None), ("V2", None, 0.08, None, 108_000_000),
        ("V3", 50_000_000, None, 4_000_000, None), ("V4", None, None, 5_000_000, 55_000_000),
        ("V5", 10, None, 1, 20)]
    for so, b, r, t, a in cases:
        row = list(base)
        row[col["Số hợp đồng"]] = so
        for k, x in (("Giá trị trước thuế (đồng)", b), ("Thuế suất VAT (%)", r), ("Tiền thuế VAT (đồng)", t),
                     ("Giá trị sau thuế (đồng)", a)):
            row[col[k]] = x
        ws.append(row)
    wb.save(tpl)
    rows = [r for r in read_workbook(tpl, db) if r.sheet == "Hợp đồng"]
    got = [(r.values["GiaTriTruocThue"], r.values["ThueSuat"], r.values["VAT"], r.values["GiaTriTongSauThue"])
           for r in rows[:4]]
    assert got == [(100_000_000, "10%", 10_000_000, 110_000_000), (100_000_000, "8%", 8_000_000, 108_000_000),
                   (50_000_000, "8%", 4_000_000, 54_000_000), (50_000_000, "10%", 5_000_000, 55_000_000)]
    assert rows[0].errors[0].startswith("Tự tính") and rows[4].status == "Lỗi"
    import_rows(db, rows)
    saved = {h["SoHopDong"]: h for h in db.active("HopDong")}
    assert saved["V2"]["GiaTriTruocThue"] == 100_000_000 and saved["V2"]["ThueSuat"] == "8%"


def test_old_template_without_new_columns_still_reads(tmp_path, company):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Hợp đồng"
    ws.append(["Số hợp đồng *", "Tên hợp đồng *", "Chiều hợp đồng *", "Loại hợp đồng",
               "Đối tác (Chủ đầu tư / Nhà cung cấp) *", "MST đối tác", "Giá trị trước thuế (đồng) *",
               "Tiền thuế VAT (đồng)", "Ngày ký *"])
    ws.append(["C1", "Cũ", "Đầu ra", None, "Ban A", None, 200_000_000, 20_000_000, "01/02/2026"])
    wb.save(tmp_path / "cu.xlsx")
    r = read_workbook(tmp_path / "cu.xlsx", company["gd"].db)[0]
    assert r.status == "Hợp lệ" and r.values["GiaTriTongSauThue"] == 220_000_000 and r.values["NgayKy"] == "2026-02-01"


def test_in_invoice_bang_ke_auto_vat():
    head = ("Số hóa đơn", "Ngày hóa đơn", "Mã số thuế người bán", "Tên người bán", "Giá trị chưa có thuế",
            "Thuế suất", "Thuế GTGT", "Tổng thanh toán")
    data = [("1", "30/04/2026", "0109266456", "A", 94_699, "8%", None, None),
            ("2", "30/04/2026", "0109266456", "A", None, "10%", None, 110_000),
            ("3", "30/04/2026", "0109266456", "A", None, None, 800, 10_800),
            ("4", "30/04/2026", "0109266456", "A", 100, None, 10, 200)]
    items = in_invoice_excel.read_rows([head] + data, "S")
    got = [(i.values["GiaTriTruocThue"], i.values["VAT"], i.values["GiaTriSauThue"], i.values["ThueSuat"])
           for i in items]
    assert got[:3] == [(94_699, 7_576, 102_275, "8%"), (100_000, 10_000, 110_000, "10%"),
                       (10_000, 800, 10_800, "8%")]
    assert items[3].errors and not items[0].errors


def test_save_and_backfill_old_data(company):
    db = company["gd"].db
    cs = ContractService(db)
    p = cs.find_or_create_partner("Nhà cung cấp X", "0101010101", "Nhà cung cấp")
    base = {"SoHoaDonGoc": "9", "TenNhaCungCap": "X", "MaThueNCC": "0101010101", "NgayNhanHD": "2026-03-01"}
    ma = cs.save_in_invoice(dict(base, GiaTriSauThue=108_000, ThueSuat="8%"))
    assert db.get("HoaDonDauVao", ma)["VAT"] == 8_000
    with pytest.raises(ValidationError):
        cs.save_in_invoice(dict(base, SoHoaDonGoc="10", GiaTriTruocThue=100, VAT=10, GiaTriSauThue=120))
    # bản cũ: có thuế suất nhưng VAT = 0, sau thuế = trước thuế
    old = db.insert("HoaDonDauVao", dict(base, SoHoaDonGoc="11", GiaTriTruocThue=50_000, VAT=0,
                                         GiaTriSauThue=50_000, ThueSuat="10%"))
    ok = db.insert("HoaDonDauVao", dict(base, SoHoaDonGoc="12", GiaTriTruocThue=50_000, VAT=0,
                                        GiaTriSauThue=50_000, ThueSuat="KCT"))
    hd = cs.save_contract({"SoHopDong": "B1", "TenHopDong": "t", "ChieuHopDong": "Đầu vào", "MaDoiTac": p,
                           "NgayKy": "2026-01-01", "GiaTriTruocThue": 1_000})
    db.update("HopDong", hd, {"GiaTriTongSauThue": 0})
    assert backfill_vat(db) == 2
    assert (db.get("HoaDonDauVao", old)["VAT"], db.get("HoaDonDauVao", old)["GiaTriSauThue"]) == (5_000, 55_000)
    assert db.get("HoaDonDauVao", ok)["VAT"] == 0
    assert db.get("HopDong", hd)["GiaTriTongSauThue"] == 1_000
    assert backfill_vat(db) == 0


def test_form_live_calculation(qtbot):
    from qlhd.ui.widgets import FormDialog, rate_field

    dlg = FormDialog("HĐ", [{"key": "GiaTriTruocThue", "label": "Trước", "kind": "money"}, rate_field(),
                            {"key": "VAT", "label": "VAT", "kind": "money"},
                            {"key": "GiaTriSauThue", "label": "Sau", "kind": "money"}])
    qtbot.addWidget(dlg)
    w = dlg.widgets
    w["ThueSuat"].setEditText("10%")
    qtbot.keyClicks(w["GiaTriTruocThue"], "2000000")
    assert (dlg.values()["VAT"], dlg.values()["GiaTriSauThue"]) == (200_000, 2_200_000)
    w["GiaTriSauThue"].clear()
    qtbot.keyClicks(w["GiaTriSauThue"], "1080000")
    w["ThueSuat"].setEditText("8%")
    v = dlg.values()
    assert (v["GiaTriTruocThue"], v["VAT"], v["GiaTriSauThue"]) == (1_000_000, 80_000, 1_080_000)


# ---------------------------------------------------------------- phí bảo vệ môi trường (hóa đơn nước sạch)
WATER = """<?xml version="1.0" encoding="UTF-8"?>
<HDon><DLHDon><TTChung><KHMSHDon>1</KHMSHDon><KHHDon>K26TNS</KHHDon><SHDon>389849</SHDon><NLap>2026-07-11</NLap>
</TTChung><NDHDon><NBan><Ten>Công ty CP Nước sạch số 2 Hà Nội</Ten><MST>0100106088</MST></NBan>
<NMua><Ten>Công ty</Ten><MST>0101234567</MST></NMua></NDHDon>
<DSHHDVu><HHDVu><TChat>1</TChat><THHDVu>Nước tiêu thụ tháng 7/2026</THHDVu><DVTinh>m3</DVTinh><SLuong>14</SLuong>
<ThTien>124600</ThTien><TSuat>5%</TSuat></HHDVu></DSHHDVu>
<TToan><TgTCThue>124600</TgTCThue><TgTThue>6230</TgTThue>{phi}<TgTTTBSo>143290</TgTTTBSo></TToan></DLHDon></HDon>"""


@pytest.mark.parametrize("args,expect", [
    ((124_600, None, None, "5%", None, "10%"), (124_600, 6_230, 12_460, 143_290)),   # trước thuế + % VAT + % phí
    ((None, None, 143_290, "5%", None, "10%"), (124_600, 6_230, 12_460, 143_290)),   # tổng thanh toán + 2 tỷ lệ
    ((124_600, 6_230, None, "", 12_460, ""), (124_600, 6_230, 12_460, 143_290)),     # nhập số tiền phí
    ((None, 6_230, 143_290, "", 12_460, ""), (124_600, 6_230, 12_460, 143_290)),
])
def test_fee_directions(args, expect):
    res = vat.complete(*args)
    assert not res.error and (res.before, res.vat, res.fee, res.after) == expect
    assert vat.complete(124_600, 6_230, 150_000, "5%", 12_460).error             # tổng không khớp


def test_bang_ke_two_rows_one_invoice():
    head = ("STT", "Số hóa đơn", "Ngày hóa đơn", "Tên người bán", "Mã số thuế người bán", "Mặt hàng",
            "Giá trị chưa có thuế", "Thuế GTGT", "Tổng thanh toán")
    data = [(1, "323404", "14/04/2026", "Công ty CP nước sạch số 2 Hà Nội", "0100106088", "Tiền nước sinh hoạt",
             154_300, 7_715, 162_015),
            (None, None, None, None, None, "Phí bảo vệ môi trường", 15_430, None, 15_430),   # dòng 2 (ô gộp)
            (2, "500", "15/04/2026", "Cửa hàng B", "0300000003", "Văn phòng phẩm", 100_000, 8_000, 108_000)]
    items = in_invoice_excel.read_rows([head] + data, "S")
    assert len(items) == 2
    w = items[0].values
    assert (w["GiaTriTruocThue"], w["VAT"], w["PhiKhac"], w["GiaTriSauThue"]) == (154_300, 7_715, 15_430, 177_445)
    assert not items[0].errors and any("dòng 3" in n for n in items[0].notes)
    assert items[1].values["PhiKhac"] == 0 and not items[1].errors               # một dòng: không có phí


def test_template_fee_columns(tmp_path):
    tpl = in_invoice_excel.write_template(tmp_path / "mau.xlsx")
    ws = load_workbook(tpl)[in_invoice_excel.SHEET]
    head = [c.value.rstrip(" *") for c in ws[1]]
    assert "Phí BVMT / phí khác" in head and "Tỷ lệ phí BVMT (%)" in head
    row = [None] * len(head)
    for k, x in (("Số hóa đơn", "9"), ("Ngày hóa đơn", "11/07/2026"), ("Tên người bán", "Nước sạch"),
                 ("Mã số thuế người bán", "0100106088"), ("Giá trị chưa có thuế", 124_600), ("Thuế suất", "5%"),
                 ("Tỷ lệ phí BVMT (%)", "10%")):
        row[head.index(k)] = x
    item = in_invoice_excel.read_rows([tuple(head), tuple(row)], "S")[0]
    assert (item.values["PhiKhac"], item.values["GiaTriSauThue"]) == (12_460, 143_290) and not item.errors


@pytest.mark.parametrize("phi", ["<DSLPhi><LPhi><TLPhi>Phí BVMT đối với NTSH</TLPhi><TPhi>12460</TPhi></LPhi></DSLPhi>",
                                 ""])
def test_water_xml_fee_and_cost_line(company, phi):
    from qlhd.core.assets import AssetService
    from qlhd.core.xml_invoice import parse_invoice
    from qlhd.ui.invoices_page import invoice_values

    db = company["gd"].db
    d = parse_invoice(WATER.format(phi=phi).encode())
    assert (d["GiaTriTruocThue"], d["VAT"], d["PhiKhac"], d["GiaTriSauThue"]) == (124_600, 6_230, 12_460, 143_290)
    values, notes = invoice_values("vao", d, "")
    ma = ContractService(db).save_in_invoice(values)
    inv = db.get("HoaDonDauVao", ma)
    assert (inv["PhiKhac"], inv["GiaTriSauThue"]) == (12_460, 143_290) and any("12.460" in n for n in notes)
    lines = AssetService(db).lines_of(ma)
    assert [(x["ThanhTien"], x["PhanLoai"], x["MaNhomDM"]) for x in lines] == [
        (124_600, "Chi phí trong kỳ", "CP06"), (12_460, "Chi phí trong kỳ", "CP06")]
    assert sum(x["ThanhTien"] for x in AssetService(db).expense_lines(2026)) == 137_060
    # bỏ phí khi sửa hóa đơn: dòng phí mất theo
    ContractService(db).save_in_invoice(dict(values, PhiKhac=0, GiaTriSauThue=130_830), ma)
    assert len(AssetService(db).lines_of(ma)) == 1


def test_form_fee_rate(qtbot):
    from qlhd.ui.widgets import FormDialog, rate_field

    dlg = FormDialog("HĐ", [{"key": "GiaTriTruocThue", "label": "Trước", "kind": "money"}, rate_field(),
                            {"key": "VAT", "label": "VAT", "kind": "money"},
                            {"key": "PhiKhac", "label": "Phí", "kind": "money"},
                            {"key": "TyLePhiKhac", "label": "Tỷ lệ phí", "kind": "combo_edit", "options": ["10%"]},
                            {"key": "GiaTriSauThue", "label": "Sau", "kind": "money"}])
    qtbot.addWidget(dlg)
    dlg.widgets["ThueSuat"].setEditText("5%")
    qtbot.keyClicks(dlg.widgets["GiaTriTruocThue"], "124600")
    dlg.widgets["TyLePhiKhac"].setEditText("10%")
    v = dlg.values()
    assert (v["VAT"], v["PhiKhac"], v["GiaTriSauThue"]) == (6_230, 12_460, 143_290)
