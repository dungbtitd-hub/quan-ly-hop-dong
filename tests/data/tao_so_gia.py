"""Tạo hai sổ giả (sổ TSCĐ, sổ theo dõi CCDC năm 2025) dùng cho kiểm thử: số liệu bịa, bố cục và cách tính như sổ
xuất từ phần mềm kế toán. Cách tính viết độc lập với qlhd/core/assets.py để kiểm thử có ý nghĩa.

Chạy lại: python tests/data/tao_so_gia.py
"""
import calendar
from datetime import date, datetime
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
NAM = 2025


def tron(a, b):
    return (2 * a + b) // (2 * b)


def ky_trong_nam(ngay: date, so_ky: int, gia_tri: int, muc: int, theo_ngay: bool):
    """Danh sách (năm, tháng, số tiền) từng kỳ."""
    out, cong = [], 0
    for k in range(so_ky):
        y, m = divmod(ngay.month - 1 + k, 12)
        tien = muc
        if k == 0 and theo_ngay and ngay.day > 1:
            dim = calendar.monthrange(ngay.year, ngay.month)[1]
            tien = tron(muc * (dim - ngay.day + 1), dim)
        if k == so_ky - 1:
            tien = gia_tri - cong
        out.append((ngay.year + y, m + 1, tien))
        cong += tien
    return out


def so_lieu(ngay, so_ky, gia_tri, muc, theo_ngay):
    ky = ky_trong_nam(ngay, so_ky, gia_tri, muc, theo_ngay)
    trong = sum(t for y, _, t in ky if y == NAM)
    luy_ke = sum(t for y, _, t in ky if y <= NAM)
    con_ky = sum(1 for y, _, _ in ky if y > NAM)
    return trong, luy_ke, gia_tri - luy_ke, con_ky


def ghi(path, tieu_de, cot, dong, cong):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([tieu_de])
    ws.append([f"Năm {NAM}"])
    ws.append(cot)
    for r in dong:
        ws.append(r)
    ws.append([f"Số dòng = {len(dong)}"] + [sum(r[i] for r in dong) if i in cong else None for i in range(1, len(cot))])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cot))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cot))
    wb.save(path)


def so_tscd():
    # mã, tên, loại, đơn vị, ngày ghi tăng, số tháng, nguyên giá, giá trị tính KH, TK
    ds = [("XE01", "Ô tô con 5 chỗ BKS 99A-000.01 (mua 06/2022)", "Phương tiện vận tải, truyền dẫn", "Chuyên viên",
           date(2023, 1, 1), 60, 600_000_000, 520_000_000, "21113"),
          ("XEDIEN01", "Xe ô tô điện 5 chỗ BKS 99A-000.02", "Phương tiện vận tải, truyền dẫn", "Ban giám đốc",
           date(2024, 3, 1), 72, 690_000_000, 690_000_000, "21113"),
          ("MAYCHU01", "Máy chủ lưu trữ dữ liệu văn phòng", "Máy móc, thiết bị", "Kế toán",
           date(2024, 7, 15), 36, 45_500_000, 45_500_000, "21112"),
          ("DT01", "Điện thoại iPhone 16 Pro Max 512GB", "Máy móc, thiết bị", "Chuyên viên",
           date(2025, 2, 27), 36, 31_818_182, 31_818_182, "21112")]
    dong = []
    for i, (ma, ten, loai, dv, ngay, h, ng, k, tk) in enumerate(ds, start=1):
        muc = tron(k, h)
        trong, luy_ke, con, con_ky = so_lieu(ngay, h, k, muc, True)
        d = datetime(ngay.year, ngay.month, ngay.day)
        dong.append([ma, ten, loai, dv, d, f"GTTS{i:05d}", d, h, con_ky, ng, k, trong, luy_ke, con, muc, tk, "2141"])
    ghi(HERE / "so_tscd_gia.xlsx", "SỔ TÀI SẢN CỐ ĐỊNH",
        ["Mã TSCĐ", "Tên TSCĐ", "Loại TSCĐ", "Đơn vị sử dụng", "Ngày ghi tăng", "Số CT ghi tăng",
         "Ngày bắt đầu tính KH", "Thời gian SD (tháng)", "Thời gian SD còn lại (tháng)", "Nguyên giá",
         "Giá trị tính KH", "Hao mòn trong kỳ", "Hao mòn lũy kế", "Giá trị còn lại", "Giá trị KH tháng",
         "TK nguyên giá", "TK khấu hao"], dong, set(range(7, 15)))


def so_ccdc():
    # mã, tên, lý do, ngày ghi tăng, số kỳ, ĐVT, SL, giá trị, mức kỳ ghi trên sổ (None = tự tính)
    ds = [("CC001/2021", "Máy tính xách tay Dell Inspiron", "Đầu kỳ", date(2023, 1, 1), 4, "Cái", 1, 2_000_002,
           500_001),
          ("CC001/2022", "Máy in laser đen trắng", "Đầu kỳ", date(2023, 1, 1), 18, "Cái", 1, 4_500_000, None),
          ("CC002/2022", "Tủ lạnh văn phòng", "Đầu kỳ", date(2023, 1, 1), 33, "Cái", 1, 6_000_001, None),
          ("CC003/2022", "Cước sử dụng đường bộ xe 99A-000.01", "Đầu kỳ", date(2023, 1, 1), 23, "Gói", 1,
           3_800_000, None),
          ("CC001/2023", "Máy quét tài liệu", "Tăng mới", date(2023, 2, 9), 24, "Cái", 1, 6_000_000, None),
          ("CC002/2023", "Màn hình vi tính 24 inch", "Tăng mới", date(2023, 4, 3), 12, "Cái", 1, 2_200_000, None),
          ("CC003/2023", "Máy tính xách tay MacBook Air M2", "Tăng mới", date(2023, 2, 28), 36, "Cái", 1,
           18_000_001, None),
          ("CC004/2023", "Sơn chống nóng mái văn phòng", "Tăng mới", date(2023, 5, 22), 24, "Gói", 1, 15_000_000,
           None),
          ("CC001/2024", "Gia hạn chữ ký số 3 năm", "Tăng mới", date(2024, 11, 18), 36, "Gói", 1, 1_300_000, None),
          ("CC002/2024", "Điện thoại Redmi Note 13", "Tăng mới", date(2024, 6, 1), 24, "Chiếc", 1, 3_500_000, None),
          ("CC003/2024", "Tấm pin mặt trời 545W", "Tăng mới", date(2025, 1, 1), 36, "Tấm", 20, 36_000_000, None),
          ("CC001/2025", "Ống kính máy ảnh 18-50mm", "Tăng mới", date(2025, 5, 16), 24, "Cái", 1, 11_000_000, None),
          ("CC002/2025", "Đồng hồ thông minh Huawei", "Tăng mới", date(2025, 6, 11), 12, "Cái", 1, 2_100_000, None),
          ("CC003/2025", "Ổ cứng ngoài 1TB", "Tăng mới", date(2025, 11, 25), 12, "Cái", 1, 2_000_000, None)]
    dong = []
    for ma, ten, ly_do, ngay, g, dvt, sl, m, muc in sorted(ds, key=lambda x: x[3]):
        muc = muc or tron(m, g)
        trong, luy_ke, con, con_ky = so_lieu(ngay, g, m, muc, False)
        dong.append([ma, ten, None, ly_do, datetime(ngay.year, ngay.month, ngay.day), ma, g, con_ky, dvt, sl, 0, sl,
                     m, muc, trong, luy_ke, con, "242"])
    ghi(HERE / "so_ccdc_gia.xlsx", "SỔ THEO DÕI CÔNG CỤ DỤNG CỤ",
        ["Mã CCDC", "Tên CCDC", "Loại CCDC", "Lý do ghi tăng", "Ngày ghi tăng", "Số CT ghi tăng", "Số kỳ phân bổ",
         "Số kỳ PB còn lại", "ĐVT", "SL ghi tăng", "Lũy kế SL đã giảm", "SL còn lại", "Giá trị CCDC",
         "Giá trị PB hàng kỳ", "Phân bổ trong kỳ", "Lũy kế đã PB", "Giá trị còn lại", "TK chờ phân bổ"], dong,
        set(range(9, 17)))


if __name__ == "__main__":
    so_tscd()
    so_ccdc()
