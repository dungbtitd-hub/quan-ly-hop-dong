from qlhd.core.excel_import import duplicate_invoice_rows, import_rows, read_workbook, write_template
from qlhd.core.pdf_invoice import parse_invoice, parse_text
from tests.helpers import sample_data

PDF_TEXT = """HÓA ĐƠN GIÁ TRỊ GIA TĂNG
Ký hiệu (Serial): 1C26TAA
Số (No.): 00000125
Ngày (Date) 05 tháng (month) 03 năm (year) 2026
Đơn vị bán hàng (Seller): CÔNG TY CỔ PHẦN ABC
Mã số thuế (Tax code): 0101234567
Tên đơn vị (Company name): BAN QUẢN LÝ DỰ ÁN XYZ
Mã số thuế (Tax code): 0109876543
Cộng tiền hàng (Total amount): 10.000.000
Tiền thuế GTGT (VAT amount): 1.000.000
Tổng cộng tiền thanh toán (Total payment): 11.000.000
"""


def test_parse_pdf_text():
    d = parse_text(PDF_TEXT)
    assert d["KyHieuMauSo"] == "1" and d["KyHieuHD"] == "C26TAA" and d["SoHoaDon"] == "125"
    assert d["NgayLap"] == "2026-03-05"
    assert d["MSTNguoiBan"] == "0101234567" and d["MSTNguoiMua"] == "0109876543"
    assert d["TenNguoiBan"].startswith("CÔNG TY CỔ PHẦN ABC")
    assert (d["GiaTriTruocThue"], d["VAT"], d["GiaTriSauThue"]) == (10_000_000, 1_000_000, 11_000_000)
    assert d["ThieuTruong"] == []


FONTS = ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf",
         "/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf"]


def test_parse_real_pdf_file(tmp_path):
    from pathlib import Path

    import pytest
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    font = next((f for f in FONTS if Path(f).exists()), None)
    if not font:
        pytest.skip("Không có font Unicode để tạo PDF thử")
    pdfmetrics.registerFont(TTFont("VN", font))
    p = tmp_path / "hd.pdf"
    c = canvas.Canvas(str(p))
    c.setFont("VN", 11)
    y = 800
    for line in PDF_TEXT.splitlines():
        c.drawString(40, y, line)
        y -= 16
    c.save()
    d = parse_invoice(p)
    assert d["SoHoaDon"] == "125" and d["GiaTriSauThue"] == 11_000_000


def test_excel_template_roundtrip(company, tmp_path):
    db = company["kt"].db
    tpl = write_template(tmp_path / "mau.xlsx")
    rows = read_workbook(tpl, db)
    contracts = [r for r in rows if r.sheet == "Hợp đồng"]
    assert len(contracts) == 2 and all(r.ok for r in contracts)
    saved = import_rows(db, rows)
    assert saved["Hợp đồng"] == 2 and saved["Hóa đơn đầu ra"] == 1 and saved["Thanh toán"] == 1, saved
    again = read_workbook(tpl, db)
    assert all(r.status == "Trùng" for r in again if r.sheet in ("Hợp đồng", "Hóa đơn đầu ra"))


def test_excel_duplicate_invoice_numbers(company, tmp_path):
    from openpyxl import load_workbook

    db = company["kt"].db
    sample_data(db)                     # đã có hóa đơn C26TAA số 12 của HĐ 01/2026/HĐTV
    tpl = write_template(tmp_path / "mau.xlsx")
    wb = load_workbook(tpl)
    ws = wb["Hóa đơn đầu ra"]
    first = [c.value for c in ws[2]]
    ws.delete_rows(2, ws.max_row)
    cols = [c.value for c in ws[1]]
    i_so, i_kh, i_hd = (next(i for i, c in enumerate(cols) if c.startswith(x))
                        for x in ("Số hóa đơn", "Ký hiệu hóa đơn", "Số hợp đồng"))
    for kh, so in (("C26TAA", "00000012"), ("C25TAA", "12"), ("", "12"), ("C26TAA", "99"), ("C26TAA", "99")):
        row = list(first)
        row[i_kh], row[i_so], row[i_hd] = kh, so, "01/2026/HĐTV"
        ws.append(row)
    wb.save(tpl)
    rows = [r for r in read_workbook(tpl, db) if r.sheet == "Hóa đơn đầu ra"]
    assert [r.status for r in rows] == ["Trùng", "Hợp lệ", "Nghi trùng", "Hợp lệ", "Trùng"], \
        [(r.status, r.errors) for r in rows]
    assert "đã có trong phần mềm" in rows[0].errors[0] and "dòng 5" in rows[4].errors[0]
    assert any(e.startswith("Lưu ý: cùng số") for e in rows[1].errors)
    assert len(duplicate_invoice_rows(rows)) == 3
    rows[2].force = True
    assert import_rows(db, rows)["Hóa đơn đầu ra"] == 3
