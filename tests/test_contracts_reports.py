import pytest

from qlhd.core.contracts import ValidationError
from qlhd.core.reports import CONTEXTS, ContractReports, ReportFilters, export_xlsx, period_range, split_periods
from tests.helpers import sample_data


def test_contract_summary(company):
    d = sample_data(company["kt"].db)
    s = d["cs"].summary(d["hd"])
    assert s["GiaTriHD"] == 110_000_000
    assert s["DaXuatHD"] == 55_000_000
    assert s["DaThanhToan"] == 35_000_000
    assert s["ConPhaiThu"] == 20_000_000
    assert s["ConPhaiXuat"] == 55_000_000
    assert s["TinhTrangThanhToan"] == "Thu một phần"
    assert s["TinhTrangXuatHD"] == "Xuất một phần"


def test_duplicates_rejected(company):
    d = sample_data(company["kt"].db)
    cs = d["cs"]
    with pytest.raises(ValidationError):
        cs.save_out_invoice({"KyHieuHD": "C26TAA", "SoHoaDon": "12", "MaHopDong": d["hd"], "NgayXuatHD": "2026-05-01"})
    with pytest.raises(ValidationError):
        cs.save_contract({"ThoiGianThucHien": 6, "SoHopDong": "01/2026/HĐTV", "TenHopDong": "x", "ChieuHopDong": "Đầu ra",
                          "MaDoiTac": d["cdt"], "NgayKy": "2026-01-01"})


def test_new_contract_requires_duration(company):
    d = sample_data(company["kt"].db)
    cs = d["cs"]
    v = {"SoHopDong": "07/2026", "TenHopDong": "Giám sát", "ChieuHopDong": "Đầu ra", "MaDoiTac": d["cdt"],
         "NgayKy": "2026-03-01", "GiaTriTruocThue": 10_000_000}
    with pytest.raises(ValidationError, match="Thời gian thực hiện"):
        cs.save_contract(dict(v))
    ma = cs.save_contract(dict(v, ThoiGianThucHien=45, DonViThoiGian="ngày"))
    assert cs.db.get("HopDong", ma)["NgayKetThucDuKien"] == "2026-04-14"
    cs.db.update("HopDong", ma, {"ThoiGianThucHien": None})          # hợp đồng cũ chưa có thời gian: sửa vẫn lưu được
    cs.save_contract(dict(cs.db.get("HopDong", ma), GhiChu="sửa"), ma)


def test_all_contexts_and_excel(company, tmp_path):
    db = company["kt"].db
    sample_data(db)
    cr = ContractReports(db)
    assert len(CONTEXTS) == 12
    f = ReportFilters(tach_nam=False, tu="2026-01-01", den="2026-12-31")
    for key, _ in CONTEXTS:
        reps = cr.build(key, f)
        out = export_xlsx(reps, tmp_path / f"{key}.xlsx", company="Cty", user="Kế toán")
        assert out.stat().st_size > 1000
    by_type = cr.build("loai_hd", f)
    assert [r.sheet for r in by_type] == ["Báo cáo đầu ra", "Báo cáo đầu vào"]     # đầu ra, đầu vào tách bảng
    assert by_type[0].rows[0]["DaThanhToan"] == 35_000_000
    merged = cr.build("loai_hd", ReportFilters(tach_nam=False, tu="2026-01-01", den="2026-12-31", gop_ra_vao=True))
    assert len(merged) == 1 and {r["Chieu"] for r in merged[0].rows} == {"Đầu ra", "Đầu vào"}
    paid = cr.build("da_thanh_toan", ReportFilters(tach_nam=False, tu="2026-04-01", den="2026-04-30"))[0]
    assert paid.total["SoTien"] == 5_000_000
    vin = cr.build("hoa_don_vao", ReportFilters(tach_nam=False, co_hd_lien_ket=False))[0]
    data = [r for r in vin.rows if not r.get("_kieu")]
    assert len(data) == 1 and data[0]["SoHopDong"] == "(Mua trực tiếp)"
    group = cr.build("nhom_cdt", f)[0]
    assert group.total["ConPhaiThu"] == 20_000_000
    assert len(cr.build("thoi_gian", f)) == 2 and len(cr.build("hoa_don_ra_vao", f)) == 2
    assert [r.sheet for r in cr.build("tong_hop", f)] == ["Hợp đồng đầu ra", "Hợp đồng đầu vào", "Công việc",
                                                           "Hóa đơn đầu ra", "Thanh toán", "Hóa đơn đầu vào"]


def test_input_invoice_listing_tax_groups(company, tmp_path):
    from openpyxl import load_workbook

    db = company["kt"].db
    d = sample_data(db)
    d["cs"].save_in_invoice({"SoHoaDonGoc": "00086289", "TenNhaCungCap": "Cây xăng 133", "MaThueNCC": "0100598136",
                             "NgayNhanHD": "2026-02-20", "GiaTriTruocThue": 600_000, "VAT": 0, "ThueSuat": "KCT",
                             "MatHang": "Xăng E10 RON 95", "NguoiMuaHang": "Bùi Tiến Dũng"})
    cr = ContractReports(db)
    rep = cr.build("hoa_don_vao", ReportFilters(tach_nam=False))[0]
    heads = [r["_nhan"] for r in rep.rows if r.get("_kieu") == "nhom"]
    assert heads[0].startswith("Nhóm HHDV mua vào : 1.") and heads[0].endswith("(2)")
    assert heads[1].startswith("Nhóm HHDV mua vào : 5.") and heads[1].endswith("(1)")
    assert rep.total["SoHoaDonGoc"] == "Số dòng = 3" and rep.total["VAT"] == 1_080_000
    only = cr.build("hoa_don_vao", ReportFilters(tach_nam=False, nhom_thue="Không chịu thuế"))[0]
    assert [r["MatHang"] for r in only.rows if not r.get("_kieu")] == ["Xăng E10 RON 95"]
    mixed = cr.build("hoa_don_vao", ReportFilters(tach_nam=False, nhom_thue="lan"))[0]
    assert not any(r.get("_kieu") for r in mixed.rows)
    assert [r["NgayNhanHD"] for r in mixed.rows] == sorted(r["NgayNhanHD"] for r in mixed.rows)
    path = export_xlsx(rep, tmp_path / "bang_ke.xlsx")
    ws = load_workbook(path).active
    assert ws.cell(7, 1).value.startswith("Nhóm HHDV mua vào : 1.")
    assert ws.cell(7, 7).value == 11_000_000      # cộng nhóm, cột "chưa có thuế"


def test_out_invoices_one_sheet_per_buyer(company, tmp_path):
    db = company["kt"].db
    d = sample_data(db)
    d["cs"].save_out_invoice({"KyHieuHD": "C26TAA", "SoHoaDon": "13", "MaHopDong": d["hd"], "NgayXuatHD": "2026-04-01",
                              "TenNguoiMua": "Trường THCS Lê Lợi", "MSTNguoiMua": "0109999999",
                              "GiaTriTruocThue": 10_000_000, "VAT": 1_000_000})
    cr = ContractReports(db)
    buyers = cr.buyers()
    assert len(buyers) == 2
    reps = cr.build("hoa_don_ra_nguoi_mua", ReportFilters(tach_nam=False))
    assert len(reps) == 2 and {len(r.rows) for r in reps} == {1}
    one = cr.build("hoa_don_ra_nguoi_mua", ReportFilters(tach_nam=False, nguoi_mua=["mst:0109999999"]))
    assert len(one) == 1 and one[0].sheet == "Trường THCS Lê Lợi"
    export_xlsx(reps, tmp_path / "nguoi_mua.xlsx")


def test_relative_period_for_quick_export():
    from datetime import date

    from qlhd.core.reports import relative_range
    today = date(2026, 1, 20)
    assert relative_range("thang_truoc", today) == ("2025-12-01", "2025-12-31")
    assert relative_range("quy_truoc", today) == ("2025-10-01", "2025-12-31")
    assert relative_range("nam_nay", today) == ("2026-01-01", "2026-12-31")
    f = ReportFilters(ky_tuong_doi="quy_nay").resolved(today)
    assert (f.tu, f.den) == ("2026-01-01", "2026-03-31")


def test_periods():
    assert [p[0] for p in split_periods("2026-01-15", "2026-03-02", "thang")] == [
        "Tháng 01/2026", "Tháng 02/2026", "Tháng 03/2026"]
    assert split_periods("2026-01-01", "2026-12-31", "quy")[3] == ("Quý 4/2026", "2026-10-01", "2026-12-31")
    from datetime import date
    assert period_range("quy", date(2026, 5, 9)) == ("2026-04-01", "2026-06-30")
    assert period_range("thang", date(2026, 12, 9)) == ("2026-12-01", "2026-12-31")


def test_reports_split_by_contract_year(company):
    db = company["kt"].db
    d = sample_data(db)
    cs = d["cs"]
    old = cs.save_contract({"ThoiGianThucHien": 3, "DonViThoiGian": "tháng", "SoHopDong": "09/2025/HĐTV",
                            "TenHopDong": "Thiết kế 2025", "ChieuHopDong": "Đầu ra", "MaDoiTac": d["cdt"],
                            "GiaTriTongSauThue": 50_000_000, "NgayKy": "2025-06-01"})
    cs.save_out_invoice({"SoHoaDon": "5", "MaHopDong": old, "NgayXuatHD": "2026-01-10", "GiaTriSauThue": 20_000_000})
    cr = ContractReports(db)
    assert cr.years(ReportFilters()) == [2025, 2026]
    rep = next(r for r in cr.build("hop_dong", ReportFilters()) if "đầu ra" in r.title)
    labels = [r.get("_nhan") or r.get("SoHopDong") for r in rep.rows]
    assert labels == ["NĂM 2025", "09/2025/HĐTV", "Cộng năm 2025", "NĂM 2026", "01/2026/HĐTV", "Cộng năm 2026"]
    assert rep.total["GiaTriHD"] == 160_000_000 and rep.total["SoHopDong"] == "TỔNG CỘNG"
    # hóa đơn xuất năm 2026 của hợp đồng ký 2025 nằm trong nhóm năm 2025
    inv = cr.build("hoa_don_ra", ReportFilters())[0]
    y25 = inv.rows[:inv.rows.index(next(r for r in inv.rows if r.get("_nhan") == "NĂM 2026"))]
    assert any(r.get("SoHoaDon") == "5" for r in y25)
    # mọi ngữ cảnh xuất được khi tách năm
    for key, _ in CONTEXTS:
        assert cr.build(key, ReportFilters())
    assert cr.build("hop_dong", ReportFilters(tach_nam=False))[0].rows[0].get("_nhan") is None


def test_advance_before_invoice_then_deducted_per_installment(company):
    """Hợp đồng chưa xuất hóa đơn vẫn ghi tạm ứng; có hóa đơn thì khấu trừ tạm ứng, phần còn lại thu đợt sau."""
    db = company["gd"].db
    d = sample_data(db)
    cs = d["cs"]
    hd = cs.save_contract({"ThoiGianThucHien": 3, "SoHopDong": "30/2026/HĐTV", "TenHopDong": "Hai đợt",
                           "ChieuHopDong": "Đầu ra", "MaDoiTac": d["cdt"], "GiaTriTruocThue": 200_000_000,
                           "VAT": 20_000_000, "NgayKy": "2026-05-01"})
    cs.save_payment({"MaHopDong": hd, "NgayThanhToan": "2026-05-05", "SoTien": 30_000_000})
    cs.save_payment({"MaHopDong": hd, "NgayThanhToan": "2026-05-20", "SoTien": 40_000_000})
    assert sum(r["SoTien"] for r in cs.advances(hd)) == 70_000_000
    inv1 = cs.save_out_invoice({"SoHoaDon": "101", "MaHopDong": hd, "NgayXuatHD": "2026-06-01",
                                "GiaTriTruocThue": 50_000_000, "VAT": 5_000_000})
    assert cs.apply_advances(hd, inv1, 50_000_000) == 50_000_000                   # khấu trừ một phần: tách lần 2
    assert sorted(p["SoTien"] for p in cs.invoice_payments(inv1)) == [20_000_000, 30_000_000]
    assert [r["SoTien"] for r in cs.advances(hd)] == [20_000_000]
    cs.save_payment({"MaHopDong": hd, "MaHDDauRa": inv1, "NgayThanhToan": "2026-06-10", "SoTien": 5_000_000})
    assert sum(p["SoTien"] for p in cs.invoice_payments(inv1)) == 55_000_000
    inv2 = cs.save_out_invoice({"SoHoaDon": "102", "MaHopDong": hd, "NgayXuatHD": "2026-08-01",
                                "GiaTriTruocThue": 150_000_000, "VAT": 15_000_000})
    assert cs.apply_advances(hd, inv2, 20_000_000) == 20_000_000 and not cs.advances(hd)
    total = sum(p["SoTien"] for p in db.active("ThanhToanHopDong", "MaHopDong = ?", (hd,)))
    assert total == 75_000_000                                                       # không sinh thêm tiền


def test_reports_aggregate_or_drill_down_to_contracts_and_invoices(company, tmp_path):
    d = sample_data(company["kt"].db)
    cs = d["cs"]
    cs.save_out_invoice({"KyHieuHD": "C26TAA", "SoHoaDon": "15", "MaHopDong": d["hd"], "NgayXuatHD": "2026-05-02",
                         "GiaTriTruocThue": 10_000_000, "VAT": 1_000_000})
    cr = ContractReports(company["kt"].db)
    base = dict(tach_nam=False, tt_thanh_toan="Còn phải thu")
    gop = cr.build("chu_dau_tu", ReportFilters(**base))[0]
    assert [r["TenDoiTac"] for r in gop.rows] == ["Ban QLDA Hà Nội"]          # gộp: chỉ dòng chủ đầu tư
    ct = cr.build("chu_dau_tu", ReportFilters(**base, muc_chi_tiet="hop_dong"))[0]
    kids = [r for r in ct.rows if r.get("_kieu") == "con"]
    assert len(kids) == 1 and "01/2026/HĐTV" in kids[0]["TenDoiTac"] and kids[0]["ConPhaiThu"] == 31_000_000
    assert ct.total["ConPhaiThu"] == gop.total["ConPhaiThu"]                   # chi tiết không cộng trùng
    sau = cr.build("chu_dau_tu", ReportFilters(**base, muc_chi_tiet="hoa_don"))[0]
    inv = [r["TenDoiTac"] for r in sau.rows if "Hóa đơn" in str(r.get("TenDoiTac"))]
    assert len(inv) == 2                                     # hóa đơn 12 còn 25 triệu, hóa đơn 15 chưa thu
    hd = cr.build("hop_dong", ReportFilters(tach_nam=False, chieu="Đầu ra", muc_chi_tiet="hoa_don"))[0]
    assert sum(1 for r in hd.rows if r.get("_kieu") == "con") == 2
    # hóa đơn theo tháng / quý: liệt kê từng hóa đơn có dòng cộng, hoặc chỉ dòng cộng mỗi kỳ
    q = cr.build("hoa_don_ra", ReportFilters(tach_nam=False, nhom_ky_hd="thang"))[0]
    assert [r["_nhan"] for r in q.rows if r.get("_kieu") == "nhom"] == ["Tháng 03/2026 (1 hóa đơn)",
                                                                         "Tháng 05/2026 (1 hóa đơn)"]
    g = cr.build("hoa_don_ra", ReportFilters(tach_nam=False, nhom_ky_hd="quy", chi_cong_ky=True))[0]
    assert [(r["KyHieuHD"], r["GiaTriSauThue"]) for r in g.rows] == [("Quý 1/2026", 55_000_000),
                                                                      ("Quý 2/2026", 11_000_000)]
    assert g.total["GiaTriSauThue"] == 66_000_000
    unpaid_in = cr.build("hoa_don_vao", ReportFilters(tach_nam=False, tt_thanh_toan="Chưa thu", nhom_ky_hd="nam"))[0]
    assert sum(1 for r in unpaid_in.rows if not r.get("_kieu")) == 2           # đầu vào chưa có sao kê
    export_xlsx([ct, sau, q, g], tmp_path / "x.xlsx")


def test_contract_value_mismatch_flag_and_fix(company):
    from qlhd.core import alerts as al
    from qlhd.core.contracts import value_filter_ok
    d = sample_data(company["kt"].db)
    cs, hd = d["cs"], d["hd"]
    # một hóa đơn chưa ghi đợt, HĐ đang thực hiện: có thể còn đợt sau, chỉ tô màu, không cảnh báo
    assert cs.summary(hd)["LechGiaTri"] == "chua_ro_dot"
    assert al.collect(cs.db, "2026-05-01", ["hd_lech_gia_tri"]) == []
    # hợp đồng đã hoàn thành: một hóa đơn chưa ghi đợt = thanh toán toàn bộ, 55 triệu lệch giá trị HĐ 110 triệu
    cs.db.update("HopDong", hd, {"TinhTrangHopDong": "Hoàn thành"})
    assert cs.summary(hd)["LechGiaTri"] == "toan_bo"
    assert value_filter_ok("toan_bo", "lech") and not value_filter_ok("toan_bo", "khop")
    assert [a.loai for a in al.collect(cs.db, "2026-05-01", ["hd_lech_gia_tri"])] == ["hd_lech_gia_tri"]
    # ghi là đợt 1: không còn lệch
    inv = dict(cs.db.get("HoaDonDauRa", d["inv"]))
    cs.save_out_invoice(dict(inv, DotThanhToan="Lần 1"), d["inv"])
    assert cs.summary(hd)["LechGiaTri"] is None
    assert al.collect(cs.db, "2026-05-01", ["hd_lech_gia_tri"]) == []
    # lần cuối mà tổng chưa bằng giá trị HĐ: lệch
    cs.save_out_invoice({"KyHieuHD": "C26TAA", "SoHoaDon": "13", "MaHopDong": hd, "NgayXuatHD": "2026-04-01",
                         "GiaTriTruocThue": 40_000_000, "VAT": 4_000_000, "DotThanhToan": "Lần cuối"})
    assert cs.summary(hd)["LechGiaTri"] == "toan_bo"
    # sửa nhanh: lấy giá trị HĐ theo tổng hóa đơn
    new = cs.value_from_invoices(hd)
    assert new == {"GiaTriTruocThue": 90_000_000, "VAT": 9_000_000, "GiaTriTongSauThue": 99_000_000}
    cs.save_contract(dict(cs.db.get("HopDong", hd), **new), hd)
    assert cs.summary(hd)["LechGiaTri"] is None
    # hóa đơn vượt giá trị HĐ
    cs.save_out_invoice({"KyHieuHD": "C26TAA", "SoHoaDon": "14", "MaHopDong": hd, "NgayXuatHD": "2026-04-02",
                         "GiaTriTruocThue": 1_000_000, "VAT": 100_000, "DotThanhToan": "Lần 3"})
    assert cs.summary(hd)["LechGiaTri"] == "vuot"


def test_new_alert_type_is_on_for_old_saved_choices(company):
    from qlhd.core import alerts as al
    db = company["kt"].db
    db.set_setting("thong_bao_loai:kt", '["hd_het_han"]')       # lựa chọn lưu ở bản cũ
    st = al.AlertSettings(db)
    assert st.app_types("kt") == ["hd_het_han", "hd_cap_nhat_tt", "hd_lech_gia_tri", "thu_thua", "tt_trung"]
    st.set_app_types("kt", ["hd_het_han"])                       # bỏ chọn ở bản mới thì giữ nguyên
    assert st.app_types("kt") == ["hd_het_han"]


def test_report_payment_filter_skips_contracts_without_invoices(company):
    db = company["kt"].db
    d = sample_data(db)
    cs = d["cs"]
    cs.save_contract({"ThoiGianThucHien": 3, "SoHopDong": "09/2026", "TenHopDong": "Chưa xuất hóa đơn",
                      "ChieuHopDong": "Đầu ra", "MaDoiTac": d["cdt"], "GiaTriTruocThue": 5_000_000,
                      "NgayKy": "2026-02-01"})
    cr = ContractReports(db)
    nos = lambda **k: {c["SoHopDong"] for c in cr.contracts(ReportFilters(tach_nam=False, **k))}  # noqa: E731
    assert "09/2026" in nos(tt_xuat_hd="Chưa xuất")
    assert "09/2026" not in nos(tt_thanh_toan="Chưa thu")           # chưa có hóa đơn: không có tình trạng thu
    assert "05/2026/MB" not in nos(tt_thanh_toan="Chưa thu")        # hợp đồng đầu vào không có tiền thu về
    assert nos(tt_thanh_toan="Thu một phần") == {"01/2026/HĐTV"}
    assert "09/2026" in nos(tt_xuat_hd="Chưa xuất", tt_thanh_toan="Đã thu đủ")   # chưa xuất: bỏ qua lọc thu


def test_duplicate_payment_needs_review_and_overpay_warns(company):
    from qlhd.core import duplicates as dupl
    from qlhd.core.alerts import collect
    from qlhd.core.contracts import DuplicateWarning
    d = sample_data(company["kt"].db)
    cs = d["cs"]
    p = {"MaHopDong": d["hd"], "MaHDDauRa": d["inv"], "NgayThanhToan": "2026-03-10", "SoTien": 30_000_000,
         "HinhThuc": "Chuyển khoản", "SoThamChieuSK": "FT123"}
    with pytest.raises(DuplicateWarning) as e:       # giống hệt lần đã ghi: không tự thêm
        cs.save_payment(dict(p))
    assert any("cùng số tham chiếu" in m for m in e.value.matches)
    assert any("Thu thừa" in m for m in e.value.matches)   # 30 + 30 > hóa đơn 55 triệu
    n = len(cs.payments(d["hd"]))
    cs.save_payment(dict(p, **{dupl.CONFIRM_KEY: True}))     # nhân viên rà soát xác nhận mới thêm
    assert len(cs.payments(d["hd"])) == n + 1
    assert [r["SoTien"] for r in dupl.duplicate_payments(cs.db)] == [30_000_000]
    kinds = {a.loai for a in collect(cs.db, "2026-05-01", ["tt_trung", "thu_thua"])}
    assert kinds == {"tt_trung", "thu_thua"}
    # lần thu khác (khác ngày, khác sao kê, không vượt) vẫn thêm bình thường
    cs.save_payment(dict(p, NgayThanhToan="2026-04-20", SoThamChieuSK="FT999", SoTien=1_000_000,
                         MaHDDauRa=None))


def test_duplicate_checks_other_data(company):
    from qlhd.core.contracts import DuplicateWarning
    d = sample_data(company["kt"].db)
    cs = d["cs"]
    cs.save_task({"MaHopDong": d["hd"], "TenNoiDungCV": "Khảo sát"})
    with pytest.raises(DuplicateWarning):
        cs.save_task({"MaHopDong": d["hd"], "TenNoiDungCV": "khảo  sát"})
    with pytest.raises(DuplicateWarning):
        cs.save_partner({"TenDoiTac": "ban qlda  hà nội", "VaiTroDoiTac": "Chủ đầu tư"})
    with pytest.raises(DuplicateWarning):             # cùng số, cùng tiền nhưng khác ký hiệu
        cs.save_out_invoice({"KyHieuHD": "C26TBB", "SoHoaDon": "0012", "MaHopDong": d["hd"],
                             "NgayXuatHD": "2026-03-05", "GiaTriTruocThue": 50_000_000, "VAT": 5_000_000})
    with pytest.raises(DuplicateWarning):
        cs.save_contract({"ThoiGianThucHien": 6, "SoHopDong": "02/2026/HĐTV", "TenHopDong": "Thiết kế trường học",
                          "ChieuHopDong": "Đầu ra", "MaDoiTac": d["cdt"], "NgayKy": "2026-02-01"})
    with pytest.raises(DuplicateWarning):
        cs.save_in_invoice({"KyHieuHD": "K26X", "SoHoaDonGoc": "88", "TenNhaCungCap": "Công ty Vật tư X",
                            "MaThueNCC": "0200000002", "NgayNhanHD": "2026-02-11", "GiaTriTruocThue": 1})
    # sửa bản ghi đã có không bị hỏi
    t = cs.tasks(d["hd"])[0]
    cs.save_task(dict(t, GhiChuCV="sửa"), t["Ma"])


def test_data_export_every_dataset_and_columns(company, tmp_path):
    from qlhd.core import data_export as dx
    db = company["kt"].db
    sample_data(db)
    for ds in dx.DATASETS:
        if ds.finance:
            continue
        rep = dx.build(db, ds.key)
        assert [c[0] for c in rep.columns] == [c[0] for c in ds.columns]
        export_xlsx(rep, tmp_path / f"{ds.key}.xlsx")
    pay = dx.build(db, "thanh_toan", ["NgayThanhToan", "SoTien", "SoThamChieuSK"], None, "2026-12-31")
    assert [c[0] for c in pay.columns] == ["NgayThanhToan", "SoTien", "SoThamChieuSK"]
    assert pay.total["SoTien"] == 35_000_000
    assert len(dx.build(db, "thanh_toan", None, "2026-04-01", None).rows) == 1
    with pytest.raises(ValueError):
        dx.build(db, "thanh_toan", ["KhongCo"])          # không có cột nào hợp lệ
    with pytest.raises(ValueError):
        dx.build(db, "thu_chi")                           # thu chi cần mở khóa
    hd = dx.build(db, "hop_dong", ["SoHopDong", "DaThanhToan"])
    assert {r["SoHopDong"]: r["DaThanhToan"] for r in hd.rows}["01/2026/HĐTV"] == 35_000_000
    rep = dx.select_columns(ContractReports(db).build("da_thanh_toan", ReportFilters(tach_nam=False)), {"SoTien"})
    assert [c[0] for c in rep[0].columns] == ["SoHopDong", "SoTien"]


def test_relative_period_until_today():
    from datetime import date
    from qlhd.core.reports import relative_range
    assert relative_range("den_nay", date(2026, 10, 26)) == ("", "2026-10-26")
    f = ReportFilters(ky_tuong_doi="den_nay").resolved(date(2026, 10, 26))
    assert f.tu is None and f.den == "2026-10-26"


def test_fully_invoiced_contract_asks_status_update_not_expiry(company):
    from qlhd.core import alerts as al
    d = sample_data(company["kt"].db)
    cs, hd = d["cs"], d["hd"]
    cs.db.update("HopDong", hd, {"NgayKetThucDuKien": "2026-04-01"})          # đã quá hạn
    mine = lambda: [a.loai for a in al.collect(cs.db, "2026-05-01", ["hd_het_han", "hd_cap_nhat_tt"])  # noqa: E731
                    if a.ma_hd == hd]
    assert mine() == ["hd_het_han"]
    cs.save_out_invoice({"KyHieuHD": "C26TAA", "SoHoaDon": "20", "MaHopDong": hd, "NgayXuatHD": "2026-04-20",
                         "GiaTriTruocThue": 50_000_000, "VAT": 5_000_000, "DotThanhToan": "Lần cuối"})
    assert mine() == ["hd_cap_nhat_tt"]                                        # xuất đủ: nhắc cập nhật tình trạng
    cs.db.update("HopDong", hd, {"TinhTrangHopDong": "Hoàn thành"})
    assert mine() == []


def test_duplicate_payment_is_not_counted_as_advance(company):
    from qlhd.core import duplicates as dupl
    d = sample_data(company["kt"].db)
    cs, hd = d["cs"], d["hd"]
    before = cs.summary(hd)["TamUngConLai"]
    cs.save_payment({"MaHopDong": hd, "NgayThanhToan": "2026-03-10", "SoTien": 30_000_000, "SoThamChieuSK": "FT123",
                     dupl.CONFIRM_KEY: True})                                    # ghi lại lần đã có, không gắn hóa đơn
    assert cs.summary(hd)["TamUngConLai"] == before
