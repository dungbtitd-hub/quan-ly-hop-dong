import msoffcrypto
import pytest
from openpyxl import load_workbook

from qlhd.core import finance_import as fi
from qlhd.core.finance import FinanceError, FinanceStore
from tests.helpers import sample_data


def _store(m):
    store, _ = FinanceStore.initialize(m.db, m.finance_path, m.user, "khoa-thu-chi-1", iterations=1000)
    return store


def _fill(path, rows):
    wb = load_workbook(path)
    ws = wb[fi.SHEET]
    ws.delete_rows(2, ws.max_row)
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_template_roundtrip_and_duplicates(company, tmp_path):
    gd = company["gd"]
    sample_data(gd.db)
    store = _store(gd)
    tpl = fi.write_template(tmp_path / "mau.xlsx", store)
    lists = load_workbook(tpl)[fi.LIST_SHEET]
    assert lists["D2"].value == "01/2026/HĐTV | Thiết kế trường học"
    assert load_workbook(tpl)[fi.SHEET]["D3"].value == lists["D2"].value
    rows = fi.read_workbook(tpl, store)
    assert [r.status for r in rows] == ["Hợp lệ"] * 3, [r.errors for r in rows]
    project_row = rows[1].values
    assert project_row["PhanLoaiChi"] == "Dự án" and project_row["TenDuAn"] == "Thiết kế trường học"
    res = fi.import_rows(store, rows)
    assert (res["Chi"], res["Thu"], res["tong_chi"]) == (2, 1, 23_000_000) and not res["loi"]
    store.lock()
    store = FinanceStore.unlock(gd.db, gd.finance_path, gd.user, "khoa-thu-chi-1")
    assert len(store.vouchers()) == 3
    assert {v["SoPhieu"][:3] for v in store.vouchers()} == {"PC-", "PT-"}
    again = fi.read_workbook(tpl, store)
    assert all(r.status == "Trùng" for r in again)


def test_errors_new_category_and_defaults(company, tmp_path):
    gd = company["gd"]
    sample_data(gd.db)
    store = _store(gd)
    path = fi.write_template(tmp_path / "mau.xlsx", store)
    _fill(path, [
        ["01/04/2026", "chi", "", "", "", "Lương", "", 1_000_000, "Lương T4"],               # tự chọn Văn phòng
        ["02/04/2026", "Chi", "Dự án", "99/XX", "", "Lương", "", 500_000, "Sai số HĐ"],     # lỗi hợp đồng
        ["03/04/2026", "Chi", "", "", "Công trình B", "Marketing > Quảng cáo", "", 2_000_000, "Chạy quảng cáo"],
        ["04/04/2026", "Thu", "", "", "", "Lãi tiền gửi", "Quỹ không có", 100_000, "Lãi"],  # lỗi quỹ
        ["05/04/2026", "Chuyển", "", "", "", "Lương", "", 100_000, "Sai loại"],              # lỗi loại
        ["06/04/2026", "Chi", "", "", "", "Lương", "", 1_000_000, "Lương T4", "", "", "PC-X1"],
        ["06/04/2026", "Chi", "", "", "", "Lương", "", 1_000_000, "Khác", "", "", "PC-X1"],  # trùng số phiếu
        ["", "Chi", "", "", "", "Lương", "", 0, ""],                                          # thiếu dữ liệu
    ])
    rows = fi.read_workbook(path, store)
    status = [r.status for r in rows]
    assert status == ["Hợp lệ", "Lỗi", "Hợp lệ", "Lỗi", "Lỗi", "Hợp lệ", "Trùng", "Lỗi"], \
        [(r.status, r.errors) for r in rows]
    assert rows[0].values["PhanLoaiChi"] == "Văn phòng"
    assert rows[2].values["PhanLoaiChi"] == "Dự án" and "khoản mục mới" in rows[2].errors[0]
    res = fi.import_rows(store, rows)
    assert res["Chi"] == 3 and res["khoan_muc_moi"] == 2
    qc = [c for c, _ in store.category_tree("Chi") if c["TenKhoanMuc"] == "Quảng cáo"][0]
    assert store.category_path(qc["Ma"]) == "Marketing › Quảng cáo" and qc["ApDungCho"] == "Dự án"


def test_password_protected_file_and_requires_unlock(company, tmp_path):
    gd = company["gd"]
    store = _store(gd)
    plain = fi.write_template(tmp_path / "mau.xlsx", store)
    enc = tmp_path / "mau_mat_khau.xlsx"
    with open(plain, "rb") as f, open(enc, "wb") as out:
        office = msoffcrypto.format.ooxml.OOXMLFile(f)
        office.encrypt("mk-file-123", out)
    assert fi.is_encrypted(enc) and not fi.is_encrypted(plain)
    with pytest.raises(fi.FileNeedsPassword):
        fi.read_workbook(enc, store)
    with pytest.raises(fi.FileNeedsPassword):
        fi.read_workbook(enc, store, "sai")
    assert len(fi.read_workbook(enc, store, "mk-file-123")) == 3
    with pytest.raises(FinanceError):
        fi.read_workbook(plain, None)


def test_contract_picker_disambiguates_same_number(company, tmp_path):
    gd = company["gd"]
    data = sample_data(gd.db)
    cs = data["cs"]
    other = cs.save_partner({"TenDoiTac": "UBND xã B", "MaSoThue": "0100000009", "VaiTroDoiTac": "Chủ đầu tư"})
    cs.save_contract({"ThoiGianThucHien": 6, "SoHopDong": "01/2026/HĐTV", "TenHopDong": "Thiết kế nhà văn hóa", "ChieuHopDong": "Đầu ra",
                      "MaDoiTac": other, "GiaTriTruocThue": 50_000_000, "NgayKy": "2026-05-01"})
    store = _store(gd)
    path = fi.write_template(tmp_path / "mau.xlsx", store)
    labels = [c.value for c in load_workbook(path)[fi.LIST_SHEET]["D"][1:] if c.value]
    assert "01/2026/HĐTV | Thiết kế nhà văn hóa" in labels and len(labels) == 2
    _fill(path, [
        ["01/05/2026", "Chi", "Dự án", "01/2026/HĐTV | Thiết kế nhà văn hóa", "", "Lương", "", 1_000, "Chọn từ danh sách"],
        ["02/05/2026", "Chi", "Dự án", "01/2026/HĐTV", "", "Lương", "", 1_000, "Gõ số trùng"],
    ])
    rows = fi.read_workbook(path, store)
    assert rows[0].ok and rows[0].values["TenDuAn"] == "Thiết kế nhà văn hóa"
    assert rows[0].values["SoHopDong"] == "01/2026/HĐTV"
    assert rows[1].status == "Lỗi" and "chọn từ danh sách" in rows[1].errors[0]
