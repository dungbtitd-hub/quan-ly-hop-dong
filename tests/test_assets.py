from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from qlhd.core import asset_catalog as dm
from qlhd.core.assets import CCDC_KEYS, TSCD_KEYS, AssetError, AssetService, register_row
from qlhd.core.contracts import ContractService
from qlhd.core.finance import FinanceStore
from qlhd.core.profit import ProfitError, ProfitService
from qlhd.core.sync import SyncService
from qlhd.core.xml_invoice import parse_invoice

DATA = Path(__file__).parent / "data"
BOOKS = [(DATA / "so_tscd_gia.xlsx", "TSCĐ", TSCD_KEYS), (DATA / "so_ccdc_gia.xlsx", "CCDC", CCDC_KEYS)]


def _book_rows(path):
    ws = openpyxl.load_workbook(path).active
    return [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0] and not str(r[0]).startswith("Số dòng")]


def _same(book_value, mine):
    if isinstance(book_value, datetime):
        book_value = book_value.date().isoformat()
    return (book_value in (None, "", 0) and mine in (None, "", 0)) or book_value == mine


@pytest.mark.parametrize("path,kind,keys", BOOKS)
def test_import_reproduces_accountant_book(tmp_path, company, path, kind, keys):
    db = company["gd"].db
    s = AssetService(db)
    res = s.import_register(path)
    book = _book_rows(path)
    assert res["nam"] == 2025 and res["them"] == len(book) and not res["lech"] and not res["bo_qua"]
    reg = s.register(kind, 2025)
    assert [r["MaTS"] for r in reg] == [b[0] for b in book]
    for b, r in zip(book, reg):
        for i, k in enumerate(keys):
            if k != "LyDo":
                assert _same(b[i], r[k]), (b[0], k, b[i], r[k])
    # các tài sản trong sổ cũng khớp khi tính lại từ đầu (không dùng mốc của sổ)
    for a in s.assets(kind):
        plain = dict(a, MocNam=None)
        assert register_row(plain, 2025)["LuyKe"] == a["MocLuyKe"]
    # xuất lại rồi nhập lại vào máy khác: vẫn đúng từng đồng
    out = tmp_path / "xuat.xlsx"
    assert s.export_register(out, kind, 2025) == len(book)
    db2 = company["kt"].db
    res2 = AssetService(db2).import_register(out)
    assert res2["them"] == len(book) and not res2["lech"]
    # nhập lại lần nữa thì cập nhật, không nhân đôi
    assert s.import_register(path)["cap_nhat"] == len(book)


def test_continue_next_year_and_cases(company):
    s = AssetService(company["gd"].db)
    for p, _, _ in BOOKS:
        s.import_register(p)
    xe = next(r for r in s.register("TSCĐ", 2026) if r["MaTS"] == "XEDIEN01")
    assert xe["TrongKy"] == 12 * 9_583_333 and xe["KyConLai"] == 38
    oto1 = next(r for r in s.register("TSCĐ", 2026) if r["MaTS"] == "XE01")
    assert oto1["KyConLai"] == 12 and oto1["ConLai"] == 207_999_988 - 12 * 8_666_667
    assert oto1["LuyKe"] + oto1["ConLai"] == 520_000_000
    # năm 2027 tài sản hết khấu hao, kỳ cuối nhận phần còn lại
    oto1 = next(r for r in s.register("TSCĐ", 2027) if r["MaTS"] == "XE01")
    assert oto1["ConLai"] == 0 and oto1["KyConLai"] == 0
    ccdc_2025 = s.register("CCDC", 2025)
    het = s.register("CCDC", 2025, "het_kh")
    con = s.register("CCDC", 2025, "con_kh")
    assert len(het) + len(con) == len(ccdc_2025)
    assert {r["MaTS"] for r in s.register("CCDC", 2025, "tang_moi")} >= {"CC001/2025"}
    assert all(r["LyDo"] == "Đầu kỳ" for r in s.register("CCDC", 2025, "dau_ky"))
    tot = s.year_totals(2025)
    assert sum(tot["TSCĐ"]) == 243_068_179 and sum(tot["CCDC"]) == 29_794_695
    ma = next(a["Ma"] for a in s.assets("CCDC") if a["MaTS"] == "CC003/2025")
    s.write_off(ma, "2026-03-15", "Hỏng")
    r = next(r for r in s.register("CCDC", 2026) if r["MaTS"] == "CC003/2025")
    assert r["ConLai"] == 0 and r["SLConLai"] == 0 and r["TrongKy"] == 1_666_666
    assert "CC003/2025" in {x["MaTS"] for x in s.register("CCDC", 2026, "ghi_giam")}


XML = """<?xml version="1.0" encoding="UTF-8"?>
<HDon><DLHDon><TTChung><KHMSHDon>1</KHMSHDon><KHHDon>C25TAA</KHHDon><SHDon>{so}</SHDon><NLap>{ngay}</NLap></TTChung>
<NDHDon><NBan><Ten>Cửa hàng ABC</Ten><MST>0109999999</MST></NBan><NMua><Ten>Công ty</Ten><MST>0101234567</MST></NMua>
</NDHDon><DSHHDVu>{lines}</DSHHDVu>
<TToan><TgTCThue>{tong}</TgTCThue><TgTThue>0</TgTThue><TgTTTBSo>{tong}</TgTTTBSo></TToan></DLHDon></HDon>"""


def _line(ten, sl, gia, tchat="1"):
    return (f"<HHDVu><TChat>{tchat}</TChat><THHDVu>{ten}</THHDVu><DVTinh>Cái</DVTinh><SLuong>{sl}</SLuong>"
            f"<DGia>{gia}</DGia><ThTien>{sl * gia}</ThTien><TSuat>10%</TSuat></HHDVu>")


def _xml_invoice(cs, so, ngay, lines):
    xml = XML.format(so=so, ngay=ngay, lines="".join(_line(*x) for x in lines), tong=sum(x[1] * x[2] for x in lines))
    d = parse_invoice(xml.encode())
    return cs.save_in_invoice({"SoHoaDonGoc": d["SoHoaDon"], "KyHieuHD": d["KyHieuHD"],
                               "TenNhaCungCap": d["TenNguoiBan"], "MaThueNCC": d["MSTNguoiBan"],
                               "NgayNhanHD": d["NgayLap"], "GiaTriTruocThue": d["GiaTriTruocThue"],
                               "MatHang": d["MatHang"], "_Dong": d["Dong"]})


def test_invoice_lines_are_classified_and_feed_registers(company):
    db = company["kt"].db
    cs, s = ContractService(db), AssetService(db)
    ma = _xml_invoice(cs, "101", "2025-02-27", [("Iphone 16 Pro Max 256GB", 1, 30_436_364),
                                               ("Ốp lưng điện thoại", 2, 150_000), ("Phí thiết kế logo", 1, 2_000_000),
                                               ("Xăng RON95", 10, 23_000), ("Chiết khấu", 1, 100_000, "3")])
    lines = s.lines_of(ma)
    assert [x["PhanLoai"] for x in lines] == ["TSCĐ", "CCDC", "Chưa phân loại", "Chi phí trong kỳ",
                                              "Chưa phân loại"]
    assert lines[-1]["ThanhTien"] == -100_000
    iphone = s.assets("TSCĐ")[0]
    assert (iphone["MaTS"], iphone["SoCTGhiTang"], iphone["SoKy"]) == ("TS001/2025", "GTTS00001", 36)
    r = register_row(iphone, 2025)
    assert (r["TrongKy"], r["KyConLai"], r["MucKy"]) == (8_514_940, 25, 845_455)    # như sổ kế toán
    op = s.assets("CCDC")[0]
    assert op["MaTS"] == "CC001/2025" and op["SoKy"] == 12 and op["NguyenGia"] == 300_000
    pending = s.pending_lines()
    assert {p["TenHang"] for p in pending} == {"Phí thiết kế logo", "Chiết khấu"} and s.pending_count() == 2
    p = next(x for x in pending if x["TenHang"] == "Phí thiết kế logo")
    with pytest.raises(AssetError):
        s.set_line_class(p["Ma"], "Chi phí trong kỳ", "NHOM_TU_DAT", None)       # danh mục khóa
    with pytest.raises(AssetError):
        s.set_line_class(p["Ma"], "TSCĐ", "CP10", 36)                             # sai phân loại của nhóm
    with pytest.raises(AssetError):
        s.set_line_class(p["Ma"], "Chờ phân bổ", "PB02", None)                    # thiếu số tháng
    s.set_line_class(p["Ma"], "Chờ phân bổ", "PB02", 18)
    a = next(x for x in s.assets("CCDC") if x["TenTS"] == "Phí thiết kế logo")
    assert a["PhanLoai"] == "Chờ phân bổ" and a["SoKy"] == 18 and a["TKKhauHao"] == "242"
    s.set_line_class(p["Ma"], "Chi phí trong kỳ", "PB02", None)
    assert not any(x["TenTS"] == "Phí thiết kế logo" for x in s.assets("CCDC"))
    ck = next(x for x in s.pending_lines() if x["TenHang"] == "Chiết khấu")
    s.set_line_class(ck["Ma"], "Chi phí trong kỳ", "CP09", None)
    assert s.pending_count() == 0
    exp = s.expense_lines(2025)
    assert sum(x["ThanhTien"] for x in exp if x["DongLaiLo"] == "6") == 230_000 - 100_000
    assert s.month_warning("TSCĐ", "TB02", 12)
    # xóa hóa đơn thì dòng và tài sản sinh ra cũng bỏ
    cs.delete_in_invoice(ma)
    assert not s.assets() and not s.lines_of(ma)


def test_manual_invoice_single_line_and_backfill(company):
    db = company["kt"].db
    cs, s = ContractService(db), AssetService(db)
    ma = cs.save_in_invoice({"SoHoaDonGoc": "5", "TenNhaCungCap": "Cây xăng", "MaThueNCC": "0100000009",
                             "NgayNhanHD": "2025-05-02", "GiaTriTruocThue": 900_000, "MatHang": "Xăng E5"})
    (ln,) = s.lines_of(ma)
    assert ln["PhanLoai"] == "Chi phí trong kỳ" and ln["MaNhomDM"] == "CP01"
    cs.save_in_invoice({"SoHoaDonGoc": "5", "TenNhaCungCap": "Cây xăng", "MaThueNCC": "0100000009",
                        "NgayNhanHD": "2025-05-02", "GiaTriTruocThue": 1_000_000, "MatHang": "Xăng E5"}, ma)
    assert [x["ThanhTien"] for x in s.lines_of(ma)] == [1_000_000]
    ma2 = cs.save_in_invoice({"SoHoaDonGoc": "6", "TenNhaCungCap": "Siêu thị", "MaThueNCC": "0100000008",
                              "NgayNhanHD": "2025-05-03", "GiaTriTruocThue": 5_000_000,
                              "MatHang": "Máy in; Giấy A4; Bút bi"})
    assert s.lines_of(ma2)[0]["PhanLoai"] == "Chưa phân loại"        # dòng gộp nhiều mặt hàng: kế toán xem
    db.execute("DELETE FROM DongHoaDonDauVao")
    assert s.backfill() == 2 and len(s.lines_of(ma)) == 1


def test_catalog_is_locked_and_matches_books():
    assert len(dm.CATALOG) == len(dm.BY_CODE)
    hits = 0
    for path, kind, _ in BOOKS:
        for r in _book_rows(path):
            price = r[9] if kind == "TSCĐ" else r[12] / (r[9] or 1)
            kq = dm.classify(r[1], price)
            assert kq.nhom is not None, r[1]
            want = {"TSCĐ"} if kind == "TSCĐ" else {"CCDC", "Chờ phân bổ"}
            assert kq.phan_loai in want, (r[1], kq.phan_loai)
            hits += 1
    assert hits == 18
    assert dm.classify("Phí thiết kế kiến trúc", 5_000_000).phan_loai == dm.CHUA_PHAN_LOAI
    assert dm.classify("Giấy A4 Double A 1 ram", 70_000).nhom.ma == "CP04"


KEY = "khoa-thu-chi-1"


def _profit(m):
    store, _ = FinanceStore.initialize(m.db, m.finance_path, m.user, KEY, iterations=1000)
    return store, ProfitService(store)


def test_profit_table_numbers(company):
    gd = company["gd"]
    cs = ContractService(gd.db)
    cdt = cs.save_partner({"TenDoiTac": "Chủ đầu tư", "VaiTroDoiTac": "Chủ đầu tư"})
    hd = cs.save_contract({"ThoiGianThucHien": 6, "SoHopDong": "01/2024", "TenHopDong": "Tư vấn", "ChieuHopDong": "Đầu ra", "MaDoiTac": cdt,
                           "GiaTriTruocThue": 3_000_000_000, "NgayKy": "2024-01-10"})
    cs.save_contract({"ThoiGianThucHien": 6, "SoHopDong": "09/2024", "TenHopDong": "Chưa xuất HĐ", "ChieuHopDong": "Đầu ra",
                      "MaDoiTac": cdt, "GiaTriTruocThue": 300_000_000, "NgayKy": "2024-11-10"})
    cs.save_out_invoice({"SoHoaDon": "1", "MaHopDong": hd, "NgayXuatHD": "2024-06-01",
                         "GiaTriTruocThue": 3_000_000_000, "VAT": 0})
    for i, (ngay, tien) in enumerate([("2024-02-01", 10_000_000), ("2024-05-01", 12_000_000),
                                      ("2024-08-01", 9_000_000), ("2024-11-01", 11_500_000)]):
        cs.save_in_invoice({"SoHoaDonGoc": str(i), "TenNhaCungCap": "Petrolimex", "MaThueNCC": "0100000001",
                            "NgayNhanHD": ngay, "GiaTriTruocThue": tien, "MatHang": "Xăng RON 95"})
    store, p = _profit(gd)
    res = p.compute(2024)
    row = {r["Nguon"] if r["Nguon"] != "nhap" else r["TenDong"]: r for r in res["rows"]}
    assert row["doanh_thu"]["GiaTri"] == 3_000_000_000
    assert row["dong:6"]["GiaTri"] == 42_500_000
    assert row["dong:6"]["Quy"] == [10_000_000, 12_000_000, 9_000_000, 11_500_000]
    assert row["do_dang_cuoi"]["GiaTri"] == -300_000_000            # gợi ý từ hợp đồng chưa xuất hóa đơn
    assert not row["dong:6a"]["Hien"]                                # thầu phụ = 0 thì ẩn
    vals = {"Chi phí thuế môn bài": 2_000_000, "Tiền lương người lao động": 1_800_000_000,
            "Tiền bảo hiểm công ty đóng (21.5%)": 50_000_000}
    for name, v in vals.items():
        p.update_row(row[name]["Ma"], {"GiaTriNhap": v})
    p.update_row(row["tscd"]["Ma"], {"GiaTriNhap": 120_000_000})
    p.update_row(row["ccdc"]["Ma"], {"GiaTriNhap": 40_000_000})
    p.update_row(row["do_dang_dau"]["Ma"], {"GiaTriNhap": 200_000_000})
    p.set_year(2024, {"TyLeThueTNDN": 20, "TNCNConPhaiNop": 30_000_000})
    res = p.compute(2024)
    assert res["chi_phi"] == 1_954_500_000 and res["loi_nhuan"] == 1_045_500_000
    assert res["thue_tndn"] == 209_100_000 and res["tong_phai_nop"] == 239_100_000
    # dòng tự thêm, sắp xếp, xóa; dòng mặc định chỉ ẩn được
    ma = p.add_row(2024, "Điều chỉnh", "Hoàn tiền", -1, 1_000_000)
    assert p.compute(2024)["tong_doanh_thu"] == 2_999_000_000
    p.move_row(ma, -1)
    p.delete_row(ma)
    with pytest.raises(ProfitError):
        p.delete_row(row["tscd"]["Ma"])
    # năm sau: chép cấu trúc, dở dang đầu kỳ gợi ý = dở dang cuối năm trước
    res25 = p.compute(2025)
    assert next(r for r in res25["rows"] if r["Nguon"] == "do_dang_dau")["GiaTri"] == 300_000_000
    p.close_year(2024)
    with pytest.raises(ProfitError):
        p.update_row(row["tscd"]["Ma"], {"GiaTriNhap": 1})
    assert p.compute(2024)["da_chot"]
    p.reopen_year(gd.user, 2024)
    path = gd.dir / "lai_lo.xlsx"
    p.export_xlsx(path, 2024)
    ws = openpyxl.load_workbook(path).active
    assert any(c.value == "Tổng chi phí phải nộp" for c in ws["B"])
    from qlhd.core.reports import export_xlsx
    export_xlsx(p.year_report(2024), gd.dir / "dt_cp.xlsx")
    # dữ liệu lãi lỗ nằm trong tệp mã hóa, xóa vĩnh viễn được
    store.save()
    assert "Tiền lương".encode() not in gd.finance_path.read_bytes()
    assert p.delete_year(2025) > 0 and 2025 not in p.years()
    assert p.delete_all() > 0 and not p.years()


def test_close_year_needs_classification(company):
    gd = company["gd"]
    ContractService(gd.db).save_in_invoice({"SoHoaDonGoc": "1", "TenNhaCungCap": "X", "MaThueNCC": "0100000001",
                                            "NgayNhanHD": "2025-03-01", "GiaTriTruocThue": 1_000_000,
                                            "MatHang": "Dịch vụ lạ"})
    _, p = _profit(gd)
    assert p.compute(2025)["cho_phan_loai"] == 1
    with pytest.raises(ProfitError):
        p.close_year(2025)


def test_lines_and_assets_sync(company, tmp_path):
    kt, gd = company["kt"], company["gd"]
    _xml_invoice(ContractService(kt.db), "7", "2025-08-15", [("Samsung Fold6 SM-F956B", 1, 26_283_333)])
    s = SyncService(kt.db, kt.dir, kt.acc)
    s.export_package(tmp_path / "g.qlhd", "dongbo-2026", user=kt.user)
    d = SyncService(gd.db, gd.dir, gd.acc)
    res = d.apply(d.read_package(tmp_path / "g.qlhd", "dongbo-2026"), user=gd.user)
    assert res["loi"] == 0
    r = AssetService(gd.db).register("CCDC", 2025)[0]
    assert (r["MaTS"], r["SoKy"], r["TrongKy"]) == ("CC001/2025", 36, 3_650_465)


def test_import_flags_duplicates(tmp_path, company):
    from qlhd.core.assets import A_ADD, A_SKIP, ST_IN_FILE, ST_MAYBE, ST_NEW, ST_SAME_CODE

    s = AssetService(company["gd"].db)
    # đã theo dõi trên phần mềm (nhập tay, mã khác) đúng món có trong sổ Excel: cùng ngày, cùng tên khác hoa thường/dấu
    s.save_asset({"LoaiSo": "CCDC", "TenTS": "ONG KINH MAY ANH 18-50mm", "NgayGhiTang": "2025-05-16", "SoKy": 24,
                  "NguyenGia": 11_000_000, "MaTS": "TAY01"})
    # sổ Excel có một dòng bị nhập hai lần
    wb = openpyxl.load_workbook(DATA / "so_ccdc_gia.xlsx")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    total = next(i for i, r in enumerate(rows) if r[0] and str(r[0]).startswith("Số dòng"))
    ws.insert_rows(total + 1)
    for c, val in enumerate(rows[total - 1], start=1):
        ws.cell(total + 1, c, val)
    ws.cell(total + 1, 1, "CC009/2025")                       # cùng ngày, cùng tên, mã khác
    path = tmp_path / "so_trung.xlsx"
    wb.save(path)

    data = s.read_register(path)
    st = {it["values"]["MaTS"]: it for it in data["rows"]}
    assert st["CC001/2025"]["status"] == ST_MAYBE and st["CC001/2025"]["action"] == A_SKIP
    assert "TAY01" in st["CC001/2025"]["trung_voi"]
    assert st["CC009/2025"]["status"] == ST_IN_FILE and st["CC009/2025"]["action"] == A_SKIP
    assert sum(it["status"] == ST_NEW for it in data["rows"]) == 13
    res = s.apply_register(data)
    assert (res["them"], res["cap_nhat"], res["bo_qua"]) == (13, 0, 2)
    assert len(s.assets("CCDC")) == 14 and not s.duplicate_warnings("CCDC")

    # nhập lại: các dòng đã có cùng mã và ngày thì cập nhật, không nhân đôi
    again = s.read_register(path)
    assert sum(it["status"] == ST_SAME_CODE for it in again["rows"]) == 13
    # người dùng cố ý thêm dòng nghi trùng: sổ hiện cảnh báo ở cả hai tài sản
    dong = next(it["dong"] for it in again["rows"] if it["values"]["MaTS"] == "CC001/2025")
    res = s.apply_register(again, {dong: A_ADD})
    assert res["them"] == 1 and res["cap_nhat"] == 13 and len(s.assets("CCDC")) == 15
    warn = s.duplicate_warnings("CCDC")
    assert len(warn) == 2 and any("TAY01" in w for w in warn.values())
    # sổ có mã đã dùng cho tài sản khác ngày ghi tăng: cũng báo nghi trùng
    s2 = AssetService(company["kt"].db)
    ma2 = s2.save_asset({"LoaiSo": "CCDC", "TenTS": "Máy khác", "NgayGhiTang": "2025-01-02", "SoKy": 12,
                         "NguyenGia": 1_000_000, "MaTS": "CC003/2025"})
    code = s2.db.get("TaiSan", ma2)["MaTS"]                            # mã tự cấp được đánh lại: CC001/2025
    it = next(it for it in s2.read_register(DATA / "so_ccdc_gia.xlsx")["rows"] if it["values"]["MaTS"] == code)
    assert it["status"] == ST_MAYBE


def _ccdc(s, name, day, value=1_200_000):
    return s.save_asset({"LoaiSo": "CCDC", "TenTS": name, "NgayGhiTang": day, "SoKy": 12, "NguyenGia": value})


def test_ccdc_codes_renumber_after_edit_and_delete(company):
    s = AssetService(company["gd"].db)
    s.import_register(DATA / "so_ccdc_gia.xlsx")                      # mã sổ kế toán năm 2025 giữ nguyên
    locked = {a["MaTS"] for a in s.assets("CCDC")}
    a = _ccdc(s, "Ghế xoay", "2026-03-10")
    b = _ccdc(s, "Bàn họp", "2026-05-02")
    c = _ccdc(s, "Tủ hồ sơ", "2026-07-20")
    code = lambda ma: s.db.get("TaiSan", ma)["MaTS"]                  # noqa: E731
    assert [code(x) for x in (a, b, c)] == ["CC001/2026", "CC002/2026", "CC003/2026"]
    s.save_asset({"NgayGhiTang": "2026-01-05"}, c)                    # sửa ngày: lên đầu năm
    assert [code(x) for x in (c, a, b)] == ["CC001/2026", "CC002/2026", "CC003/2026"]
    assert s.db.get("TaiSan", c)["SoCTGhiTang"] == "CC001/2026"
    s.delete_asset(c)
    assert [code(x) for x in (a, b)] == ["CC001/2026", "CC002/2026"]
    d = _ccdc(s, "Máy hủy giấy", "2025-12-30")                        # năm 2025: nối tiếp sau mã của sổ kế toán
    n = max(int(m[2:5]) for m in locked if m.endswith("/2025"))
    assert code(d) == f"CC{n + 1:03d}/2025"
    assert locked <= {x["MaTS"] for x in s.assets("CCDC")}


def test_extra_ccdc_template_round_trip(company, tmp_path):
    from qlhd.core import ccdc_excel
    from qlhd.core.assets import A_ADD, A_SKIP, ST_MAYBE, ST_NEW

    s = AssetService(company["gd"].db)
    _ccdc(s, "Máy in mini", "2026-04-01", 2_000_000)
    path = ccdc_excel.write_template(tmp_path / "mau.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb[ccdc_excel.SHEET]
    assert wb.active.title == ccdc_excel.SHEET and str(ws.cell(1, 2).value).startswith("Tên CCDC")
    rows = [["", "Quạt cây", "", "", "Được tặng", datetime(2026, 2, 3), "", 6, "Cái", 3, 900_000, "", ""],
            ["", "Máy in mini", "CCDC", "", "", "01/04/2026", "", 12, "Cái", 1, 2_000_000, "", ""],
            ["", "Kệ sách", "Chờ phân bổ", "", "", "2026-08-15", "PN05", 24, "Cái", 1, 4_800_000, "", ""],
            ["", "Thiếu giá", "", "", "", "2026-08-15", "", 12, "", 1, None, "", ""]]
    for i, r in enumerate(rows, start=3):
        for j, v in enumerate(r, start=1):
            ws.cell(i, j, v)
    wb.save(path)
    data = s.read_extra_ccdc(path)
    st = {it["values"]["TenTS"]: it for it in data["rows"]}
    assert set(st) == {"Quạt cây", "Máy in mini", "Kệ sách"}                # dòng ví dụ bỏ qua, dòng lỗi báo lỗi
    assert len(data["loi"]) == 1 and "giá trị" in data["loi"][0]
    assert st["Quạt cây"]["status"] == ST_NEW and st["Máy in mini"]["status"] == ST_MAYBE
    assert st["Máy in mini"]["action"] == A_SKIP and st["Kệ sách"]["values"]["PhanLoai"] == dm.CHO_PB
    res = s.apply_extra_ccdc(data)
    assert (res["them"], res["bo_qua"]) == (2, 1)
    reg = {r["TenTS"]: r for r in s.register("CCDC", 2026)}
    assert [reg[n]["MaTS"] for n in ("Quạt cây", "Máy in mini", "Kệ sách")] == ["CC001/2026", "CC002/2026",
                                                                               "CC003/2026"]
    assert reg["Kệ sách"]["SoCTGhiTang"] == "PN05" and reg["Quạt cây"]["SoLuong"] == 3
    assert reg["Quạt cây"]["TrongKy"] == 900_000 and "Được tặng" in reg["Quạt cây"]["GhiChu"]
    assert A_ADD


def test_remove_fully_depreciated(company):
    s = AssetService(company["gd"].db)
    old = _ccdc(s, "Bút trình chiếu", "2024-01-10")                     # 12 tháng: hết phân bổ cuối 2024
    mid = _ccdc(s, "Ổ cứng", "2025-06-01")                              # hết phân bổ giữa 2026
    _ccdc(s, "Chuột", "2026-02-01")
    profit_2024 = s.year_totals(2024)
    assert [r["Ma"] for r in s.fully_depreciated("CCDC", 2026)] == [old]
    with pytest.raises(AssetError):
        s.remove_depreciated([mid], 2026)
    assert s.remove_depreciated([old], 2026) == 1
    names = lambda y, case="toan_bo": {r["TenTS"] for r in s.register("CCDC", y, case)}  # noqa: E731
    assert "Bút trình chiếu" not in names(2026) and "Bút trình chiếu" not in names(2027)
    assert "Bút trình chiếu" in names(2025) and names(2026, "da_xoa") == {"Bút trình chiếu"}
    assert s.year_totals(2024) == profit_2024 and not s.fully_depreciated("CCDC", 2026)
    assert [r["Ma"] for r in s.fully_depreciated("CCDC", 2027)] == [mid]
    s.restore_to_register(old)
    assert "Bút trình chiếu" in names(2026)


def test_unclear_lines_are_only_flagged_never_auto_picked(company):
    db = company["kt"].db
    cs, s = ContractService(db), AssetService(db)
    ma = _xml_invoice(cs, "201", "2025-03-03", [("Laptop và giấy vệ sinh", 1, 12_000_000),
                                               ("Thay lốp ô tô", 4, 2_000_000), ("Máy chiếu Epson", 1, 29_000_000),
                                               ("Máy chủ lưu trữ dữ liệu văn phòng", 1, 80_000_000)])
    lines = {x["TenHang"]: x for x in s.lines_of(ma)}
    mixed = lines["Laptop và giấy vệ sinh"]
    assert mixed["PhanLoai"] == dm.CHUA_PHAN_LOAI and mixed["MaNhomDM"]       # chỉ gợi ý nhóm, không tự chọn
    assert lines["Thay lốp ô tô"]["PhanLoai"] == dm.CHI_PHI and lines["Thay lốp ô tô"]["MaNhomDM"] == "CP03"
    assert lines["Máy chủ lưu trữ dữ liệu văn phòng"]["PhanLoai"] == dm.TSCD
    flags = {r["TenHang"]: r["MaCanhBao"] for r in s.review_lines()}
    assert flags["Laptop và giấy vệ sinh"] == "goi_y" and flags["Máy chiếu Epson"] == "nguong"
    assert "Thay lốp ô tô" not in flags
    before = dict(db.get("DongHoaDonDauVao", lines["Máy chiếu Epson"]["Ma"]))
    assert s.confirm_lines([lines["Máy chiếu Epson"]["Ma"], mixed["Ma"]]) == 1     # dòng chưa phân loại: không xác nhận
    after = db.get("DongHoaDonDauVao", lines["Máy chiếu Epson"]["Ma"])
    assert after["PhanLoai"] == before["PhanLoai"] and after["TuDong"] == 0
    assert "Máy chiếu Epson" not in {r["TenHang"] for r in s.review_lines()}


def test_old_auto_classification_is_flagged_after_upgrade_without_changing_it(company):
    db = company["kt"].db
    cs, s = ContractService(db), AssetService(db)
    ma = _xml_invoice(cs, "202", "2025-03-04", [("Sửa chữa máy điều hòa", 1, 3_000_000)])
    ln = s.lines_of(ma)[0]
    db.update("DongHoaDonDauVao", ln["Ma"], {"PhanLoai": dm.CCDC, "MaNhomDM": "TB07", "SoThang": 12, "TuDong": 1})
    r = next(x for x in s.review_lines() if x["Ma"] == ln["Ma"])         # bản cũ tự xếp nhầm vào CCDC
    assert r["MaCanhBao"] == "nghi" and "CP03" in r["CanhBao"]
    assert db.get("DongHoaDonDauVao", ln["Ma"])["PhanLoai"] == dm.CCDC    # chỉ cảnh báo, không tự đổi


def test_recurring_bills_are_not_duplicates_and_confirm_clears_warning(company):
    db = company["kt"].db
    cs, s = ContractService(db), AssetService(db)
    a = _xml_invoice(cs, "301", "2025-04-02", [("Tiền nước sinh hoạt", 1, 144_400)])
    b = _xml_invoice(cs, "302", "2025-05-02", [("Tiền nước sinh hoạt", 1, 144_400)])    # tháng sau, cùng số tiền
    assert not [r for r in s.review_lines() if r["MaCanhBao"] == "trung"]
    c = _xml_invoice(cs, "303", "2025-05-02", [("Tiền nước sinh hoạt", 1, 144_400)])    # cùng ngày: nghi nhập trùng
    dup = [r for r in s.review_lines() if r["MaCanhBao"] == "trung"]
    assert [r["MaHoaDon"] for r in dup] == [c]
    assert s.confirm_lines([dup[0]["Ma"]]) == 1                    # kế toán xác nhận đúng: dòng thôi hiện
    assert not [r for r in s.review_lines() if r["MaHoaDon"] in (a, b, c)]


def test_fixed_asset_depreciation_stops_on_write_off_day(company):
    s = AssetService(company["kt"].db)
    ma = s.save_asset({"LoaiSo": "TSCĐ", "TenTS": "Máy chủ", "NgayGhiTang": "2025-01-01", "SoKy": 36,
                       "NguyenGia": 36_000_000})
    s.write_off(ma, "2025-03-16", "Thanh lý")
    a = s.db.get("TaiSan", ma)
    from qlhd.core.assets import schedule
    sc = schedule(a)
    # T1, T2 đủ tháng; T3 chỉ 15 ngày (01–15/3), thôi trích từ ngày 16/3
    assert [v for _, v in sc.periods] == [1_000_000, 1_000_000, 483_871]
