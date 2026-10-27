"""File Excel mẫu kê khai công cụ dụng cụ nhập thêm ngoài hóa đơn đầu vào (mua lẻ không hóa đơn, được tặng,
điều chuyển, kiểm kê thừa…) và đọc file đó vào sổ CCDC. Cột nhận theo tên tiêu đề, không cần đúng thứ tự."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from . import asset_catalog as dm
from .invoice_match import fold

# (khóa, tiêu đề, bắt buộc, kiểu, các cụm từ nhận diện tiêu đề — đã bỏ dấu, chữ thường; cụm cụ thể đứng trước)
COLUMNS = [
    ("MaTS", "Mã CCDC", False, "text", ("ma ccdc", "ma cong cu")),
    ("TenTS", "Tên CCDC", True, "text", ("ten ccdc", "ten cong cu", "ten")),
    ("PhanLoai", "Loại (CCDC / Chờ phân bổ)", False, "kind", ("loai",)),
    ("MaNhomDM", "Nhóm danh mục", False, "group", ("nhom",)),
    ("LyDo", "Lý do ghi tăng", False, "text", ("ly do",)),
    ("NgayGhiTang", "Ngày ghi tăng", True, "date", ("ngay",)),
    ("SoCTGhiTang", "Số CT ghi tăng", False, "text", ("so ct", "chung tu")),
    ("SoKy", "Số kỳ phân bổ (tháng)", True, "int", ("so ky", "thang")),
    ("DVT", "ĐVT", False, "text", ("dvt", "don vi tinh")),
    ("SoLuong", "Số lượng", False, "qty", ("so luong",)),
    ("NguyenGia", "Giá trị CCDC (đồng)", True, "money", ("gia tri ccdc", "nguyen gia", "thanh tien", "gia tri")),
    ("TKKhauHao", "TK chờ phân bổ", False, "text", ("tk", "tai khoan")),
    ("GhiChu", "Ghi chú", False, "text", ("ghi chu",)),
]
LABELS = {c[0]: c[1] for c in COLUMNS}
SHEET = "Kê khai CCDC"
KINDS = {"ccdc": dm.CCDC, "cong cu dung cu": dm.CCDC, "cho phan bo": dm.CHO_PB,
         "chi phi cho phan bo": dm.CHO_PB}
GROUPS = [n for n in dm.CATALOG if n.kieu in ("GIA_TRI", "CCDC", "CHO_PB")]

GUIDE = [
    "HƯỚNG DẪN KÊ KHAI CÔNG CỤ DỤNG CỤ NHẬP THÊM NGOÀI",
    "",
    "Dùng cho CCDC không đi qua hóa đơn đầu vào trong phần mềm: mua lẻ không có hóa đơn, được tặng, điều chuyển,"
    " kiểm kê phát hiện thừa… CCDC có hóa đơn đầu vào thì nhập hóa đơn, phần mềm tự đưa vào sổ.",
    "1. Điền từ dòng 2 của sheet 'Kê khai CCDC', mỗi dòng một CCDC. Cột có dấu * là bắt buộc. Xóa dòng ví dụ"
    " trước khi nhập.",
    "2. Mã CCDC: để trống thì phần mềm tự đánh CC001/2026, CC002/2026… theo ngày ghi tăng, nối tiếp mã đã có"
    " trong sổ (mã tự đánh lại khi thêm, sửa, xóa CCDC).",
    "3. Loại: 'CCDC' hoặc 'Chờ phân bổ' (để trống là CCDC). Nhóm danh mục: chọn trong danh sách (để trống được).",
    "4. Số kỳ phân bổ: số tháng phân bổ, thường không quá 36 tháng. Giá trị CCDC: tổng giá trị của dòng"
    " (đủ số lượng), phân bổ đều theo tháng từ tháng ghi tăng.",
    "5. Trong phần mềm: Tài sản, CCDC → Sổ CCDC → Nhập CCDC kê khai từ Excel… → xem trước → Nhập vào sổ.",
    "   Dòng nghi trùng với CCDC đã có (cùng ngày ghi tăng, cùng tên hoặc cùng giá trị) được tô màu và mặc định"
    " bỏ qua; dòng trùng mã CCDC đã có thì mặc định cập nhật.",
]
EXAMPLE = ["", "Bàn làm việc gỗ 1m2", "CCDC", "", "Mua lẻ không hóa đơn", date(2026, 5, 15), "", 12, "Cái", 2,
           3600000, "242", "Ví dụ, xóa dòng này trước khi nhập"]


def group_label(n) -> str:
    return f"{n.ma} – {n.ten}"


def write_template(path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    guide = wb.active
    guide.title = "Hướng dẫn"
    for i, line in enumerate(GUIDE, start=1):
        guide.cell(i, 1, line).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    guide.column_dimensions["A"].width = 130
    lists = wb.create_sheet("Danh mục")
    for i, n in enumerate(GROUPS, start=1):
        lists.cell(i, 1, group_label(n))
    lists.column_dimensions["A"].width = 70
    ws = wb.create_sheet(SHEET)
    wb.active = wb.index(ws)
    for c, (key, label, required, kind, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, c, label + (" *" if required else ""))
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FCE4D6" if required else "DCE8F5")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        letter = cell.column_letter
        ws.column_dimensions[letter].width = {"TenTS": 34, "MaNhomDM": 40, "GhiChu": 30, "LyDo": 24}.get(key, 15)
        ws.cell(2, c, EXAMPLE[c - 1])
        fmt = "#,##0" if kind == "money" else "DD/MM/YYYY" if kind == "date" else None
        if fmt:
            for r in range(2, 2001):
                ws[f"{letter}{r}"].number_format = fmt
        if kind in ("kind", "group"):
            formula = '"CCDC,Chờ phân bổ"' if kind == "kind" else f"'Danh mục'!$A$1:$A${len(GROUPS)}"
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}2000")
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    path = Path(path)
    wb.save(path)
    return path


def _header_map(row: list) -> dict[str, int]:
    """{khóa: chỉ số cột} theo tên tiêu đề; mỗi cột nhận một khóa, khóa đứng trước được ưu tiên."""
    out: dict[str, int] = {}
    heads = [fold(str(x or "")).strip(" *") for x in row]
    for key, _, _, _, words in COLUMNS:
        for w in words:
            j = next((j for j, h in enumerate(heads) if h and w in h and j not in out.values()), None)
            if j is not None:
                out[key] = j
                break
    return out


def _day(x) -> str:
    if isinstance(x, datetime):
        return x.date().isoformat()
    if isinstance(x, date):
        return x.isoformat()
    s = str(x).strip()
    if "/" in s:
        d, m, y = s.split("/")[:3]
        return date(int(y[:4]), int(m), int(d)).isoformat()
    return date.fromisoformat(s[:10]).isoformat()


def _num(x) -> float:
    if x in (None, ""):
        return 0
    if isinstance(x, str):
        x = x.replace(".", "").replace(",", ".") if x.count(".") > 1 or "," in x else x
    return float(x)


def read_rows(path: str | Path) -> tuple[list[tuple[int, dict]], list[str]]:
    """Đọc file kê khai: ([(số dòng Excel, giá trị TaiSan)], [lỗi])."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    head_row, cols = None, {}
    for r in range(1, min(ws.max_row, 15) + 1):
        cols = _header_map([ws.cell(r, c).value for c in range(1, ws.max_column + 1)])
        if {"TenTS", "NgayGhiTang", "NguyenGia"} <= cols.keys():
            head_row = r
            break
    if head_row is None:
        raise ValueError("Không nhận ra file kê khai CCDC: cần dòng tiêu đề có Tên CCDC, Ngày ghi tăng, "
                         "Giá trị CCDC (dùng file mẫu của phần mềm).")
    rows, errors = [], []
    for r in range(head_row + 1, ws.max_row + 1):
        raw = {k: ws.cell(r, j + 1).value for k, j in cols.items()}
        raw = {k: (x.strip() if isinstance(x, str) else x) for k, x in raw.items()}
        if not any(x not in (None, "") for x in raw.values()):
            continue
        if "ví dụ" in str(raw.get("GhiChu") or "").lower():
            continue
        try:
            rows.append((r, _values(raw)))
        except (TypeError, ValueError) as e:
            errors.append(f"Dòng {r} ({raw.get('TenTS') or ''}): {e}")
    return rows, errors


def _values(raw: dict) -> dict:
    name = str(raw.get("TenTS") or "").strip()
    if not name:
        raise ValueError("thiếu tên CCDC")
    if raw.get("NgayGhiTang") in (None, ""):
        raise ValueError("thiếu ngày ghi tăng")
    try:
        day = _day(raw["NgayGhiTang"])
    except (TypeError, ValueError):
        raise ValueError(f"ngày ghi tăng không hợp lệ: {raw['NgayGhiTang']}") from None
    so_ky = int(_num(raw.get("SoKy")))
    if not 1 <= so_ky <= 600:
        raise ValueError("số kỳ phân bổ (tháng) phải từ 1 đến 600")
    gia = int(round(_num(raw.get("NguyenGia"))))
    if gia <= 0:
        raise ValueError("thiếu giá trị CCDC")
    kind = fold(str(raw.get("PhanLoai") or "")).strip()
    if kind and kind not in KINDS:
        raise ValueError(f"loại không hợp lệ: {raw['PhanLoai']} (CCDC hoặc Chờ phân bổ)")
    g = str(raw.get("MaNhomDM") or "").strip()
    nhom = dm.BY_CODE.get(g.split("–")[0].split("-")[0].strip().upper()) if g else dm.find_group(name)
    if g and nhom is None:
        raise ValueError(f"nhóm không có trong danh mục: {g}")
    qty = _num(raw.get("SoLuong")) or 1
    note = str(raw.get("GhiChu") or "").strip()
    reason = str(raw.get("LyDo") or "").strip()
    return {"LoaiSo": "CCDC", "MaTS": str(raw.get("MaTS") or "").strip() or None, "TenTS": name,
            "PhanLoai": KINDS.get(kind, dm.CCDC), "MaNhomDM": nhom.ma if nhom else None,
            "NgayGhiTang": day, "NgayBatDauKH": day, "SoCTGhiTang": str(raw.get("SoCTGhiTang") or "").strip() or None,
            "SoKy": so_ky, "DVT": str(raw.get("DVT") or "").strip() or None,
            "SoLuong": int(qty) if qty == int(qty) else qty, "NguyenGia": gia, "GiaTriTinhKH": gia,
            "TKKhauHao": str(raw.get("TKKhauHao") or "").strip() or "242",
            "GhiChu": "; ".join(x for x in (reason, note, "Kê khai nhập thêm ngoài hóa đơn") if x)}
