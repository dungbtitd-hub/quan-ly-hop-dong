"""Sổ tài sản cố định (TSCĐ), sổ theo dõi công cụ dụng cụ (CCDC) và phân loại từng dòng hóa đơn đầu vào.

Cách tính giống sổ kế toán của công ty (đã đối chiếu từng đồng với sổ năm 2025):
- Mức mỗi kỳ = Giá trị (TSCĐ: giá trị tính khấu hao) / số kỳ, làm tròn đến đồng; kỳ cuối nhận phần còn lại.
- CCDC: phân bổ đủ tháng, tính từ tháng ghi tăng.
- TSCĐ: tháng đầu tính theo số ngày còn lại của tháng kể cả ngày bắt đầu (mua 27/02 thì tính 2/28 mức tháng),
  các tháng sau đủ mức; kỳ cuối nhận phần còn lại.
- Giá trị còn lại = giá trị − lũy kế; số kỳ còn lại = số kỳ − số kỳ đã tính đến hết năm.
- Tài sản nhập từ sổ cũ giữ đúng số liệu của sổ đến hết năm đó (mốc), các năm sau tính tiếp từ giá trị còn lại.
"""
from __future__ import annotations

import re
import calendar
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from . import asset_catalog as dm
from .db import Database

CASES = [("toan_bo", "Tổng hợp toàn bộ (như sổ kế toán)"), ("con_kh", "Còn khấu hao / phân bổ"),
         ("het_kh", "Đã hết khấu hao / phân bổ"), ("dau_ky", "Đầu kỳ (chuyển từ các năm trước)"),
         ("tang_moi", "Tăng mới trong năm"), ("co_ps", "Có khấu hao / phân bổ trong năm"),
         ("ghi_giam", "Đã ghi giảm"), ("da_xoa", "Đã xóa khỏi sổ (hết khấu hao / phân bổ)")]

TSCD_COLS = ["Mã TSCĐ", "Tên TSCĐ", "Loại TSCĐ", "Đơn vị sử dụng", "Ngày ghi tăng", "Số CT ghi tăng",
             "Ngày bắt đầu tính KH", "Thời gian SD (tháng)", "Thời gian SD còn lại (tháng)", "Nguyên giá",
             "Giá trị tính KH", "Hao mòn trong kỳ", "Hao mòn lũy kế", "Giá trị còn lại", "Giá trị KH tháng",
             "TK nguyên giá", "TK khấu hao"]
TSCD_KEYS = ["MaTS", "TenTS", "LoaiTS", "DonViSuDung", "NgayGhiTang", "SoCTGhiTang", "NgayBatDauKH", "SoKy",
             "KyConLai", "NguyenGia", "GiaTriTinhKH", "TrongKy", "LuyKe", "ConLai", "MucKy", "TKNguyenGia", "TKKhauHao"]
CCDC_COLS = ["Mã CCDC", "Tên CCDC", "Loại CCDC", "Lý do ghi tăng", "Ngày ghi tăng", "Số CT ghi tăng", "Số kỳ phân bổ",
             "Số kỳ PB còn lại", "ĐVT", "SL ghi tăng", "Lũy kế SL đã giảm", "SL còn lại", "Giá trị CCDC",
             "Giá trị PB hàng kỳ", "Phân bổ trong kỳ", "Lũy kế đã PB", "Giá trị còn lại", "TK chờ phân bổ"]
CCDC_KEYS = ["MaTS", "TenTS", "LoaiTS", "LyDo", "NgayGhiTang", "SoCTGhiTang", "SoKy", "KyConLai", "DVT", "SoLuong",
             "SLDaGiam", "SLConLai", "NguyenGia", "MucKy", "TrongKy", "LuyKe", "ConLai", "TKKhauHao"]
# cột cộng ở dòng "Số dòng = N" (theo sổ kế toán)
TSCD_SUM = {"SoKy", "KyConLai", "NguyenGia", "GiaTriTinhKH", "TrongKy", "LuyKe", "ConLai", "MucKy"}
CCDC_SUM = {"SoLuong", "SLDaGiam", "SLConLai", "NguyenGia", "MucKy", "TrongKy", "LuyKe", "ConLai"}
EXTRA_COLS = [("PhanLoai", "Phân loại"), ("NhomDM", "Nhóm danh mục"), ("LyDoTS", "Đầu kỳ / Tăng mới"),
              ("ThangTrongNam", "Số tháng KH/PB trong năm"), ("NgayGhiGiam", "Ngày ghi giảm"),
              ("HoaDon", "Hóa đơn đầu vào"), ("GhiChu", "Ghi chú")]


FEE_NOTE = "Phí ngoài VAT của hóa đơn"
CCDC_CODE = re.compile(r"CC(\d{3,})/(\d{4})")

# Nhập sổ từ Excel: tình trạng từng dòng và cách xử lý
ST_NEW, ST_SAME_CODE, ST_MAYBE, ST_IN_FILE = "Mới", "Đã có (trùng mã)", "Có thể trùng", "Trùng trong file"
A_ADD, A_UPDATE, A_SKIP = "them", "cap_nhat", "bo_qua"
ACTION_LABELS = {A_ADD: "Thêm mới", A_UPDATE: "Cập nhật bản đã có", A_SKIP: "Bỏ qua"}


class AssetError(Exception):
    pass


def _same_asset(a: dict, b: dict) -> bool:
    """Cùng ngày ghi tăng và (cùng tên sau khi bỏ dấu, hoa thường, hoặc cùng giá trị)."""
    if str(a["NgayGhiTang"])[:10] != str(b["NgayGhiTang"])[:10]:
        return False
    if dm.normalize(a["TenTS"]) == dm.normalize(b["TenTS"]):
        return True
    return int(a.get("NguyenGia") or 0) > 0 and int(a.get("NguyenGia") or 0) == int(b.get("NguyenGia") or 0)


def _dmy(iso: str) -> str:
    return "/".join(reversed(str(iso)[:10].split("-")))


# ------------------------------------------------------------------ tính toán
def div_round(a: int, b: int) -> int:
    """a / b làm tròn nửa lên (như Excel ROUND), không dùng số thực."""
    if a < 0:
        return -div_round(-a, b)
    return (2 * a + b) // (2 * b)


def month_index(d: str | date) -> int:
    d = date.fromisoformat(str(d)[:10]) if not isinstance(d, date) else d
    return d.year * 12 + d.month - 1


def _base_value(a) -> int:
    return int(a["GiaTriTinhKH"] or a["NguyenGia"] or 0) if a["LoaiSo"] == "TSCĐ" else int(a["NguyenGia"] or 0)


def rate(a) -> int:
    """Giá trị khấu hao / phân bổ mỗi kỳ."""
    return int(a["MucKy"]) if a["MucKy"] else div_round(_base_value(a), int(a["SoKy"]))


@dataclass
class Schedule:
    value: int
    rate: int
    periods: list = field(default_factory=list)     # [(chỉ số tháng, số tiền)] theo từng kỳ
    extra: dict = field(default_factory=dict)       # số tiền thêm theo tháng (ghi giảm CCDC, chênh lệch mốc)
    moc: tuple | None = None                        # (tháng cuối năm mốc, số kỳ còn lại tại mốc)
    so_ky: int = 0

    def amount_between(self, a: int, b: int) -> int:
        s = sum(v for m, v in self.periods if a <= m <= b)
        return s + sum(v for m, v in self.extra.items() if a <= m <= b)

    def total_to(self, m_end: int) -> int:
        return self.amount_between(-1, m_end)

    def months_between(self, a: int, b: int) -> int:
        return sum(1 for m, v in self.periods if a <= m <= b)

    def remaining_periods(self, m_end: int) -> int:
        if self.moc and m_end >= self.moc[0]:
            return max(self.moc[1] - sum(1 for m, _ in self.periods if self.moc[0] < m <= m_end), 0)
        return max(self.so_ky - sum(1 for m, _ in self.periods if m <= m_end), 0)


def _plain_periods(a, value: int, n: int, m: int, start: str) -> list:
    d = date.fromisoformat(start[:10])
    first = month_index(d)
    out, cum = [], 0
    for k in range(n):
        amt = m
        if k == 0 and a["LoaiSo"] == "TSCĐ" and d.day > 1:
            dim = calendar.monthrange(d.year, d.month)[1]
            amt = div_round(m * (dim - d.day + 1), dim)
        amt = value - cum if k == n - 1 else min(amt, value - cum)
        out.append((first + k, amt))
        cum += amt
    return out


def schedule(a) -> Schedule:
    """Lịch khấu hao / phân bổ từng kỳ của một tài sản."""
    value, n = _base_value(a), int(a["SoKy"])
    m = rate(a)
    start = (a["NgayBatDauKH"] if a["LoaiSo"] == "TSCĐ" and a["NgayBatDauKH"] else a["NgayGhiTang"])
    sc = Schedule(value, m, _plain_periods(a, value, n, m, start), so_ky=n)
    if a["MocNam"]:
        y = int(a["MocNam"])
        jan, dec = y * 12, y * 12 + 11
        con_ky = int(a["MocKyConLai"] or 0)
        luy_ke, trong_ky = int(a["MocLuyKe"] or 0), int(a["MocTrongKy"] or 0)
        before = [(mi, v) for mi, v in sc.periods if mi < jan]
        during = [(mi, v) for mi, v in sc.periods if jan <= mi <= dec]
        diff_before = (luy_ke - trong_ky) - sum(v for _, v in before)
        diff_during = trong_ky - sum(v for _, v in during)
        if diff_before:
            sc.extra[jan - 1] = sc.extra.get(jan - 1, 0) + diff_before
        if diff_during:
            sc.extra[dec] = sc.extra.get(dec, 0) + diff_during
        after, rest, cum = [], value - luy_ke, 0
        for k in range(con_ky):
            amt = rest - cum if k == con_ky - 1 else min(m, rest - cum)
            after.append((dec + 1 + k, amt))
            cum += amt
        if not con_ky and rest:
            sc.extra[dec + 1] = sc.extra.get(dec + 1, 0) + rest
        sc.periods = before + during + after
        sc.moc = (dec, con_ky)
    if a["NgayGhiGiam"]:
        g = month_index(a["NgayGhiGiam"])
        later = sum(v for mi, v in sc.periods if mi > g) + sum(v for mi, v in sc.extra.items() if mi > g)
        sc.periods = [(mi, v) for mi, v in sc.periods if mi <= g]
        sc.extra = {mi: v for mi, v in sc.extra.items() if mi <= g}
        if a["LoaiSo"] == "TSCĐ":       # TT 45/2013 điều 9: thôi trích khấu hao từ ngày tài sản giảm (tính theo ngày)
            d = date.fromisoformat(str(a["NgayGhiGiam"])[:10])
            dim = calendar.monthrange(d.year, d.month)[1]
            part = div_round(m * (d.day - 1), dim)
            sc.periods = [(mi, min(v, part)) if mi == g else (mi, v) for mi, v in sc.periods]
        if a["LoaiSo"] == "CCDC" and later:      # CCDC hỏng, mất: giá trị còn lại phân bổ hết vào tháng ghi giảm
            sc.extra[g] = sc.extra.get(g, 0) + later
    return sc


def register_row(a, year: int, sc: Schedule | None = None) -> dict:
    """Một dòng sổ của năm `year`, đủ các cột như sổ kế toán."""
    sc = sc or schedule(a)
    jan, dec = year * 12, year * 12 + 11
    r = dict(a)
    r["TrongKy"] = sc.amount_between(jan, dec)
    r["LuyKe"] = sc.total_to(dec)
    r["ConLai"] = sc.value - r["LuyKe"]
    r["KyConLai"] = sc.remaining_periods(dec)
    r["MucKy"] = sc.rate
    r["ThangTrongNam"] = sc.months_between(jan, dec)
    giam = bool(a["NgayGhiGiam"]) and a["NgayGhiGiam"] <= f"{year}-12-31"
    r["SLDaGiam"] = a["SoLuong"] if giam else 0
    r["SLConLai"] = 0 if giam else a["SoLuong"]
    r["LyDoTS"] = r["LyDo"] = "Đầu kỳ" if a["NgayGhiTang"] < f"{year}-01-01" else "Tăng mới"
    nhom = dm.BY_CODE.get(a["MaNhomDM"] or "")
    r["NhomDM"] = f"{nhom.ma} – {nhom.ten}" if nhom else ""
    for k in ("SoLuong", "SLDaGiam", "SLConLai"):
        if isinstance(r[k], float) and r[k].is_integer():
            r[k] = int(r[k])
    return r


def _in_case(r: dict, a, year: int, case: str) -> bool:
    y0, y1 = f"{year}-01-01", f"{year}-12-31"
    if a["NgayGhiTang"] > y1:
        return False
    hidden = bool(a["XoaKhoiSoNam"]) and year >= a["XoaKhoiSoNam"]
    if case == "da_xoa" or hidden:
        return case == "da_xoa" and hidden
    giam = a["NgayGhiGiam"]
    if case == "ghi_giam":
        return bool(giam) and giam <= y1
    if giam and giam < y0:
        return False
    if case == "con_kh":
        return r["ConLai"] != 0 and not (giam and giam <= y1)
    if case == "het_kh":
        return r["ConLai"] == 0
    if case == "dau_ky":
        return a["NgayGhiTang"] < y0
    if case == "tang_moi":
        return a["NgayGhiTang"] >= y0
    if case == "co_ps":
        return r["TrongKy"] != 0
    return True


# ------------------------------------------------------------------ dịch vụ
# Cảnh báo dòng hóa đơn đầu vào: mã → (nhãn, màu). Chỉ tô màu để kế toán biết, phần mềm không tự đổi phân loại.
FLAGS = {
    "chua_phan_loai": ("Chưa phân loại", "danger"),
    "goi_y": ("Có gợi ý, chờ xác nhận", "warn"),
    "thieu_thang": ("Thiếu số tháng", "warn"),
    "nguong": ("Chưa rõ CCDC hay TSCĐ", "accent"),
    "nghi": ("Nghi phân loại sai", "warn"),
    "trung": ("Nghi trùng", "danger"),
}
REVIEW_AFTER_UPGRADE = ("goi_y", "nguong", "nghi", "trung")     # rà soát dòng phần mềm đã tự phân loại


def line_flag(r: dict, dup: dict | None = None) -> tuple[str | None, str]:
    """(mã cảnh báo, nội dung) của một dòng đã đọc kèm hóa đơn; (None, "") nếu không có gì cần xem lại."""
    pl, nhom, unit = r["PhanLoai"], dm.BY_CODE.get(r.get("MaNhomDM") or ""), int(r.get("DonGiaTinh") or 0)
    if pl == dm.CHUA_PHAN_LOAI:
        if nhom:
            kq = dm.classify(r["TenHang"], unit)
            why = kq.ly_do or "Phần mềm chưa chắc chắn."
            return "goi_y", f"Gợi ý {nhom.ma} – {nhom.ten}. {why} Chọn nhóm để xác nhận."
        return "chua_phan_loai", "Không khớp danh mục: chọn phân loại và nhóm."
    if pl in dm.TAI_SAN and not r.get("SoThang"):
        return "thieu_thang", "Nhập số tháng khấu hao / phân bổ."
    if not r.get("TuDong"):
        return None, ""             # kế toán đã phân loại / bấm "Xác nhận đúng": thôi cảnh báo
    if dup is not None:
        return "trung", (f"Cùng người bán, tên hàng, số tiền, ngày với hóa đơn số {dup.get('SoHoaDonGoc')}: "
                         "kiểm tra nhập trùng.")
    if pl in (dm.TSCD, dm.CCDC) and dm.near_threshold(unit):
        return "nguong", (f"Đơn giá {unit:,.0f} đ sát ngưỡng 30 triệu: xác nhận là CCDC hay TSCĐ.").replace(",", ".")
    kq = dm.classify(r["TenHang"], unit, co_hop_dong=bool(r.get("MaHopDong")))
    if kq.phan_loai == dm.CHUA_PHAN_LOAI and kq.ly_do:
        return "nghi", kq.ly_do + " Phân loại đang có do phần mềm tự chọn."
    if kq.nhom is not None and nhom is not None and kq.nhom.ma != nhom.ma and kq.phan_loai != pl:
        return "nghi", (f"Theo danh mục hiện tại nên là {kq.nhom.ma} – {kq.nhom.ten} ({kq.phan_loai}), đang là "
                        f"{nhom.ma} ({pl}).")
    if pl == dm.CHI_PHI and unit >= dm.NGUONG_TSCD and not r.get("MaHopDong") and nhom is not None \
            and nhom.ma not in dm.SERVICE_GROUPS:
        return "nghi", "Chi phí trong kỳ có đơn giá từ 30 triệu, không gắn hợp đồng: kiểm tra có phải tài sản."
    return None, ""


class AssetService:
    def __init__(self, db: Database):
        self.db = db

    # ---- dòng hóa đơn đầu vào
    def lines_of(self, ma_hd: str) -> list:
        return self.db.active("DongHoaDonDauVao", "MaHoaDon = ?", (ma_hd,), "STT")

    def ensure_lines(self, ma_hd: str, lines: list[dict] | None = None) -> None:
        """Tạo / cập nhật các dòng hàng của một hóa đơn đầu vào rồi tự phân loại.
        lines: dòng đọc từ XML; không có thì dùng một dòng gộp từ Mặt hàng và giá trị trước thuế."""
        inv = self.db.get("HoaDonDauVao", ma_hd)
        if inv is None:
            return
        old = self.lines_of(ma_hd)
        fee_lines = [o for o in old if o["GhiChu"] == FEE_NOTE]
        old = [o for o in old if o["GhiChu"] != FEE_NOTE]
        self._sync_fee_line(inv, fee_lines, len(lines or old or [1]) + 1)
        if lines:
            if old and len(old) == len(lines) and all(o["TenHang"] == n["TenHang"] and o["ThanhTien"] == n["ThanhTien"]
                                                       for o, n in zip(old, lines)):
                return
            for o in old:
                self._drop_line(o)
            for i, ln in enumerate(lines, start=1):
                self._add_line(inv, dict(ln, STT=i))
        else:
            ten = (inv["MatHang"] or "").strip() or f"Hóa đơn {inv['SoHoaDonGoc']} – {inv['TenNhaCungCap']}"
            tong = int(inv["GiaTriTruocThue"] or 0)
            if not old:
                self._add_line(inv, {"STT": 1, "TenHang": ten, "ThanhTien": tong, "ThueSuat": inv["ThueSuat"],
                                     "GhiChu": "Dòng gộp (hóa đơn không có chi tiết từng mặt hàng)"})
            elif len(old) == 1 and (old[0]["TenHang"] != ten or old[0]["ThanhTien"] != tong) \
                    and old[0]["SoLuong"] is None:
                o = old[0]
                vals = {"TenHang": ten, "ThanhTien": tong}
                if o["TuDong"]:
                    vals.update(self._auto_class(inv, ten, tong, None, None))
                self.db.update("DongHoaDonDauVao", o["Ma"], vals, commit=False)
                self._sync_asset(self.db.get("DongHoaDonDauVao", o["Ma"]), inv)
        self.db.conn.commit()

    def _sync_fee_line(self, inv, fee_lines: list, stt: int) -> None:
        """Phí ngoài VAT của hóa đơn (VD phí bảo vệ môi trường trên hóa đơn nước): một dòng riêng để tính chi phí."""
        fee = int(inv["PhiKhac"] or 0) if "PhiKhac" in inv.keys() else 0
        ten = (inv["TenPhiKhac"] or "").strip() or "Phí bảo vệ môi trường"
        if not fee:
            for o in fee_lines:
                self._drop_line(o)
            return
        if fee_lines:
            o = fee_lines[0]
            for extra in fee_lines[1:]:
                self._drop_line(extra)
            if o["ThanhTien"] != fee or o["TenHang"] != ten:
                vals = {"TenHang": ten, "ThanhTien": fee}
                if o["TuDong"]:
                    vals.update(self._auto_class(inv, ten, fee, None, None))
                self.db.update("DongHoaDonDauVao", o["Ma"], vals, commit=False)
            return
        self._add_line(inv, {"STT": stt, "TenHang": ten, "ThanhTien": fee, "ThueSuat": "KCT", "GhiChu": FEE_NOTE})

    def backfill(self) -> int:
        """Hóa đơn đầu vào nhập trước khi có module này: tạo dòng gộp để phân loại."""
        rows = self.db.query("SELECT Ma FROM HoaDonDauVao h WHERE DaXoa = 0 AND NOT EXISTS "
                             "(SELECT 1 FROM DongHoaDonDauVao d WHERE d.MaHoaDon = h.Ma AND d.DaXoa = 0)")
        for r in rows:
            self.ensure_lines(r["Ma"])
        self.renumber_ccdc()                 # kể cả tài sản nhận từ máy khác qua gói đồng bộ
        return len(rows)

    def _auto_class(self, inv, ten: str, thanh_tien: int, so_luong, don_gia) -> dict:
        unit = don_gia or (int(thanh_tien / so_luong) if so_luong else thanh_tien)
        if so_luong is None and ";" in ten:     # dòng gộp nhiều mặt hàng: để kế toán tách, phân loại
            kq = dm.KetQua(dm.CHUA_PHAN_LOAI, None, None)
        else:
            kq = dm.classify(ten, unit, co_hop_dong=bool(inv["MaHopDong"]))
        return {"PhanLoai": kq.phan_loai, "MaNhomDM": kq.nhom.ma if kq.nhom else None, "SoThang": kq.so_thang,
                "TuDong": 1}

    def _add_line(self, inv, ln: dict) -> str:
        v = {"MaHoaDon": inv["Ma"], "STT": ln.get("STT", 1), "TenHang": ln["TenHang"], "DVT": ln.get("DVT") or None,
             "SoLuong": ln.get("SoLuong"), "DonGia": ln.get("DonGia") or None,
             "ThanhTien": int(ln.get("ThanhTien") or 0), "ThueSuat": ln.get("ThueSuat") or None,
             "GhiChu": ln.get("GhiChu")}
        v.update(self._auto_class(inv, v["TenHang"], v["ThanhTien"], v["SoLuong"], v["DonGia"]))
        ma = self.db.insert("DongHoaDonDauVao", v, commit=False)
        self._sync_asset(self.db.get("DongHoaDonDauVao", ma), inv)
        return ma

    def _drop_line(self, line) -> None:
        for ts in self.db.active("TaiSan", "MaDongHD = ?", (line["Ma"],)):
            self.db.soft_delete("TaiSan", ts["Ma"], commit=False)
        self.db.soft_delete("DongHoaDonDauVao", line["Ma"], commit=False)

    def delete_invoice_lines(self, ma_hd: str) -> None:
        for ln in self.lines_of(ma_hd):
            self._drop_line(ln)
        self.db.conn.commit()

    @staticmethod
    def unit_price(line) -> int:
        if line["DonGia"]:
            return int(line["DonGia"])
        return int(line["ThanhTien"] / line["SoLuong"]) if line["SoLuong"] else int(line["ThanhTien"] or 0)

    def is_pending(self, line) -> bool:
        return line["PhanLoai"] == dm.CHUA_PHAN_LOAI or (line["PhanLoai"] in dm.TAI_SAN and not line["SoThang"])

    def _line_rows(self, where: str = "", params=()) -> list[dict]:
        sql = ("SELECT d.*, h.SoHoaDonGoc, h.KyHieuHD, h.NgayNhanHD, h.TenNhaCungCap, h.MaThueNCC, h.MaHopDong "
               "FROM DongHoaDonDauVao d JOIN HoaDonDauVao h ON h.Ma = d.MaHoaDon "
               "WHERE d.DaXoa = 0 AND h.DaXoa = 0")
        if where:
            sql += f" AND ({where})"
        out = []
        for r in self.db.query(sql + " ORDER BY h.NgayNhanHD, h.SoHoaDonGoc, d.STT", params):
            r = dict(r)
            nhom = dm.BY_CODE.get(r["MaNhomDM"] or "")
            r["NhomDM"] = f"{nhom.ma} – {nhom.ten}" if nhom else ""
            r["DonGiaTinh"] = self.unit_price(r)
            r["CanPhanLoai"] = self.is_pending(r)
            out.append(r)
        self._flag(out)
        return out

    def _flag(self, rows: list[dict]) -> None:
        """Gắn cảnh báo cho từng dòng (chỉ tô màu, không đổi dữ liệu): MaCanhBao, CanhBao, _mau."""
        seen: dict[tuple, dict] = {}
        for r in rows:
            # hóa đơn định kỳ (nước, xăng hằng tháng) cùng số tiền nhưng khác ngày thì không phải trùng
            key = (r.get("MaThueNCC") or r.get("TenNhaCungCap"), dm.normalize(r["TenHang"]), int(r["ThanhTien"] or 0),
                   r.get("NgayNhanHD"))
            other = seen.get(key)
            dup = other if other and other["MaHoaDon"] != r["MaHoaDon"] and key[2] else None
            seen.setdefault(key, r)
            code, text = line_flag(r, dup)
            r["MaCanhBao"], r["CanhBao"] = code, text
            r["_mau"] = FLAGS[code][1] if code else None

    def review_lines(self, codes: tuple[str, ...] | None = None) -> list[dict]:
        """Dòng cần kế toán xem lại: chưa phân loại, có gợi ý chờ xác nhận, thiếu số tháng, sát ngưỡng CCDC/TSCĐ,
        nghi phân loại sai, nghi trùng."""
        return [r for r in self._line_rows() if r["MaCanhBao"] and (codes is None or r["MaCanhBao"] in codes)]

    def review_count(self) -> int:
        return len(self.review_lines())

    def confirm_lines(self, mas: list[str]) -> int:
        """Kế toán xác nhận phân loại phần mềm tự chọn là đúng: thôi cảnh báo, giữ nguyên phân loại (thành phân loại
        tay, được ưu tiên khi đồng bộ). Dòng chưa phân loại hoặc thiếu số tháng thì phải phân loại, không xác nhận được."""
        n = 0
        for ma in mas:
            ln = self.db.get("DongHoaDonDauVao", ma)
            if ln is None or ln["DaXoa"] or self.is_pending(ln):
                continue
            self.db.update("DongHoaDonDauVao", ma, {"TuDong": 0}, commit=False)
            n += 1
        self.db.conn.commit()
        return n

    def pending_lines(self) -> list[dict]:
        return [r for r in self._line_rows("d.PhanLoai IN ('Chưa phân loại', 'TSCĐ', 'CCDC', 'Chờ phân bổ')")
                if r["CanPhanLoai"]]

    def pending_count(self) -> int:
        return self.db.one(
            "SELECT COUNT(*) n FROM DongHoaDonDauVao d JOIN HoaDonDauVao h ON h.Ma = d.MaHoaDon "
            "WHERE d.DaXoa = 0 AND h.DaXoa = 0 AND (d.PhanLoai = 'Chưa phân loại' OR "
            "(d.PhanLoai IN ('TSCĐ', 'CCDC', 'Chờ phân bổ') AND COALESCE(d.SoThang, 0) = 0))")["n"]

    def lines(self, tu: str | None = None, den: str | None = None, phan_loai: str | None = None) -> list[dict]:
        where, p = ["h.NgayNhanHD BETWEEN ? AND ?"], [tu or "0000-01-01", den or "9999-12-31"]
        if phan_loai:
            where.append("d.PhanLoai = ?")
            p.append(phan_loai)
        return self._line_rows(" AND ".join(where), p)

    def set_line_class(self, ma_dong: str, phan_loai: str, ma_nhom: str | None, so_thang: int | None) -> None:
        """Kế toán phân loại một dòng: nhóm chỉ chọn trong danh mục cài sẵn; thời gian nhập tự do."""
        line = self.db.get("DongHoaDonDauVao", ma_dong)
        if line is None or line["DaXoa"]:
            raise AssetError("Không tìm thấy dòng hóa đơn.")
        if phan_loai not in dm.PHAN_LOAI or phan_loai == dm.CHUA_PHAN_LOAI:
            raise AssetError("Hãy chọn phân loại.")
        nhom = dm.BY_CODE.get(ma_nhom or "")
        if nhom is None:
            raise AssetError("Hãy chọn nhóm trong danh mục (không tự nhập nhóm mới).")
        if phan_loai not in dm.allowed_classes(nhom, self.unit_price(line)):
            raise AssetError(f"Nhóm {nhom.ma} không xếp vào \"{phan_loai}\" được. Chọn: "
                             + ", ".join(dm.allowed_classes(nhom, self.unit_price(line))) + ".")
        if phan_loai in dm.TAI_SAN:
            if not so_thang or not 1 <= int(so_thang) <= 600:
                raise AssetError("Nhập thời gian khấu hao / phân bổ (số tháng, từ 1 đến 600).")
        else:
            so_thang = None
        self.db.update("DongHoaDonDauVao", ma_dong, {"PhanLoai": phan_loai, "MaNhomDM": nhom.ma,
                                                     "SoThang": int(so_thang) if so_thang else None, "TuDong": 0},
                       commit=False)
        inv = self.db.get("HoaDonDauVao", line["MaHoaDon"])
        self._sync_asset(self.db.get("DongHoaDonDauVao", ma_dong), inv)
        self.renumber_ccdc(commit=False)
        self.db.conn.commit()

    def set_invoice_class(self, ma_hd: str, phan_loai: str, ma_nhom: str, so_thang: int | None) -> int:
        """Phân loại mọi dòng hàng của một hóa đơn theo loại người dùng ghi sẵn (VD cột Loại hàng hóa trong file Excel);
        dòng phí ngoài VAT giữ phân loại riêng. Dòng nào không xếp được thì để nguyên chờ kế toán."""
        n = 0
        for ln in self.lines_of(ma_hd):
            if ln["GhiChu"] == FEE_NOTE:
                continue
            try:
                if phan_loai in dm.TAI_SAN and not so_thang:      # thiếu số tháng: ghi loại, nhóm; vẫn hiện chờ nhập tháng
                    self.db.update("DongHoaDonDauVao", ln["Ma"], {"PhanLoai": phan_loai, "MaNhomDM": ma_nhom,
                                                                 "SoThang": None, "TuDong": 0})
                else:
                    self.set_line_class(ln["Ma"], phan_loai, ma_nhom, so_thang)
                n += 1
            except AssetError:
                pass
        return n

    @staticmethod
    def month_warning(phan_loai: str, ma_nhom: str | None, so_thang: int | None) -> str:
        """Cảnh báo (không chặn) khi thời gian nằm ngoài khung."""
        if not so_thang:
            return ""
        if phan_loai in (dm.CCDC, dm.CHO_PB) and so_thang > 36:
            return "CCDC, chi phí chờ phân bổ thường phân bổ tối đa 36 tháng."
        nhom = dm.BY_CODE.get(ma_nhom or "")
        if phan_loai == dm.TSCD and nhom and nhom.khung:
            lo, hi = (int(x) for x in nhom.khung.replace("năm", "").strip().split("–"))
            if not lo * 12 <= so_thang <= hi * 12:
                return f"Khung khấu hao nhóm {nhom.nhom_khung}: {nhom.khung} ({lo * 12}–{hi * 12} tháng)."
        return ""

    # ---- tài sản
    def renumber_ccdc(self, commit: bool = True) -> int:
        """Đánh lại mã CCDC phần mềm tự cấp (CC001/2026…) cho liên tục, theo thứ tự ngày ghi tăng trong từng năm.
        Giữ nguyên mã của tài sản nhập từ sổ kế toán (có mốc sổ) và mã người dùng tự đặt khác dạng CCnnn/năm;
        mã tự cấp nối tiếp sau số lớn nhất của sổ kế toán trong năm đó. Trả về số tài sản đổi mã."""
        rows = [dict(r) for r in self.db.active("TaiSan", "LoaiSo = 'CCDC'")]
        locked: dict[str, int] = {}
        auto: dict[str, list[dict]] = {}
        for r in rows:
            m = CCDC_CODE.fullmatch(r["MaTS"] or "")
            if r["MocNam"] is not None:
                if m:
                    locked[m.group(2)] = max(locked.get(m.group(2), 0), int(m.group(1)))
            elif m or not r["MaTS"]:
                auto.setdefault(str(r["NgayGhiTang"])[:4], []).append(r)
        changed = 0
        for year, items in auto.items():
            def key(r):
                m = CCDC_CODE.fullmatch(r["MaTS"] or "")
                return (str(r["NgayGhiTang"]), int(m.group(1)) if m else 10 ** 6, r["Ma"])
            for i, r in enumerate(sorted(items, key=key), start=locked.get(year, 0) + 1):
                code = f"CC{i:03d}/{year}"
                if r["MaTS"] != code:
                    v = {"MaTS": code}
                    if not r["SoCTGhiTang"] or r["SoCTGhiTang"] == r["MaTS"]:
                        v["SoCTGhiTang"] = code
                    self.db.update("TaiSan", r["Ma"], v, commit=False)
                    changed += 1
        if commit and changed:
            self.db.conn.commit()
        return changed

    def _next_code(self, loai_so: str, year: str) -> tuple[str, str]:
        if loai_so == "CCDC":
            n = 0
            for r in self.db.query("SELECT MaTS, SoCTGhiTang FROM TaiSan WHERE LoaiSo = 'CCDC' AND DaXoa = 0"):
                for code in (r["MaTS"], r["SoCTGhiTang"]):
                    if code and code.startswith("CC") and code.endswith("/" + year) and code[2:-5].isdigit():
                        n = max(n, int(code[2:-5]))
            code = f"CC{n + 1:03d}/{year}"
            return code, code
        n, k = 0, 0
        for r in self.db.query("SELECT MaTS, SoCTGhiTang FROM TaiSan WHERE LoaiSo = 'TSCĐ'"):
            ct = r["SoCTGhiTang"] or ""
            if ct.startswith("GTTS") and ct[4:].isdigit():
                k = max(k, int(ct[4:]))
            code = r["MaTS"] or ""
            if code.startswith("TS") and code.endswith("/" + year) and code[2:-5].isdigit():
                n = max(n, int(code[2:-5]))
        return f"TS{n + 1:03d}/{year}", f"GTTS{k + 1:05d}"

    def _sync_asset(self, line, inv) -> None:
        ts = self.db.one("SELECT * FROM TaiSan WHERE MaDongHD = ? AND DaXoa = 0", (line["Ma"],))
        if line["DaXoa"] or line["PhanLoai"] not in dm.TAI_SAN or not line["SoThang"]:
            if ts:
                self.db.soft_delete("TaiSan", ts["Ma"], commit=False)
            return
        loai_so = "TSCĐ" if line["PhanLoai"] == dm.TSCD else "CCDC"
        ngay = inv["NgayNhanHD"][:10]
        v = {"LoaiSo": loai_so, "PhanLoai": line["PhanLoai"], "MaNhomDM": line["MaNhomDM"],
             "SoKy": int(line["SoThang"]), "NguyenGia": int(line["ThanhTien"]), "GiaTriTinhKH": int(line["ThanhTien"])}
        if ts is None or ts["LoaiSo"] != loai_so:
            ma_ts, so_ct = self._next_code(loai_so, ngay[:4])
            if loai_so == "TSCĐ":
                loai_ts, tk_ng, tk_kh = dm.fixed_asset_kind(line["MaNhomDM"])
            else:
                loai_ts, tk_ng, tk_kh = None, None, "242"
            v.update(MaTS=ma_ts, SoCTGhiTang=so_ct, LoaiTS=loai_ts, TKNguyenGia=tk_ng, TKKhauHao=tk_kh)
        if ts is None:
            v.update(TenTS=line["TenHang"], NgayGhiTang=ngay, NgayBatDauKH=ngay, DVT=line["DVT"] or "Cái",
                     SoLuong=line["SoLuong"] or 1, MaDongHD=line["Ma"])
            self.db.insert("TaiSan", v, commit=False)
        else:
            self.db.update("TaiSan", ts["Ma"], v, commit=False)

    def assets(self, loai_so: str | None = None) -> list:
        where, p = ("LoaiSo = ?", (loai_so,)) if loai_so else ("", ())
        return self.db.active("TaiSan", where, p, "NgayGhiTang, ThuTu, MaTS")

    def save_asset(self, values: dict, ma: str | None = None) -> str:
        """Thêm tay / sửa một tài sản trong sổ."""
        v = {k: values[k] for k in values if k in self.db.columns("TaiSan") and k != "Ma"}
        cur = dict(self.db.get("TaiSan", ma)) if ma else {}
        merged = {**cur, **v}
        for f, label in (("TenTS", "Tên tài sản"), ("NgayGhiTang", "Ngày ghi tăng")):
            if not merged.get(f):
                raise AssetError(f"Chưa nhập {label}.")
        if not merged.get("SoKy") or int(merged["SoKy"]) < 1:
            raise AssetError("Số kỳ (tháng) khấu hao / phân bổ phải lớn hơn 0.")
        if int(merged.get("NguyenGia") or 0) <= 0:
            raise AssetError("Chưa nhập nguyên giá / giá trị.")
        if merged.get("MaNhomDM") and merged["MaNhomDM"] not in dm.BY_CODE:
            raise AssetError("Nhóm phải chọn trong danh mục.")
        loai_so = merged.get("LoaiSo") or "CCDC"
        v["LoaiSo"] = loai_so
        if loai_so == "TSCĐ":
            v["PhanLoai"] = "TSCĐ"
            if not merged.get("GiaTriTinhKH"):
                v["GiaTriTinhKH"] = merged["NguyenGia"]
            v.setdefault("NgayBatDauKH", merged.get("NgayBatDauKH") or merged["NgayGhiTang"])
        else:
            v["PhanLoai"] = merged.get("PhanLoai") if merged.get("PhanLoai") in (dm.CCDC, dm.CHO_PB) else dm.CCDC
            v["GiaTriTinhKH"] = merged["NguyenGia"]
            v.setdefault("TKKhauHao", merged.get("TKKhauHao") or "242")
        if not merged.get("MaTS"):
            v["MaTS"], ct = self._next_code(loai_so, str(merged["NgayGhiTang"])[:4])
            v.setdefault("SoCTGhiTang", merged.get("SoCTGhiTang") or ct)
        if ma:
            self.db.update("TaiSan", ma, v, commit=False)
        else:
            ma = self.db.insert("TaiSan", v, commit=False)
        self.renumber_ccdc(commit=False)
        self.db.conn.commit()
        return ma

    def write_off(self, ma: str, ngay: str, ly_do: str = "") -> None:
        """Ghi giảm (thanh lý, hỏng, mất). CCDC: giá trị còn lại phân bổ hết vào tháng ghi giảm."""
        a = self.db.get("TaiSan", ma)
        if a is None:
            raise AssetError("Không tìm thấy tài sản.")
        if ngay < a["NgayGhiTang"]:
            raise AssetError("Ngày ghi giảm phải sau ngày ghi tăng.")
        self.db.update("TaiSan", ma, {"NgayGhiGiam": ngay, "LyDoGhiGiam": ly_do or None})

    def undo_write_off(self, ma: str) -> None:
        self.db.update("TaiSan", ma, {"NgayGhiGiam": None, "LyDoGhiGiam": None})

    def delete_asset(self, ma: str) -> None:
        a = self.db.get("TaiSan", ma)
        if a and a["MaDongHD"]:
            raise AssetError("Tài sản này sinh ra từ hóa đơn đầu vào. Muốn bỏ thì phân loại lại dòng hóa đơn "
                             "(ví dụ thành Chi phí trong kỳ).")
        self.db.soft_delete("TaiSan", ma, commit=False)
        self.renumber_ccdc(commit=False)
        self.db.conn.commit()

    def fully_depreciated(self, loai_so: str, year: int) -> list[dict]:
        """Tài sản đã hết khấu hao / phân bổ từ trước năm `year` (cả năm không còn số phát sinh), còn đang hiện
        trong sổ năm đó: có thể xóa khỏi sổ. Tài sản hết khấu hao trong năm thì sang năm sau mới xóa được,
        để sổ năm có phát sinh vẫn đủ dòng."""
        out = []
        for a in self.assets(loai_so):
            r = register_row(a, year)
            if _in_case(r, a, year, "het_kh") and r["TrongKy"] == 0 and a["NgayGhiTang"] < f"{year}-01-01":
                out.append(r)
        return out

    def remove_depreciated(self, mas: list[str], year: int) -> int:
        """Xóa khỏi sổ từ năm `year` trở đi các tài sản đã hết khấu hao; sổ các năm trước và bảng lãi lỗ giữ
        nguyên. Khôi phục được (restore_to_register)."""
        ok = {r["Ma"] for loai in ("TSCĐ", "CCDC") for r in self.fully_depreciated(loai, year)}
        bad = [m for m in mas if m not in ok]
        if bad:
            raise AssetError("Chỉ xóa được tài sản đã hết khấu hao / phân bổ từ trước năm đang xem.")
        for m in mas:
            self.db.update("TaiSan", m, {"XoaKhoiSoNam": year}, commit=False)
        self.db.conn.commit()
        return len(mas)

    def restore_to_register(self, ma: str) -> None:
        self.db.update("TaiSan", ma, {"XoaKhoiSoNam": None})

    # ---- sổ theo năm
    def register(self, loai_so: str, year: int, case: str = "toan_bo") -> list[dict]:
        out = []
        invoices = {r["Ma"]: r for r in self.db.query(
            "SELECT d.Ma, h.KyHieuHD, h.SoHoaDonGoc, h.TenNhaCungCap FROM DongHoaDonDauVao d "
            "JOIN HoaDonDauVao h ON h.Ma = d.MaHoaDon")}
        for a in self.assets(loai_so):
            r = register_row(a, year)
            if not _in_case(r, a, year, case):
                continue
            h = invoices.get(a["MaDongHD"] or "")
            r["HoaDon"] = f"{h['KyHieuHD'] or ''} số {h['SoHoaDonGoc']} – {h['TenNhaCungCap']}".strip() if h else ""
            out.append(r)
        return out

    def year_totals(self, year: int) -> dict:
        """Khấu hao TSCĐ, phân bổ CCDC (gồm chờ phân bổ) trong năm, theo quý – đưa vào bảng lãi lỗ."""
        res = {"TSCĐ": [0, 0, 0, 0], "CCDC": [0, 0, 0, 0]}
        for a in self.assets():
            sc = schedule(a)
            for q in range(4):
                a0 = year * 12 + q * 3
                res[a["LoaiSo"]][q] += sc.amount_between(a0, a0 + 2)
        return res

    def expense_lines(self, year: int) -> list[dict]:
        """Dòng hóa đơn đầu vào là chi phí trong kỳ của năm (theo ngày hóa đơn)."""
        rows = self.lines(f"{year}-01-01", f"{year}-12-31", dm.CHI_PHI)
        for r in rows:
            nhom = dm.BY_CODE.get(r["MaNhomDM"] or "")
            r["DongLaiLo"] = nhom.dong if nhom else ""
            r["Quy"] = (int(r["NgayNhanHD"][5:7]) - 1) // 3 + 1
        return rows

    def year_check(self, year: int) -> int:
        """Số dòng hóa đơn trong năm còn chờ kế toán phân loại (chưa chốt năm được)."""
        return sum(1 for r in self.pending_lines() if r["NgayNhanHD"][:4] == str(year))

    # ---- nhập sổ cũ từ Excel
    def read_register(self, path: str | Path, year: int | None = None) -> dict:
        """Đọc sổ TSCĐ / CCDC từ Excel và xếp từng dòng: mới, trùng mã (cập nhật), có thể trùng, trùng trong file."""
        import openpyxl

        wb = openpyxl.load_workbook(str(path), data_only=True)
        ws = wb.active
        head_row, kind = None, None
        for r in range(1, min(ws.max_row, 15) + 1):
            first = str(ws.cell(r, 1).value or "").strip().lower()
            if first.startswith("mã tscđ") or first.startswith("mã tscd"):
                head_row, kind = r, "TSCĐ"
                break
            if first.startswith("mã ccdc"):
                head_row, kind = r, "CCDC"
                break
        if head_row is None:
            raise AssetError("Không nhận ra sổ: cần dòng tiêu đề bắt đầu bằng \"Mã TSCĐ\" hoặc \"Mã CCDC\" "
                             "(sổ xuất từ phần mềm kế toán).")
        if year is None:
            for r in range(1, head_row):
                for c in range(1, 4):
                    s = str(ws.cell(r, c).value or "")
                    digits = "".join(ch for ch in s if ch.isdigit())
                    if "năm" in s.lower() and len(digits) == 4:
                        year = int(digits)
            if year is None:
                raise AssetError("Không thấy năm của sổ (dòng \"Năm 2025\" phía trên tiêu đề).")
        res = {"loai": kind, "nam": year, "rows": [], "loi": []}
        seen: list[dict] = []
        for r in range(head_row + 1, ws.max_row + 1):
            row = [ws.cell(r, c).value for c in range(1, 19)]
            if not row[0] or str(row[0]).startswith("Số dòng"):
                continue
            try:
                v = self._from_book_row(kind, row, year)
            except (TypeError, ValueError) as e:
                res["loi"].append(f"Dòng {r} ({row[0]}): {e}")
                continue
            it = {"dong": r, "values": v, "trung_voi": "", "ma_cu": None}
            twin = next((s for s in seen if s["values"]["MaTS"] == v["MaTS"] or _same_asset(s["values"], v)), None)
            old = self.db.one("SELECT * FROM TaiSan WHERE DaXoa = 0 AND LoaiSo = ? AND MaTS = ? AND NgayGhiTang = ?",
                              (kind, v["MaTS"], v["NgayGhiTang"]))
            if twin is not None:
                it.update(status=ST_IN_FILE, action=A_SKIP, trung_voi=f"Dòng {twin['dong']} trong file: "
                          f"{twin['values']['MaTS']} – {twin['values']['TenTS']}")
            elif old is not None:
                it.update(status=ST_SAME_CODE, action=A_UPDATE, ma_cu=old["Ma"],
                          trung_voi=f"{old['MaTS']} – {old['TenTS']} (đã có, cùng mã và ngày ghi tăng)")
            else:
                dups = self.find_duplicates(v) or self.db.active("TaiSan", "LoaiSo = ? AND MaTS = ?",
                                                                 (kind, v["MaTS"]))
                if dups:
                    d = dups[0]
                    it.update(status=ST_MAYBE, action=A_SKIP, ma_cu=d["Ma"],
                              trung_voi=f"{d['MaTS']} – {d['TenTS']}; ghi tăng {_dmy(d['NgayGhiTang'])}; "
                                        f"giá trị {d['NguyenGia']:,} đ".replace(",", "."))
                else:
                    it.update(status=ST_NEW, action=A_ADD)
            seen.append(it)
            res["rows"].append(it)
        return res

    def apply_register(self, data: dict, actions: dict[int, str] | None = None) -> dict:
        """Ghi các dòng đã đọc theo cách xử lý từng dòng (mặc định: mới thì thêm, trùng mã thì cập nhật,
        có thể trùng / trùng trong file thì bỏ qua). actions: {số dòng Excel: 'them' / 'cap_nhat' / 'bo_qua'}."""
        kind, year = data["loai"], data["nam"]
        res = {"loai": kind, "nam": year, "them": 0, "cap_nhat": 0, "bo_qua": 0, "lech": [], "loi": list(data["loi"])}
        order = (self.db.one("SELECT MAX(ThuTu) m FROM TaiSan")["m"] or 0)
        for it in data["rows"]:
            act = (actions or {}).get(it["dong"], it["action"])
            v = dict(it["values"])
            if act == A_SKIP or (act == A_UPDATE and not it["ma_cu"]):
                res["bo_qua"] += 1
                continue
            plain = dict(v, MocNam=None, NgayGhiGiam=None)
            pr = register_row(plain, year)
            if (pr["TrongKy"], pr["LuyKe"], pr["KyConLai"]) != (v["MocTrongKy"], v["MocLuyKe"], v["MocKyConLai"]):
                res["lech"].append(f"{v['MaTS']}: sổ ghi trong kỳ {v['MocTrongKy']:,}, lũy kế {v['MocLuyKe']:,}; "
                                   f"cách tính chuẩn ra {pr['TrongKy']:,}, {pr['LuyKe']:,}. Giữ theo sổ."
                                   .replace(",", "."))
            if act == A_UPDATE:
                self.db.update("TaiSan", it["ma_cu"], v, commit=False)
                res["cap_nhat"] += 1
            else:
                order += 1
                v["ThuTu"] = order
                self.db.insert("TaiSan", v, commit=False)
                res["them"] += 1
        self.db.conn.commit()
        return res

    def read_extra_ccdc(self, path: str | Path) -> dict:
        """Đọc file kê khai CCDC nhập thêm ngoài hóa đơn (mẫu của phần mềm), xếp dòng như khi nhập sổ cũ.
        Khác sổ kế toán: không có mốc số liệu, phần mềm tự tính phân bổ và tự đánh mã."""
        from . import ccdc_excel

        try:
            items, errors = ccdc_excel.read_rows(path)
        except ValueError as e:
            raise AssetError(str(e)) from None
        years = sorted({v["NgayGhiTang"][:4] for _, v in items})
        res = {"loai": "CCDC", "nam": int(years[-1]) if years else date.today().year, "rows": [], "loi": errors,
               "ke_khai": True}
        seen: list[dict] = []
        for line, v in items:
            it = {"dong": line, "values": v, "trung_voi": "", "ma_cu": None}
            twin = next((s for s in seen if (v["MaTS"] and s["values"]["MaTS"] == v["MaTS"])
                         or _same_asset(s["values"], v)), None)
            old = self.db.one("SELECT * FROM TaiSan WHERE DaXoa = 0 AND LoaiSo = 'CCDC' AND MaTS = ?",
                              (v["MaTS"],)) if v["MaTS"] else None
            dups = self.find_duplicates(v)
            if twin is not None:
                it.update(status=ST_IN_FILE, action=A_SKIP, trung_voi=f"Dòng {twin['dong']} trong file: "
                          f"{twin['values']['TenTS']}")
            elif old is not None:
                it.update(status=ST_SAME_CODE, action=A_UPDATE, ma_cu=old["Ma"],
                          trung_voi=f"{old['MaTS']} – {old['TenTS']} (đã có mã này)")
            elif dups:
                d = dups[0]
                it.update(status=ST_MAYBE, action=A_SKIP, ma_cu=d["Ma"],
                          trung_voi=f"{d['MaTS']} – {d['TenTS']}; ghi tăng {_dmy(d['NgayGhiTang'])}; "
                                    f"giá trị {d['NguyenGia']:,} đ".replace(",", "."))
            else:
                it.update(status=ST_NEW, action=A_ADD)
            seen.append(it)
            res["rows"].append(it)
        return res

    def apply_extra_ccdc(self, data: dict, actions: dict[int, str] | None = None) -> dict:
        """Ghi các dòng kê khai đã đọc: thêm mới hoặc cập nhật bản đã có; mã CCDC tự đánh lại sau khi ghi."""
        res = {"loai": "CCDC", "nam": data["nam"], "them": 0, "cap_nhat": 0, "bo_qua": 0, "lech": [],
               "loi": list(data["loi"])}
        for it in data["rows"]:
            act = (actions or {}).get(it["dong"], it["action"])
            v = dict(it["values"])
            if act == A_SKIP or (act == A_UPDATE and not it["ma_cu"]):
                res["bo_qua"] += 1
                continue
            try:
                if act == A_UPDATE:
                    if not v["MaTS"]:
                        v.pop("MaTS")
                    self.save_asset(v, it["ma_cu"])
                    res["cap_nhat"] += 1
                else:
                    self.save_asset(v)
                    res["them"] += 1
            except AssetError as e:
                res["loi"].append(f"Dòng {it['dong']} ({v['TenTS']}): {e}")
        return res

    def import_register(self, path: str | Path, year: int | None = None) -> dict:
        """Đọc và ghi luôn theo cách xử lý mặc định."""
        return self.apply_register(self.read_register(path, year))

    def find_duplicates(self, v: dict, exclude: str | None = None) -> list:
        """Tài sản đã có trong sổ có thể là cùng một món: cùng ngày ghi tăng và cùng tên (bỏ dấu, hoa thường)
        hoặc cùng giá trị."""
        rows = self.db.active("TaiSan", "LoaiSo = ? AND NgayGhiTang = ? AND Ma != ?",
                              (v.get("LoaiSo") or "CCDC", str(v["NgayGhiTang"])[:10], exclude or ""))
        return [r for r in rows if _same_asset(dict(r), v)]

    def duplicate_warnings(self, loai_so: str) -> dict[str, str]:
        """{mã bản ghi: "Có thể trùng với …"} cho các tài sản trong sổ nghi trùng nhau."""
        out: dict[str, str] = {}
        groups: dict[str, list] = {}
        for a in self.assets(loai_so):
            groups.setdefault(a["NgayGhiTang"], []).append(a)
        for same_day in groups.values():
            for i, a in enumerate(same_day):
                for b in same_day[i + 1:]:
                    if _same_asset(dict(a), dict(b)):
                        out.setdefault(a["Ma"], f"Có thể trùng với {b['MaTS']} – {b['TenTS']}")
                        out.setdefault(b["Ma"], f"Có thể trùng với {a['MaTS']} – {a['TenTS']}")
        return out

    @staticmethod
    def _from_book_row(kind: str, row: list, year: int) -> dict:
        def day(x):
            if isinstance(x, datetime):
                return x.date().isoformat()
            if isinstance(x, date):
                return x.isoformat()
            s = str(x).strip()
            if "/" in s:
                d, m, y = s.split("/")[:3]
                return date(int(y[:4]), int(m), int(d)).isoformat()
            return date.fromisoformat(s[:10]).isoformat()

        def num(x):
            if x in (None, ""):
                return 0
            return int(round(float(str(x).replace(",", "")) if isinstance(x, str) else float(x)))

        def text(x):
            return None if x in (None, "") else str(x).strip()

        nhom = dm.find_group(str(row[1] or ""))
        if kind == "TSCĐ":
            v = {"LoaiSo": "TSCĐ", "PhanLoai": "TSCĐ", "MaTS": text(row[0]), "TenTS": text(row[1]),
                 "LoaiTS": text(row[2]), "DonViSuDung": text(row[3]), "NgayGhiTang": day(row[4]),
                 "SoCTGhiTang": text(row[5]), "NgayBatDauKH": day(row[6] or row[4]), "SoKy": num(row[7]),
                 "MocKyConLai": num(row[8]), "NguyenGia": num(row[9]), "GiaTriTinhKH": num(row[10]) or num(row[9]),
                 "MocTrongKy": num(row[11]), "MocLuyKe": num(row[12]), "MucKy": num(row[14]) or None,
                 "TKNguyenGia": text(row[15]), "TKKhauHao": text(row[16]), "SoLuong": 1}
        else:
            cho_pb = nhom is not None and nhom.kieu == "CHO_PB"
            v = {"LoaiSo": "CCDC", "PhanLoai": dm.CHO_PB if cho_pb else dm.CCDC, "MaTS": text(row[0]),
                 "TenTS": text(row[1]), "LoaiTS": text(row[2]), "NgayGhiTang": day(row[4]),
                 "SoCTGhiTang": text(row[5]), "SoKy": num(row[6]), "MocKyConLai": num(row[7]), "DVT": text(row[8]),
                 "SoLuong": float(row[9] or 1), "NguyenGia": num(row[12]), "GiaTriTinhKH": num(row[12]),
                 "MucKy": num(row[13]) or None, "MocTrongKy": num(row[14]), "MocLuyKe": num(row[15]),
                 "TKKhauHao": text(row[17]) or "242"}
            v["NgayBatDauKH"] = v["NgayGhiTang"]
        if not v["MaTS"] or not v["TenTS"]:
            raise ValueError("thiếu mã hoặc tên")
        if v["SoKy"] < 1:
            raise ValueError("số kỳ phải lớn hơn 0")
        v["MaNhomDM"] = nhom.ma if nhom else None
        v["MocNam"] = year
        v["GhiChu"] = f"Nhập từ sổ năm {year}"
        return v

    # ---- xuất Excel
    def export_register(self, path: str | Path, loai_so: str, year: int, case: str = "toan_bo",
                        extra: bool = False) -> int:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        rows = self.register(loai_so, year, case)
        cols, keys, sums = ((TSCD_COLS, TSCD_KEYS, TSCD_SUM) if loai_so == "TSCĐ" else (CCDC_COLS, CCDC_KEYS, CCDC_SUM))
        heads = list(cols) + ([t for _, t in EXTRA_COLS] if extra else [])
        all_keys = list(keys) + ([k for k, _ in EXTRA_COLS] if extra else [])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sổ TSCĐ" if loai_so == "TSCĐ" else "Sổ CCDC"
        n = len(heads)
        ws.cell(1, 1, "SỔ TÀI SẢN CỐ ĐỊNH" if loai_so == "TSCĐ" else "SỔ THEO DÕI CÔNG CỤ DỤNG CỤ")
        ws.cell(2, 1, f"Năm {year}" + ("" if case == "toan_bo" else f" – {dict(CASES)[case]}"))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
        ws.cell(1, 1).font = Font(bold=True, size=14)
        ws.cell(1, 1).alignment = ws.cell(2, 1).alignment = Alignment(horizontal="center")
        ws.cell(2, 1).font = Font(italic=True)
        thin = Side(style="thin", color="999999")
        border = Border(top=thin, bottom=thin, left=thin, right=thin)
        grey = PatternFill("solid", fgColor="EDEDED")
        for j, h in enumerate(heads, start=1):
            c = ws.cell(3, j, h)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
            c.fill = grey if j > len(cols) else PatternFill("solid", fgColor="D9E2F3")
        money = {"NguyenGia", "GiaTriTinhKH", "TrongKy", "LuyKe", "ConLai", "MucKy"}
        for i, r in enumerate(rows, start=4):
            for j, k in enumerate(all_keys, start=1):
                v = r.get(k)
                if k in ("NgayGhiTang", "NgayBatDauKH", "NgayGhiGiam") and v:
                    v = datetime.fromisoformat(v[:10])
                c = ws.cell(i, j, v)
                c.border = border
                if isinstance(v, datetime):
                    c.number_format = "dd/mm/yyyy"
                elif k in money:
                    c.number_format = "#,##0"
        last = 4 + len(rows)
        ws.cell(last, 1, f"Số dòng = {len(rows)}").font = Font(bold=True)
        for j, k in enumerate(all_keys, start=1):
            if k in sums:
                c = ws.cell(last, j, f"=SUBTOTAL(9,{get_column_letter(j)}4:{get_column_letter(j)}{last - 1})"
                            if rows else 0)
                c.font = Font(bold=True)
                c.number_format = "#,##0"
            ws.cell(last, j).border = border
        widths = {"TenTS": 46, "MaTS": 22, "LoaiTS": 26, "NhomDM": 34, "HoaDon": 34, "GhiChu": 30}
        for j, k in enumerate(all_keys, start=1):
            ws.column_dimensions[get_column_letter(j)].width = widths.get(k, 15)
        ws.freeze_panes = "C4"
        ws.auto_filter.ref = f"A3:{get_column_letter(n)}{max(last - 1, 3)}"
        wb.save(str(path))
        return len(rows)

    def export_lines(self, path: str | Path, year: int, *, company: str = "", user: str = "") -> int:
        """Bảng kê dòng hóa đơn đầu vào của năm và phân loại, để kế toán soát."""
        from .reports import Report, export_xlsx

        rows = self.lines(f"{year}-01-01", f"{year}-12-31")
        cols = [("NgayNhanHD", "Ngày hóa đơn", "date"), ("KyHieuHD", "Ký hiệu", "text"),
                ("SoHoaDonGoc", "Số HĐ", "text"), ("TenNhaCungCap", "Người bán", "text"),
                ("TenHang", "Tên hàng hóa, dịch vụ", "text"), ("DVT", "ĐVT", "text"), ("SoLuong", "SL", "text"),
                ("DonGiaTinh", "Đơn giá", "money"), ("ThanhTien", "Thành tiền chưa thuế", "money"),
                ("PhanLoai", "Phân loại", "text"), ("NhomDM", "Nhóm danh mục", "text"),
                ("SoThang", "Số tháng KH/PB", "int")]
        export_xlsx(Report(f"Phân loại dòng hóa đơn đầu vào năm {year}", cols, rows, sheet=f"Dòng HĐ đầu vào {year}"),
                    path, company=company, user=user)
        return len(rows)
