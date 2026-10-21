"""Nhập hàng loạt phiếu chi và phiếu thu khác từ file Excel vào module Thu chi.

Chỉ chạy được khi Thu chi đang mở khóa (cần FinanceStore). File Excel có thể đặt mật khẩu mở file:
phần mềm giải mã trong bộ nhớ, không ghi bản giải mã ra đĩa. Khoản thu từ hợp đồng không nhập ở đây
(lấy tự động từ Thanh toán hợp đồng, nhập qua file mẫu hợp đồng).
"""
from __future__ import annotations

import io
import re
from datetime import date
from pathlib import Path

from .excel_import import ImportRow, _norm, _read_sheet
from .finance import DEFAULT_CATEGORIES, FinanceError, FinanceStore

SHEET = "Phiếu thu chi"
LIST_SHEET = "Danh mục"

# (tiêu đề cột, khóa, bắt buộc, kiểu)
VOUCHER_COLS = [
    ("Ngày", "NgayPhatSinh", True, "date"),
    ("Loại (Chi / Thu)", "LoaiThuChi", True, "loai"),
    ("Phân loại chi (Dự án / Văn phòng)", "PhanLoaiChi", False, "phanloai"),
    ("Số hợp đồng dự án", "SoHopDong", False, "text"),
    ("Tên dự án (khi chưa có hợp đồng)", "TenDuAn", False, "text"),
    ("Khoản mục", "KhoanMuc", True, "text"),
    ("Quỹ", "Quy", False, "text"),
    ("Số tiền (đồng)", "SoTien", True, "money"),
    ("Nội dung", "NoiDung", True, "text"),
    ("Người nhận / người nộp", "DoiTuong", False, "text"),
    ("Số chứng từ tham chiếu", "ThamChieu", False, "text"),
    ("Số phiếu", "SoPhieu", False, "text"),
]

EXAMPLES = [
    [date(2026, 3, 31), "Chi", "Văn phòng", "", "", "Nhân sự > Lương", "Tài khoản ngân hàng", 15000000,
     "Lương tháng 3/2026", "Nhân viên văn phòng", "UNC 031", "", ],
    [date(2026, 3, 15), "Chi", "Dự án", "01/2026/HĐTV", "", "Chi phí dự án > Vật tư, thiết bị", "Tiền mặt",
     8000000, "Mua vật tư công trình", "Cửa hàng VLXD Minh Anh", "", ""],
    [date(2026, 3, 20), "Thu", "", "", "", "Thu khác > Lãi tiền gửi", "Tài khoản ngân hàng", 350000,
     "Lãi tiền gửi tháng 3", "Ngân hàng", "", ""],
]

GUIDE = [
    "HƯỚNG DẪN NHẬP PHIẾU THU CHI TỪ EXCEL",
    "",
    "1. Điền từ dòng 2 của sheet 'Phiếu thu chi', giữ nguyên dòng tiêu đề và thứ tự cột. Xóa các dòng ví dụ.",
    "2. Cột có dấu * là bắt buộc. Loại: 'Chi' (phiếu chi) hoặc 'Thu' (thu khác, không phải tiền hợp đồng).",
    "3. Tiền thu từ hợp đồng KHÔNG nhập ở đây: phần mềm tự lấy từ Thanh toán hợp đồng"
    " (nhập bằng file mẫu hợp đồng, sheet Thanh toán).",
    "4. Phiếu chi chọn Phân loại 'Dự án' hoặc 'Văn phòng'. Bỏ trống thì phần mềm tự chọn: có số hợp đồng"
    " hoặc tên dự án là Dự án, còn lại là Văn phòng.",
    "5. Chi dự án: ghi Số hợp đồng đầu ra (phải có trong phần mềm), hoặc ghi Tên dự án nếu chưa có hợp đồng.",
    "6. Khoản mục ghi đầy đủ theo cấp, ngăn cách bằng dấu >, VD 'Nhân sự > Lương'. Ghi tên mục cuối cũng được"
    " nếu tên đó không bị trùng. Khoản mục chưa có sẽ được tự thêm.",
    "7. Quỹ ghi đúng tên quỹ trong phần mềm (xem sheet Danh mục). Bỏ trống thì dùng quỹ ngân hàng mặc định.",
    "8. Số tiền nhập dạng số, đơn vị đồng. Ngày nhập dạng ngày của Excel hoặc dd/mm/yyyy. Số phiếu bỏ trống"
    " thì tự sinh (PC-2026-0001 / PT-2026-0001).",
    "9. Phần mềm kiểm tra trùng: trùng Số phiếu, hoặc trùng cả ngày, loại, số tiền và nội dung với phiếu đã có.",
    "10. Bảo mật: nên đặt mật khẩu mở file (Excel: File > Info > Protect Workbook > Encrypt with Password)."
    " Phần mềm đọc được file có mật khẩu. Sau khi nhập có thể chọn xóa file Excel nguồn.",
    "11. Trong phần mềm: Thu chi (mở khóa) > Sổ thu chi > Nhập từ Excel > chọn file > xem trước > Lưu.",
]

_SEP = re.compile(r"\s*[>›]\s*")


class FileNeedsPassword(FinanceError):
    """File Excel có mật khẩu mở file mà chưa nhập (hoặc nhập sai) mật khẩu."""


def _category_paths(store: FinanceStore | None) -> dict[str, list[str]]:
    if store is None:
        out = {"Chi": [], "Thu": []}
        for (loai, _), groups in DEFAULT_CATEGORIES.items():
            for g, children in groups.items():
                out[loai] += [g] + [f"{g} > {c}" for c in children]
        return out
    out = {}
    for loai in ("Chi", "Thu"):
        out[loai] = [store.category_path(c["Ma"]).replace(" › ", " > ") for c, _ in store.category_tree(loai)]
    return out


def write_template(path: str | Path, store: FinanceStore | None = None) -> Path:
    """Tạo file mẫu. Có store thì sheet Danh mục liệt kê khoản mục, quỹ, dự án hiện có (không có số tiền)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    guide = wb.active
    guide.title = "Hướng dẫn"
    for i, line in enumerate(GUIDE, start=1):
        guide.cell(i, 1, line).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    guide.column_dimensions["A"].width = 125

    paths = _category_paths(store)
    funds = [f["TenQuy"] for f in store.funds()] if store else ["Tiền mặt", "Tài khoản ngân hàng"]
    projects = [(p["SoHopDong"], p["TenHopDong"]) for p in store.projects()] if store else []
    lists = wb.create_sheet(LIST_SHEET)
    heads = ["Khoản mục chi", "Khoản mục thu", "Quỹ", "Số hợp đồng dự án", "Tên dự án"]
    for c, h in enumerate(heads, start=1):
        lists.cell(1, c, h).font = Font(bold=True)
        lists.column_dimensions[lists.cell(1, c).column_letter].width = 40 if c in (1, 2, 5) else 22
    for r, v in enumerate(paths["Chi"], start=2):
        lists.cell(r, 1, v)
    for r, v in enumerate(paths["Thu"], start=2):
        lists.cell(r, 2, v)
    for r, v in enumerate(funds, start=2):
        lists.cell(r, 3, v)
    for r, (so, ten) in enumerate(projects, start=2):
        lists.cell(r, 4, so)
        lists.cell(r, 5, ten)

    ws = wb.create_sheet(SHEET, 1)
    req_fill = PatternFill("solid", fgColor="FCE4D6")
    opt_fill = PatternFill("solid", fgColor="DCE8F5")
    for c, (label, _, required, kind) in enumerate(VOUCHER_COLS, start=1):
        cell = ws.cell(1, c, label + (" *" if required else ""))
        cell.font = Font(bold=True)
        cell.fill = req_fill if required else opt_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = {"date": 13, "money": 16, "loai": 11}.get(kind, 26)
    ws.column_dimensions["F"].width = 36
    ws.column_dimensions["I"].width = 36
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    for r, values in enumerate(EXAMPLES, start=2):
        for c, v in enumerate(values, start=1):
            ws.cell(r, c, v)
    for r in range(2, 2001):
        ws[f"A{r}"].number_format = "DD/MM/YYYY"
        ws[f"H{r}"].number_format = "#,##0"
    validations = [
        ("B", '"Chi,Thu"'),
        ("C", '"Dự án,Văn phòng"'),
        ("G", f"'{LIST_SHEET}'!$C$2:$C${max(len(funds), 1) + 1}"),
    ]
    if projects:
        validations.append(("D", f"'{LIST_SHEET}'!$D$2:$D${len(projects) + 1}"))
    for col, formula in validations:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.showErrorMessage = col in ("B", "C")
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}2000")
    path = Path(path)
    wb.save(path)
    return path


# ---------------------------------------------------------------- đọc & kiểm tra
def is_encrypted(path: str | Path) -> bool:
    import msoffcrypto

    with open(path, "rb") as f:
        try:
            return msoffcrypto.OfficeFile(f).is_encrypted()
        except Exception:  # noqa: BLE001 - không phải file Office được mã hóa
            return False


def _open_workbook(path: str | Path, password: str | None):
    from openpyxl import load_workbook

    if not is_encrypted(path):
        return load_workbook(str(path), data_only=True, read_only=True)
    import msoffcrypto

    if not password:
        raise FileNeedsPassword("File Excel có mật khẩu mở file. Nhập mật khẩu để đọc.")
    buf = io.BytesIO()
    with open(path, "rb") as f:
        office = msoffcrypto.OfficeFile(f)
        try:
            office.load_key(password=password)
            office.decrypt(buf)
        except Exception:  # noqa: BLE001
            raise FileNeedsPassword("Mật khẩu file Excel không đúng.") from None
    buf.seek(0)
    return load_workbook(buf, data_only=True, read_only=True)


def read_workbook(path: str | Path, store: FinanceStore, password: str | None = None) -> list[ImportRow]:
    if store is None:
        raise FinanceError("Mở khóa Thu chi trước khi nhập dữ liệu.")
    wb = _open_workbook(path, password)
    if SHEET not in wb.sheetnames:
        wb.close()
        raise FinanceError(f"File không có sheet '{SHEET}'. Dùng file mẫu nhập thu chi.")
    rows = _read_sheet(wb[SHEET], VOUCHER_COLS, SHEET)
    wb.close()
    if not rows:
        raise FinanceError(f"Sheet '{SHEET}' chưa có dữ liệu.")
    check(rows, store)
    return rows


def _key(s: str) -> str:
    return _norm(s).lower()


class _Lookup:
    def __init__(self, store: FinanceStore):
        self.store = store
        self.paths: dict[tuple[str, str], str] = {}
        self.leaves: dict[tuple[str, str], list[str]] = {}
        for loai in ("Chi", "Thu"):
            for c, _ in store.category_tree(loai, include_hidden=True):
                parts = [_key(p) for p in store.category_path(c["Ma"]).split(" › ")]
                self.paths[(loai, " > ".join(parts))] = c["Ma"]
                self.leaves.setdefault((loai, parts[-1]), []).append(c["Ma"])
        self.funds = {_key(f["TenQuy"]): f["Ma"] for f in store.funds(include_hidden=True)}
        self.projects: dict[str, list] = {}
        for p in store.projects():
            self.projects.setdefault(_key(p["SoHopDong"]), []).append(p)

    def category(self, loai: str, text: str) -> tuple[str | None, list[str]]:
        """Trả về (Mã khoản mục hoặc None nếu sẽ thêm mới, các phần của đường dẫn)."""
        parts = [_norm(p) for p in _SEP.split(text) if _norm(p)]
        if not parts:
            raise FinanceError("Thiếu Khoản mục")
        full = " > ".join(_key(p) for p in parts)
        if (loai, full) in self.paths:
            return self.paths[(loai, full)], parts
        if len(parts) == 1:
            hits = self.leaves.get((loai, _key(parts[0])), [])
            if len(hits) == 1:
                return hits[0], parts
            if len(hits) > 1:
                raise FinanceError(f"Có nhiều khoản mục tên '{parts[0]}', ghi đầy đủ dạng 'Mục cha > {parts[0]}'")
        return None, parts


def _norm_loai(v: str) -> str | None:
    k = _key(v)
    if k in ("chi", "phiếu chi"):
        return "Chi"
    if k in ("thu", "thu khác", "phiếu thu"):
        return "Thu"
    return None


def _norm_phan_loai(v: str) -> str | None:
    k = _key(v).replace("chi ", "", 1) if _key(v).startswith("chi ") else _key(v)
    if k in ("dự án", "du an"):
        return "Dự án"
    if k in ("văn phòng", "van phong", "hoạt động văn phòng"):
        return "Văn phòng"
    return None


def check(rows: list[ImportRow], store: FinanceStore) -> None:
    """Chuẩn hóa từng dòng, gắn mã khoản mục / quỹ / dự án, đánh dấu Lỗi / Trùng."""
    lk = _Lookup(store)
    default_fund = store.default_fund("Chuyển khoản")
    existing_no = {_key(r["SoPhieu"]) for r in store.vouchers()}
    existing_sig = {(r["NgayPhatSinh"], r["LoaiThuChi"], r["SoTien"], _key(r["NoiDung"])) for r in store.vouchers()}
    seen_no: set[str] = set()
    seen_sig: set[tuple] = set()
    for r in rows:
        if not r.ok:
            continue
        v = r.values
        errors, notes = [], []
        loai = _norm_loai(v["LoaiThuChi"])
        if not loai:
            errors.append("Loại phải là Chi hoặc Thu")
        v["LoaiThuChi"] = loai
        if loai == "Chi":
            pl = _norm_phan_loai(v["PhanLoaiChi"]) if v["PhanLoaiChi"] else (
                "Dự án" if v["SoHopDong"] or v["TenDuAn"] else "Văn phòng")
            if not pl:
                errors.append("Phân loại chi phải là Dự án hoặc Văn phòng")
            v["PhanLoaiChi"] = pl
        else:
            v["PhanLoaiChi"] = None
        v["MaDuAn"] = None
        if v["PhanLoaiChi"] == "Dự án":
            if v["SoHopDong"]:
                hits = lk.projects.get(_key(v["SoHopDong"]), [])
                if len(hits) == 1:
                    v["MaDuAn"], v["TenDuAn"] = hits[0]["Ma"], hits[0]["TenHopDong"]
                elif hits:
                    errors.append(f"Có nhiều hợp đồng đầu ra số {v['SoHopDong']}")
                else:
                    errors.append(f"Không tìm thấy hợp đồng đầu ra số {v['SoHopDong']}")
            elif not v["TenDuAn"]:
                errors.append("Chi dự án phải ghi Số hợp đồng hoặc Tên dự án")
        else:
            v["TenDuAn"] = None
        if loai:
            try:
                v["MaKhoanMuc"], v["_parts"] = lk.category(loai, v["KhoanMuc"])
                if v["MaKhoanMuc"] is None:
                    notes.append(f"Sẽ thêm khoản mục mới: {' > '.join(v['_parts'])}")
            except FinanceError as e:
                errors.append(str(e))
        if v["Quy"]:
            v["MaQuy"] = lk.funds.get(_key(v["Quy"]))
            if not v["MaQuy"]:
                errors.append(f"Không có quỹ '{v['Quy']}'")
        else:
            v["MaQuy"] = default_fund
        if v["SoTien"] is not None and v["SoTien"] <= 0:
            errors.append("Số tiền phải lớn hơn 0")
        if errors:
            r.status, r.errors = "Lỗi", errors
            continue
        no = _key(v["SoPhieu"])
        sig = (v["NgayPhatSinh"], loai, v["SoTien"], _key(v["NoiDung"]))
        if no and no in existing_no:
            r.status, r.errors = "Trùng", [f"Số phiếu {v['SoPhieu']} đã có"]
        elif no and no in seen_no:
            r.status, r.errors = "Trùng", ["Số phiếu trùng với dòng khác trong file"]
        elif sig in existing_sig:
            r.status, r.errors = "Trùng", ["Đã có phiếu cùng ngày, loại, số tiền và nội dung"]
        elif sig in seen_sig:
            r.status, r.errors = "Trùng", ["Trùng ngày, loại, số tiền, nội dung với dòng khác trong file"]
        else:
            r.errors = notes
        if no:
            seen_no.add(no)
        seen_sig.add(sig)


def _ensure_category(store: FinanceStore, lk: _Lookup, loai: str, parts: list[str], ap_dung: str) -> str:
    parent = None
    for i in range(len(parts)):
        key = (loai, " > ".join(_key(p) for p in parts[:i + 1]))
        if key not in lk.paths:
            lk.paths[key] = store.add_category(parts[i], loai, parent, ap_dung)
        parent = lk.paths[key]
    return parent


def import_rows(store: FinanceStore, rows: list[ImportRow]) -> dict:
    """Lưu các dòng hợp lệ trong một lần ghi tệp mã hóa. Trả về số phiếu đã lưu và lỗi (nếu có)."""
    lk = _Lookup(store)
    saved = {"Chi": 0, "Thu": 0, "tong_chi": 0, "tong_thu": 0, "khoan_muc_moi": 0, "loi": []}
    before = len(store.categories(include_hidden=True))
    for r in rows:
        if not r.ok:
            continue
        v = r.values
        try:
            ma_km = v["MaKhoanMuc"] or _ensure_category(
                store, lk, v["LoaiThuChi"], v["_parts"],
                v["PhanLoaiChi"] if v["LoaiThuChi"] == "Chi" else "Cả hai")
            values = store._validate({
                "NgayPhatSinh": v["NgayPhatSinh"], "LoaiThuChi": v["LoaiThuChi"], "PhanLoaiChi": v["PhanLoaiChi"],
                "MaDuAn": v["MaDuAn"], "TenDuAn": v["TenDuAn"], "MaKhoanMuc": ma_km, "MaQuy": v["MaQuy"],
                "SoTien": v["SoTien"], "NoiDung": v["NoiDung"], "DoiTuong": v["DoiTuong"] or None,
                "ThamChieu": v["ThamChieu"] or None})
            values["SoPhieu"] = v["SoPhieu"] or store.next_voucher_no(values["LoaiThuChi"], values["NgayPhatSinh"])
            values["NguoiTao"] = values["NguoiSua"] = store.username
            store.fdb.insert("PhieuThuChi", values, commit=False)
        except FinanceError as e:
            r.status, r.errors = "Lỗi", [str(e)]
            saved["loi"].append(f"Dòng {r.line}: {e}")
            continue
        saved[values["LoaiThuChi"]] += 1
        saved["tong_chi" if values["LoaiThuChi"] == "Chi" else "tong_thu"] += values["SoTien"]
    saved["khoan_muc_moi"] = len(store.categories(include_hidden=True)) - before
    store.save()
    store.db.log(store.username, f"Nhập {saved['Chi'] + saved['Thu']} phiếu thu chi từ Excel")
    return saved
