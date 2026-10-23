"""Nhập hàng loạt hợp đồng, hóa đơn đầu ra và thanh toán từ file Excel mẫu.

Quy trình: tải file mẫu → điền → chọn file → xem trước (đánh dấu dòng lỗi / trùng) → xác nhận lưu.
Loại hợp đồng và đối tác chưa có trong danh mục được tự thêm khi lưu.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from . import vat
from .contracts import CONTRACT_STATUSES, DURATION_UNITS, PAY_METHODS, ContractService, ValidationError
from .db import Database

RATE_OPTIONS = ("10%", "8%", "5%", "0%", "KCT", "KKKNT")
VAT_AFTER = {"Hợp đồng": "GiaTriTongSauThue", "Hóa đơn đầu ra": "GiaTriSauThue"}

# (tiêu đề cột, khóa, bắt buộc, kiểu)
CONTRACT_COLS = [
    ("Số hợp đồng", "SoHopDong", True, "text"),
    ("Tên hợp đồng", "TenHopDong", True, "text"),
    ("Chiều hợp đồng", "ChieuHopDong", True, "chieu"),
    ("Loại hợp đồng", "LoaiHD", False, "text"),
    ("Đối tác (Chủ đầu tư / Nhà cung cấp)", "TenDoiTac", True, "text"),
    ("MST đối tác", "MaThueDoiTac", False, "text"),
    ("Giá trị trước thuế (đồng)", "GiaTriTruocThue", False, "money"),
    ("Thuế suất VAT (%)", "ThueSuat", False, "rate"),
    ("Tiền thuế VAT (đồng)", "VAT", False, "money"),
    ("Giá trị sau thuế (đồng)", "GiaTriTongSauThue", False, "money"),
    ("Ngày ký", "NgayKy", True, "date"),
    ("Ngày bắt đầu dự kiến", "NgayBatDauDuKien", False, "date"),
    ("Thời gian thực hiện", "ThoiGianThucHien", True, "int"),
    ("Đơn vị thời gian", "DonViThoiGian", False, "unit"),
    ("Ngày kết thúc dự kiến", "NgayKetThucDuKien", False, "date"),
    ("Tình trạng", "TinhTrangHopDong", False, "status"),
    ("Ghi chú", "GhiChu", False, "text"),
]
INVOICE_COLS = [
    ("Số hợp đồng", "SoHopDong", True, "text"),
    ("Ký hiệu mẫu số", "KyHieuMauSo", False, "text"),
    ("Ký hiệu hóa đơn", "KyHieuHD", False, "text"),
    ("Số hóa đơn", "SoHoaDon", True, "text"),
    ("Ngày xuất hóa đơn", "NgayXuatHD", True, "date"),
    ("Giá trị trước thuế (đồng)", "GiaTriTruocThue", False, "money"),
    ("Thuế suất VAT (%)", "ThueSuat", False, "rate"),
    ("Tiền thuế VAT (đồng)", "VAT", False, "money"),
    ("Giá trị sau thuế (đồng)", "GiaTriSauThue", False, "money"),
    ("Ghi chú", "GhiChuHDDR", False, "text"),
]
PAYMENT_COLS = [
    ("Số hợp đồng", "SoHopDong", True, "text"),
    ("Ngày tiền về", "NgayThanhToan", True, "date"),
    ("Số tiền (đồng)", "SoTien", True, "money"),
    ("Hình thức", "HinhThuc", False, "method"),
    ("Số tham chiếu sao kê", "SoThamChieuSK", False, "text"),
    ("Số hóa đơn liên quan", "SoHoaDon", False, "text"),
    ("Ghi chú", "GhiChu", False, "text"),
]
SHEETS = [("Hợp đồng", CONTRACT_COLS), ("Hóa đơn đầu ra", INVOICE_COLS), ("Thanh toán", PAYMENT_COLS)]

EXAMPLES = {
    "Hợp đồng": [
        ["01/2026/HĐTV", "Tư vấn thiết kế trường THCS A", "Đầu ra", "Tư vấn thiết kế", "Ban QLDA quận Ba Đình",
         "0101234567", 100000000, "8%", None, None, date(2026, 1, 15), date(2026, 2, 1), 11, "tháng", None,
         "Đang thực hiện", "Ví dụ, xóa dòng này trước khi nhập"],
        ["05/2026/MB", "Mua vật tư văn phòng", "Đầu vào", "Mua sắm", "Công ty TNHH Vật tư B", "0309876543",
         None, "10%", None, 22000000, date(2026, 2, 3), None, 30, "ngày", None, "Hoàn thành", "Ví dụ, xóa dòng này trước khi nhập"],
    ],
    "Hóa đơn đầu ra": [
        ["01/2026/HĐTV", "1", "C26TAA", "0000012", date(2026, 3, 1), 50000000, None, 4000000, None, "Ví dụ"],
    ],
    "Thanh toán": [
        ["01/2026/HĐTV", date(2026, 3, 10), 30000000, "Chuyển khoản", "FT26069123456", "0000012", "Ví dụ"],
    ],
}

GUIDE = [
    "HƯỚNG DẪN NHẬP DỮ LIỆU TỪ EXCEL",
    "",
    "1. Điền dữ liệu từ dòng 2 của từng sheet, giữ nguyên dòng tiêu đề và thứ tự cột.",
    "2. Cột có dấu * là bắt buộc. Xóa các dòng ví dụ trước khi nhập.",
    "3. Sheet 'Hợp đồng': mỗi dòng một hợp đồng. Chiều hợp đồng chọn 'Đầu ra' (ký với chủ đầu tư) hoặc"
    " 'Đầu vào' (ký với nhà cung cấp).",
    "4. Sheet 'Hóa đơn đầu ra' và 'Thanh toán' không bắt buộc; nối với hợp đồng qua cột Số hợp đồng"
    " (hợp đồng phải có trong file hoặc đã có trong phần mềm).",
    "5. Số tiền nhập dạng số, đơn vị đồng (VD 100000000). Ngày nhập dạng ngày của Excel hoặc dd/mm/yyyy.",
    "   Giá trị: chỉ cần 2 trong 4 cột Trước thuế / Thuế suất VAT / Tiền thuế VAT / Sau thuế, phần mềm tự tính"
    " phần còn lại: trước thuế + thuế suất; sau thuế + thuế suất; trước thuế + tiền VAT; sau thuế + tiền VAT."
    " Thuế suất ghi 10%, 8%, 5%, 0%, KCT… Điền đủ cả ba số tiền thì phần mềm kiểm tra trước thuế + VAT = sau thuế.",
    "   Thời gian thực hiện (bắt buộc): số ngày hoặc số tháng; Đơn vị thời gian ghi 'ngày' hoặc 'tháng' (để trống là"
    " ngày). Phần mềm tự tính Ngày kết thúc dự kiến từ ngày bắt đầu (trống thì từ ngày ký), tính cả ngày bắt đầu.",
    "6. Loại hợp đồng, đối tác chưa có trong phần mềm sẽ được tự thêm vào danh mục.",
    "7. Phần mềm kiểm tra trùng: hợp đồng trùng Số hợp đồng + đối tác, hóa đơn trùng Ký hiệu + Số hóa đơn"
    " (cả với hóa đơn đã có lẫn các dòng khác trong file) và báo ngay khi chọn file. Dòng trùng hoặc lỗi"
    " không được lưu. Hóa đơn trùng số nhưng thiếu ký hiệu được đánh dấu Nghi trùng để bạn tự quyết định.",
    "8. Trong phần mềm: Hợp đồng → Nhập từ Excel → chọn file → kiểm tra bảng xem trước → bấm Lưu các dòng hợp lệ.",
]


def write_template(path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    guide = wb.active
    guide.title = "Hướng dẫn"
    for i, line in enumerate(GUIDE, start=1):
        guide.cell(i, 1, line).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    guide.column_dimensions["A"].width = 120
    req_fill = PatternFill("solid", fgColor="FCE4D6")
    opt_fill = PatternFill("solid", fgColor="DCE8F5")
    for name, cols in SHEETS:
        ws = wb.create_sheet(name)
        for c, (label, _, required, kind) in enumerate(cols, start=1):
            cell = ws.cell(1, c, label + (" *" if required else ""))
            cell.font = Font(bold=True)
            cell.fill = req_fill if required else opt_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[cell.column_letter].width = 18 if kind in ("money", "date", "rate") else 26
            if kind == "chieu":
                cell.comment = Comment("Đầu ra hoặc Đầu vào", "QLHD")
        ws.row_dimensions[1].height = 34
        ws.freeze_panes = "A2"
        for r, values in enumerate(EXAMPLES[name], start=2):
            for c, v in enumerate(values, start=1):
                cell = ws.cell(r, c, v)
                kind = cols[c - 1][3]
                if kind == "money":
                    cell.number_format = "#,##0"
                elif kind == "date":
                    cell.number_format = "DD/MM/YYYY"
        for c, (_, _, _, kind) in enumerate(cols, start=1):
            letter = ws.cell(1, c).column_letter
            options = {"chieu": ["Đầu ra", "Đầu vào"], "status": list(CONTRACT_STATUSES),
                       "unit": list(DURATION_UNITS),
                       "method": list(PAY_METHODS), "rate": list(RATE_OPTIONS)}.get(kind)
            if options:
                dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                                    showErrorMessage=kind != "rate")
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}2000")
            elif kind == "money":
                for r in range(2, 2001):
                    ws[f"{letter}{r}"].number_format = "#,##0"
            elif kind == "date":
                for r in range(2, 2001):
                    ws[f"{letter}{r}"].number_format = "DD/MM/YYYY"
    path = Path(path)
    wb.save(path)
    return path


# ---------------------------------------------------------------- đọc & kiểm tra
@dataclass
class ImportRow:
    sheet: str
    line: int
    values: dict
    errors: list[str] = field(default_factory=list)
    status: str = "Hợp lệ"          # Hợp lệ / Trùng / Nghi trùng / Lỗi
    force: bool = False             # người dùng chọn vẫn lưu dòng "Nghi trùng"

    @property
    def ok(self) -> bool:
        return self.status == "Hợp lệ" or (self.status == "Nghi trùng" and self.force)


def invoice_no(so) -> str:
    """Số hóa đơn để so trùng: bỏ khoảng trắng và số 0 đứng đầu (00000123 và 123 là một)."""
    s = "".join(str(so or "").split()).lower()
    return s.lstrip("0") or s


def _to_date(v) -> str | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"ngày không hợp lệ: {s}")


def _to_money(v) -> int:
    if v in (None, ""):
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = str(v).strip().replace(" ", "").replace(".", "").replace(",", "")
    if not s.lstrip("-").isdigit():
        raise ValueError(f"số tiền không hợp lệ: {v}")
    return int(s)


def _to_int(v) -> int:
    if v in (None, ""):
        return 0
    if isinstance(v, float) and v.is_integer():
        return int(v)
    s = str(v).strip().split()[0] if str(v).strip() else ""
    if not s.isdigit() or int(s) <= 0:
        raise ValueError(f"cần số nguyên dương: {v}")
    return int(s)


def _norm(s) -> str:
    return " ".join(str(s or "").split()).strip()


def _header_index(header: tuple, cols) -> list[int | None]:
    """Vị trí từng cột theo tên tiêu đề (file mẫu bản cũ ít cột hơn vẫn đọc đúng); không khớp thì theo thứ tự."""
    names = [" ".join(str(h or "").replace("*", "").split()).lower() for h in header]
    idx = [names.index(label.lower()) if label.lower() in names else None for label, *_ in cols]
    if sum(i is not None for i in idx) * 2 < len(cols):
        return list(range(len(cols)))
    return idx


def _read_sheet(ws, cols, sheet_name) -> list[ImportRow]:
    out = []
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    idx = _header_index(header, cols)
    for r_idx, row in enumerate(rows, start=2):
        if row is None or all(v in (None, "") for v in row):
            continue
        row = list(row)
        cells = [row[i] if i is not None and i < len(row) else None for i in idx]
        values, errors, blank = {}, [], set()
        for (label, key, required, kind), v in zip(cols, cells):
            if v in (None, ""):
                blank.add(key)
            try:
                if kind == "money":
                    values[key] = _to_money(v)
                elif kind == "date":
                    values[key] = _to_date(v)
                elif kind == "rate":
                    values[key] = vat.parse_rate(v)
                elif kind == "int":
                    values[key] = _to_int(v)
                elif kind == "unit":
                    values[key] = _norm(v).lower() or "ngày"
                    if values[key] not in DURATION_UNITS:
                        raise ValueError("ghi 'ngày' hoặc 'tháng'")
                else:
                    values[key] = _norm(v) if v not in (None, "") else ""
            except ValueError as e:
                errors.append(f"{label}: {e}")
                values[key] = None
            if required and values.get(key) in (None, "", 0) and not (kind == "money" and v == 0):
                errors.append(f"Thiếu {label}")
            if kind == "chieu" and values.get(key) and values[key] not in ("Đầu ra", "Đầu vào"):
                errors.append("Chiều hợp đồng phải là Đầu ra hoặc Đầu vào")
            if kind == "status" and values.get(key) and values[key] not in CONTRACT_STATUSES:
                errors.append(f"Tình trạng không hợp lệ: {values[key]}")
            if kind == "method" and values.get(key) and values[key] not in PAY_METHODS:
                errors.append(f"Hình thức không hợp lệ: {values[key]}")
        notes = []
        if sheet_name in VAT_AFTER and not errors:
            notes = _complete_vat(values, blank, VAT_AFTER[sheet_name], errors)
        item = ImportRow(sheet_name, r_idx, values, errors + notes)
        if errors:
            item.status = "Lỗi"
        out.append(item)
    return out


def _complete_vat(v: dict, blank: set, after_key: str, errors: list[str]) -> list[str]:
    """Tự tính trước thuế / VAT / sau thuế / thuế suất từ hai dữ kiện; trả về các lưu ý cho màn xem trước."""
    def given(k):
        return None if k in blank else v[k]

    before, after = given("GiaTriTruocThue"), given(after_key)
    if not before and not after:
        errors.append("Thiếu giá trị: cần Giá trị trước thuế hoặc Giá trị sau thuế")
        return []
    res = vat.complete(before, given("VAT"), after, v.get("ThueSuat") or "")
    if res.error:
        errors.append(res.error)
        return []
    v["GiaTriTruocThue"], v["VAT"], v[after_key], v["ThueSuat"] = res.before, res.vat, res.after, res.rate
    notes = [f"Lưu ý: {n}" for n in res.notes]
    auto = {"Trước thuế": res.before, "VAT": res.vat, "Sau thuế": res.after}
    shown = [f"{k} {fmt_vnd(auto[k])}" if k in auto else f"{k} {res.rate}" for k in res.computed]
    if shown:
        notes.insert(0, "Tự tính: " + ", ".join(shown))
    return notes


def fmt_vnd(n: int) -> str:
    return f"{n:,}".replace(",", ".")


def read_workbook(path: str | Path, db: Database) -> list[ImportRow]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    rows: list[ImportRow] = []
    for name, cols in SHEETS:
        if name in wb.sheetnames:
            rows += _read_sheet(wb[name], cols, name)
    wb.close()
    if not rows:
        raise ValidationError("File không có dữ liệu ở các sheet Hợp đồng / Hóa đơn đầu ra / Thanh toán.")
    _check(rows, db)
    return rows


def _contract_key(so, doi_tac_mst, doi_tac_ten) -> tuple:
    return (so.lower(), (doi_tac_mst or doi_tac_ten or "").lower())


def _check(rows: list[ImportRow], db: Database) -> None:
    existing = {_contract_key(r["SoHopDong"], r["MaThueDoiTac"], r["TenDoiTac"])
                for r in db.active("HopDong")}
    existing_no = {r["SoHopDong"].lower() for r in db.active("HopDong")}
    file_no: dict[str, int] = {}
    seen = set()
    for r in rows:
        if r.sheet != "Hợp đồng" or not r.ok:
            continue
        v = r.values
        k = _contract_key(v["SoHopDong"], v["MaThueDoiTac"], v["TenDoiTac"])
        if k in existing:
            r.status, r.errors = "Trùng", ["Hợp đồng đã có trong phần mềm"]
        elif k in seen:
            r.status, r.errors = "Trùng", ["Trùng với dòng khác trong file"]
        else:
            seen.add(k)
            file_no[v["SoHopDong"].lower()] = file_no.get(v["SoHopDong"].lower(), 0) + 1
    known_no = existing_no | set(file_no)
    contract_no = {r["Ma"]: r["SoHopDong"] for r in db.query("SELECT Ma, SoHopDong FROM HopDong")}
    by_no: dict[str, list[dict]] = {}       # số hóa đơn → các hóa đơn đã có / dòng trước trong file
    for inv in db.active("HoaDonDauRa"):
        by_no.setdefault(invoice_no(inv["SoHoaDon"]), []).append(
            {"kh": (inv["KyHieuHD"] or "").upper(), "so": inv["SoHoaDon"],
             "noi": f"đã có trong phần mềm (ngày {_vn(inv['NgayXuatHD'])}, HĐ {contract_no.get(inv['MaHopDong'], '')})"})
    for r in rows:
        if r.sheet == "Hợp đồng" or not r.ok:
            continue
        so = r.values["SoHopDong"].lower()
        if so not in known_no:
            r.status, r.errors = "Lỗi", ["Không tìm thấy hợp đồng có số này"]
            continue
        if file_no.get(so, 0) > 1:
            r.status, r.errors = "Lỗi", ["Số hợp đồng này có nhiều dòng trong file, không xác định được hợp đồng"]
            continue
        if r.sheet == "Hóa đơn đầu ra":
            _check_invoice_no(r, by_no)
        if r.sheet == "Thanh toán" and r.values["SoTien"] <= 0:
            r.status, r.errors = "Lỗi", ["Số tiền phải lớn hơn 0"]


def _vn(d: str | None) -> str:
    return datetime.strptime(d[:10], "%Y-%m-%d").strftime("%d/%m/%Y") if d else ""


def _check_invoice_no(r: ImportRow, by_no: dict[str, list[dict]]) -> None:
    """Trùng ký hiệu + số: bỏ qua. Trùng số mà một bên thiếu ký hiệu: Nghi trùng (người dùng quyết định).
    Trùng số nhưng khác ký hiệu: vẫn hợp lệ, kèm lời nhắc."""
    kh, so = (r.values["KyHieuHD"] or "").upper(), r.values["SoHoaDon"]
    same = by_no.setdefault(invoice_no(so), [])
    exact = [x for x in same if x["kh"] == kh]
    unsure = [x for x in same if x["kh"] != kh and not (x["kh"] and kh)]
    other = [x for x in same if x["kh"] != kh and x["kh"] and kh]
    label = lambda x: f"hóa đơn {(x['kh'] + ' ') if x['kh'] else ''}số {x['so']} {x['noi']}"  # noqa: E731
    if exact:
        r.status, r.errors = "Trùng", ["Trùng số hóa đơn: " + label(exact[0])]
    elif unsure:
        r.status, r.errors = "Nghi trùng", ["Trùng số hóa đơn, thiếu ký hiệu để phân biệt: " + label(unsure[0])]
    elif other:
        r.errors.append("Lưu ý: cùng số với " + label(other[0]) + ", khác ký hiệu")
    same.append({"kh": kh, "so": so, "noi": f"ở dòng {r.line} trong file"})


def duplicate_invoice_rows(rows: list[ImportRow]) -> list[ImportRow]:
    return [r for r in rows if r.sheet == "Hóa đơn đầu ra" and r.status in ("Trùng", "Nghi trùng")]


def import_rows(db: Database, rows: list[ImportRow]) -> dict:
    """Lưu các dòng hợp lệ, trả về số lượng đã lưu theo từng sheet."""
    cs = ContractService(db)
    types = {t["TenLoaiHD"].lower(): t["Ma"] for t in cs.contract_types()}
    saved = {"Hợp đồng": 0, "Hóa đơn đầu ra": 0, "Thanh toán": 0, "loi": []}
    by_no: dict[str, str] = {}
    for r in rows:
        if r.sheet != "Hợp đồng" or not r.ok:
            continue
        v = dict(r.values)
        loai = v.pop("LoaiHD")
        if loai:
            if loai.lower() not in types:
                types[loai.lower()] = cs.save_contract_type(loai)
            v["MaLoaiHD"] = types[loai.lower()]
        v["MaDoiTac"] = cs.find_or_create_partner(v["TenDoiTac"], v["MaThueDoiTac"] or None,
                                                  "Chủ đầu tư" if v["ChieuHopDong"] == "Đầu ra" else "Nhà cung cấp")
        v["TinhTrangHopDong"] = v["TinhTrangHopDong"] or "Chưa bắt đầu"
        try:
            by_no[v["SoHopDong"].lower()] = cs.save_contract(v)
            saved["Hợp đồng"] += 1
        except ValidationError as e:
            r.status, r.errors = "Lỗi", [str(e)]
            saved["loi"].append(f"Hợp đồng dòng {r.line}: {e}")

    def contract_for(so: str) -> str | None:
        if so.lower() in by_no:
            return by_no[so.lower()]
        # so sánh bằng Python: lower() của SQLite không đổi chữ có dấu (Đ, Ư…)
        hit = [h for h in db.active("HopDong") if h["SoHopDong"].lower() == so.lower()]
        return hit[0]["Ma"] if len(hit) == 1 else None

    inv_by_no: dict[str, str] = {}
    for r in rows:
        if r.sheet != "Hóa đơn đầu ra" or not r.ok:
            continue
        v = dict(r.values)
        ma_hd = contract_for(v.pop("SoHopDong"))
        try:
            if not ma_hd:
                raise ValidationError("Không tìm thấy hợp đồng")
            v["MaHopDong"] = ma_hd
            inv_by_no[invoice_no(v["SoHoaDon"])] = cs.save_out_invoice(v)
            saved["Hóa đơn đầu ra"] += 1
        except ValidationError as e:
            r.status, r.errors = "Lỗi", [str(e)]
            saved["loi"].append(f"Hóa đơn dòng {r.line}: {e}")
    for r in rows:
        if r.sheet != "Thanh toán" or not r.ok:
            continue
        v = dict(r.values)
        ma_hd = contract_for(v.pop("SoHopDong"))
        so_hd = v.pop("SoHoaDon")
        try:
            if not ma_hd:
                raise ValidationError("Không tìm thấy hợp đồng")
            v["MaHopDong"] = ma_hd
            if so_hd:
                ma_inv = inv_by_no.get(invoice_no(so_hd))
                if not ma_inv:
                    hit = [i for i in db.active("HoaDonDauRa", "MaHopDong = ?", (ma_hd,))
                           if invoice_no(i["SoHoaDon"]) == invoice_no(so_hd)]
                    ma_inv = hit[0]["Ma"] if hit else None
                v["MaHDDauRa"] = ma_inv
            v["HinhThuc"] = v["HinhThuc"] or "Chuyển khoản"
            cs.save_payment(v)
            saved["Thanh toán"] += 1
        except ValidationError as e:
            r.status, r.errors = "Lỗi", [str(e)]
            saved["loi"].append(f"Thanh toán dòng {r.line}: {e}")
    return saved
