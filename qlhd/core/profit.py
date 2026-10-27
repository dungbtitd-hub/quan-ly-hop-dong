"""Bảng tổng hợp lãi lỗ theo năm, báo cáo doanh thu và chi phí năm.

Số liệu bảng lãi lỗ (dòng nhập tay, tỷ lệ thuế, số chốt) nằm trong tệp thu chi mã hóa: chỉ Giám đốc và Người quản
lý mở được, chỉ đi theo Gói Giám đốc, xóa vĩnh viễn được bất cứ lúc nào (từng dòng, từng năm, toàn bộ).
Phần tự tính lấy từ dữ liệu chung: hóa đơn đầu ra (doanh thu), sổ TSCĐ / CCDC (khấu hao, phân bổ), dòng hóa đơn
đầu vào là chi phí trong kỳ (theo dòng của danh mục), hợp đồng đầu ra chưa xuất hóa đơn (gợi ý dở dang).
"""
from __future__ import annotations

import json
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from . import asset_catalog as dm
from .assets import AssetService
from .db import now_iso

GROUPS = ("Doanh thu", "Điều chỉnh", "Chi phí", "Thuế")
REVENUE_MODES = [("xuat_hd", "Hóa đơn đầu ra đã xuất trong năm"),
                 ("da_tt", "Chỉ phần hóa đơn xuất trong năm đã thanh toán đến 31/12")]
SOURCES = {
    "nhap": "Nhập tay",
    "doanh_thu": "Tự động: hóa đơn đầu ra (giá trị chưa thuế)",
    "tscd": "Tự động: khấu hao trong năm từ sổ TSCĐ",
    "ccdc": "Tự động: phân bổ trong năm từ sổ CCDC (gồm chi phí chờ phân bổ)",
    "dong:6": "Tự động: hóa đơn đầu vào chi phí trong kỳ (xăng xe, VPP, cước, điện thoại…)",
    "dong:7": "Tự động: hóa đơn đầu vào chi phí trong kỳ (sạc xe, tiếp khách, công tác, CPN…)",
    "dong:6a": "Tự động: vật tư, thầu phụ dự án (hóa đơn gắn hợp đồng đầu vào)",
    "do_dang_dau": "Gợi ý: dở dang chuyển sang của năm trước; sửa được",
    "do_dang_cuoi": "Gợi ý: hợp đồng đầu ra ký trong năm chưa xuất hóa đơn đến 31/12; sửa được",
}
# (nhóm, tên dòng, dấu, nguồn, dòng hệ thống)
DEFAULT_ROWS = [
    ("Doanh thu", "Doanh thu", 1, "doanh_thu", 1),
    ("Điều chỉnh", "Lãi TG", -1, "nhap", 0),
    ("Điều chỉnh", "Tiền cty đấu thầu chuyển lại", -1, "nhap", 0),
    ("Chi phí", "Chi phí thuế môn bài", 1, "nhap", 1),
    ("Chi phí", "Tiền lương người lao động", 1, "nhap", 1),
    ("Chi phí", "Tiền bảo hiểm công ty đóng (21.5%)", 1, "nhap", 1),
    ("Chi phí", "Phân bổ TSCĐ", 1, "tscd", 1),
    ("Chi phí", "Phân bổ công cụ dụng cụ", 1, "ccdc", 1),
    ("Chi phí", "Chi phí xăng xe, VPP, cước đường bộ", 1, "dong:6", 1),
    ("Chi phí", "Chi phí sạc xe, tiếp khách, CPN", 1, "dong:7", 1),
    ("Chi phí", "Chi phí vật tư, thầu phụ dự án", 1, "dong:6a", 1),
    ("Chi phí", "Chi phí dở dang đầu kỳ", 1, "do_dang_dau", 1),
    ("Chi phí", "Chi phí dở dang để sang năm sau", -1, "do_dang_cuoi", 1),
    ("Thuế", "Thuế TNCN phải nộp sau quyết toán", 1, "nhap", 1),
]
HIDE_WHEN_ZERO = {"dong:6a"}       # dòng hiếm dùng: tự ẩn khi bằng 0


class ProfitError(Exception):
    pass


def _round(x: Decimal) -> int:
    return int(x.quantize(Decimal(1), rounding=ROUND_HALF_UP))


def _quarter(day: str) -> int:
    return (int(day[5:7]) - 1) // 3


class ProfitService:
    """store: FinanceStore đang mở (nơi lưu số liệu bảng lãi lỗ); db: CSDL chính."""

    def __init__(self, store, db=None):
        self.store = store
        self.fdb = store.fdb
        self.db = db or store.db
        self.assets = AssetService(self.db)

    # ------------------------------------------------------------ cấu trúc năm
    def years(self) -> list[int]:
        return [r["Nam"] for r in self.fdb.query("SELECT DISTINCT Nam FROM LaiLoNam WHERE DaXoa = 0 ORDER BY Nam DESC")]

    def year_info(self, year: int, *, create: bool = True):
        self._dedupe(year)
        r = self.fdb.one("SELECT * FROM LaiLoNam WHERE DaXoa = 0 AND Nam = ? ORDER BY NgaySua DESC", (year,))
        if r is None and create:
            self._seed(year)
            r = self.fdb.one("SELECT * FROM LaiLoNam WHERE DaXoa = 0 AND Nam = ?", (year,))
        return r

    def _seed(self, year: int) -> None:
        """Năm mới: chép cấu trúc dòng của năm gần nhất trước đó (không chép số), chưa có thì dùng mẫu."""
        prev = self.fdb.one("SELECT * FROM LaiLoNam WHERE DaXoa = 0 AND Nam < ? ORDER BY Nam DESC", (year,))
        self.fdb.insert("LaiLoNam", {"Nam": year, "TyLeThueTNDN": prev["TyLeThueTNDN"] if prev else 20,
                                     "CheDoDoanhThu": prev["CheDoDoanhThu"] if prev else "xuat_hd"}, commit=False)
        if prev:
            rows = [(r["Nhom"], r["TenDong"], r["Dau"], r["Nguon"], r["HeThong"], r["An"])
                    for r in self.rows(prev["Nam"])]
        else:
            rows = [(*x, 0) for x in DEFAULT_ROWS]
        for i, (nhom, ten, dau, nguon, he_thong, an) in enumerate(rows, start=1):
            self.fdb.insert("LaiLoDong", {"Nam": year, "ThuTu": i, "Nhom": nhom, "TenDong": ten, "Dau": dau,
                                          "Nguon": nguon, "HeThong": he_thong, "An": an}, commit=False)
        self.store.save()

    def _dedupe(self, year: int) -> None:
        """Hai máy cùng tạo bảng của một năm trước khi đồng bộ: giữ bản sửa sau cùng."""
        changed = False
        seen = set()
        for r in self.fdb.query("SELECT Ma FROM LaiLoNam WHERE DaXoa = 0 AND Nam = ? ORDER BY NgaySua DESC", (year,)):
            if seen:
                self._hard_delete("LaiLoNam", r["Ma"])
                changed = True
            seen.add(r["Ma"])
        keys = set()
        for r in self.fdb.query("SELECT * FROM LaiLoDong WHERE DaXoa = 0 AND Nam = ? AND HeThong = 1 "
                                "ORDER BY NgaySua DESC", (year,)):
            k = (r["Nhom"], r["Nguon"], r["TenDong"] if r["Nguon"] == "nhap" else "")
            if k in keys and r["NgayTao"] == r["NgaySua"]:
                self._hard_delete("LaiLoDong", r["Ma"])
                changed = True
            keys.add(k)
        if changed:
            self.store.save()

    def rows(self, year: int) -> list:
        return self.fdb.query("SELECT * FROM LaiLoDong WHERE DaXoa = 0 AND Nam = ? ORDER BY "
                              "CASE Nhom WHEN 'Doanh thu' THEN 0 WHEN 'Điều chỉnh' THEN 1 WHEN 'Chi phí' THEN 2 "
                              "ELSE 3 END, ThuTu", (year,))

    def _check_open(self, year: int) -> None:
        info = self.year_info(year)
        if info["DaChot"]:
            raise ProfitError(f"Năm {year} đã chốt. Mở chốt (Giám đốc) rồi mới sửa được.")

    def add_row(self, year: int, nhom: str, ten: str, dau: int = 1, gia_tri: int | None = None) -> str:
        self._check_open(year)
        if nhom not in GROUPS:
            raise ProfitError("Nhóm không hợp lệ.")
        if not ten.strip():
            raise ProfitError("Chưa nhập tên dòng.")
        n = self.fdb.one("SELECT MAX(ThuTu) m FROM LaiLoDong WHERE Nam = ? AND Nhom = ?", (year, nhom))["m"] or 0
        ma = self.fdb.insert("LaiLoDong", {"Nam": year, "ThuTu": n + 1, "Nhom": nhom, "TenDong": ten.strip(),
                                           "Dau": -1 if dau < 0 else 1, "Nguon": "nhap", "GiaTriNhap": gia_tri})
        self.store.save()
        return ma

    def update_row(self, ma: str, values: dict) -> None:
        r = self.fdb.get("LaiLoDong", ma)
        if r is None:
            raise ProfitError("Không tìm thấy dòng.")
        self._check_open(r["Nam"])
        v = {k: values[k] for k in ("TenDong", "GiaTriNhap", "Dau", "An", "GhiChu") if k in values}
        if "TenDong" in v and not str(v["TenDong"]).strip():
            raise ProfitError("Chưa nhập tên dòng.")
        if "Dau" in v and r["HeThong"] and r["Nguon"] != "nhap":
            v.pop("Dau")                       # dấu của dòng tự tính cố định
        self.fdb.update("LaiLoDong", ma, v)
        self.store.save()

    def move_row(self, ma: str, step: int) -> None:
        r = self.fdb.get("LaiLoDong", ma)
        self._check_open(r["Nam"])
        same = [x for x in self.rows(r["Nam"]) if x["Nhom"] == r["Nhom"]]
        i = next(k for k, x in enumerate(same) if x["Ma"] == ma)
        j = i + step
        if not 0 <= j < len(same):
            return
        same[i], same[j] = same[j], same[i]
        for k, x in enumerate(same, start=1):
            if x["ThuTu"] != k:
                self.fdb.update("LaiLoDong", x["Ma"], {"ThuTu": k}, commit=False)
        self.store.save()

    def delete_row(self, ma: str) -> None:
        r = self.fdb.get("LaiLoDong", ma)
        if r is None:
            return
        self._check_open(r["Nam"])
        if r["HeThong"]:
            raise ProfitError("Dòng mặc định chỉ ẩn được, không xóa. Dòng tự thêm thì xóa được.")
        self._hard_delete("LaiLoDong", ma)
        self.store.save()

    def set_year(self, year: int, values: dict) -> None:
        info = self.year_info(year)
        if info["DaChot"]:
            raise ProfitError(f"Năm {year} đã chốt. Mở chốt rồi mới sửa được.")
        v = {k: values[k] for k in ("TyLeThueTNDN", "ThueTNDNNhap", "TNCNConPhaiNop", "CheDoDoanhThu", "GhiChu")
             if k in values}
        if "TyLeThueTNDN" in v and not 0 <= float(v["TyLeThueTNDN"]) <= 100:
            raise ProfitError("Tỷ lệ thuế TNDN từ 0 đến 100%.")
        self.fdb.update("LaiLoNam", info["Ma"], v)
        self.store.save()

    # ------------------------------------------------------------ xóa vĩnh viễn
    def _hard_delete(self, table: str, ma: str) -> None:
        self.fdb.conn.execute(f"DELETE FROM {table} WHERE Ma = ?", (ma,))
        self.fdb.conn.execute("INSERT OR REPLACE INTO XoaVinhVien (Bang, Ma, NgayXoa) VALUES (?, ?, ?)",
                              (table, ma, now_iso()))

    def delete_year(self, year: int) -> int:
        n = 0
        for t in ("LaiLoDong", "LaiLoNam"):
            for r in self.fdb.query(f"SELECT Ma FROM {t} WHERE Nam = ?", (year,)):
                self._hard_delete(t, r["Ma"])
                n += 1
        self.store._compact()
        self.db.log(self.store.username, f"Xóa vĩnh viễn bảng lãi lỗ năm {year}")
        return n

    def delete_all(self) -> int:
        n = 0
        for y in {r["Nam"] for r in self.fdb.query("SELECT Nam FROM LaiLoNam UNION SELECT Nam FROM LaiLoDong")}:
            for t in ("LaiLoDong", "LaiLoNam"):
                for r in self.fdb.query(f"SELECT Ma FROM {t} WHERE Nam = ?", (y,)):
                    self._hard_delete(t, r["Ma"])
                    n += 1
        self.store._compact()
        self.db.log(self.store.username, "Xóa vĩnh viễn toàn bộ bảng lãi lỗ")
        return n

    # ------------------------------------------------------------ số tự động
    def revenue(self, year: int, mode: str = "xuat_hd") -> tuple[int, list[int], list[dict]]:
        y0, y1 = f"{year}-01-01", f"{year}-12-31"
        invs = self.db.query(
            "SELECT h.*, c.SoHopDong, c.TenHopDong FROM HoaDonDauRa h LEFT JOIN HopDong c ON c.Ma = h.MaHopDong "
            "WHERE h.DaXoa = 0 AND h.TinhChatHD != 'Đã hủy' AND h.NgayXuatHD BETWEEN ? AND ? ORDER BY h.NgayXuatHD",
            (y0, y1))
        q, rows = [0, 0, 0, 0], []
        for h in invs:
            paid = self.db.one("SELECT COALESCE(SUM(SoTien), 0) s FROM ThanhToanHopDong WHERE DaXoa = 0 "
                               "AND MaHDDauRa = ? AND NgayThanhToan <= ?", (h["Ma"], y1))["s"]
            if mode == "da_tt":
                sau = int(h["GiaTriSauThue"] or 0)
                amt = (h["GiaTriTruocThue"] if sau and paid >= sau else
                       _round(Decimal(h["GiaTriTruocThue"]) * min(paid, sau) / sau) if sau else 0)
            else:
                amt = int(h["GiaTriTruocThue"] or 0)
            q[_quarter(h["NgayXuatHD"])] += amt
            rows.append(dict(h, DaThu3112=paid, TinhDoanhThu=amt))
        return sum(q), q, rows

    def wip_suggestion(self, year: int) -> tuple[int, list[dict]]:
        """Hợp đồng đầu ra ký trong năm, đến 31/12 chưa xuất hóa đơn nào: giá trị trước thuế."""
        rows = [dict(r) for r in self.db.query(
            "SELECT * FROM HopDong c WHERE DaXoa = 0 AND ChieuHopDong = 'Đầu ra' AND TinhTrangHopDong != 'Hủy' "
            "AND NgayKy BETWEEN ? AND ? AND NOT EXISTS (SELECT 1 FROM HoaDonDauRa h WHERE h.MaHopDong = c.Ma "
            "AND h.DaXoa = 0 AND h.TinhChatHD != 'Đã hủy' AND h.NgayXuatHD <= ?) ORDER BY NgayKy",
            (f"{year}-01-01", f"{year}-12-31", f"{year}-12-31"))]
        return sum(int(r["GiaTriTruocThue"] or 0) for r in rows), rows

    def _wip_end_value(self, year: int) -> int:
        """Số dở dang để sang năm sau của một năm (đã chốt / nhập tay / gợi ý)."""
        info = self.fdb.one("SELECT * FROM LaiLoNam WHERE DaXoa = 0 AND Nam = ? ORDER BY NgaySua DESC", (year,))
        if info is None:
            return 0
        if info["DaChot"] and info["SoLieuChot"]:
            for r in json.loads(info["SoLieuChot"])["rows"]:
                if r["Nguon"] == "do_dang_cuoi":
                    return abs(r["GiaTri"])
        r = self.fdb.one("SELECT GiaTriNhap FROM LaiLoDong WHERE DaXoa = 0 AND Nam = ? AND Nguon = 'do_dang_cuoi'",
                         (year,))
        if r is not None and r["GiaTriNhap"] is not None:
            return int(r["GiaTriNhap"])
        return self.wip_suggestion(year)[0]

    def auto_values(self, year: int, mode: str) -> dict[str, tuple[int, list[int] | None]]:
        dt, q, _ = self.revenue(year, mode)
        out: dict[str, tuple[int, list[int] | None]] = {"doanh_thu": (dt, q)}
        tot = self.assets.year_totals(year)
        out["tscd"] = (sum(tot["TSCĐ"]), tot["TSCĐ"])
        out["ccdc"] = (sum(tot["CCDC"]), tot["CCDC"])
        by_line: dict[str, list[int]] = {}
        for r in self.assets.expense_lines(year):
            by_line.setdefault(f"dong:{r['DongLaiLo']}", [0, 0, 0, 0])[r["Quy"] - 1] += int(r["ThanhTien"] or 0)
        for key in ("dong:6", "dong:7", "dong:6a"):
            qq = by_line.pop(key, [0, 0, 0, 0])
            out[key] = (sum(qq), qq)
        for key, qq in by_line.items():          # nhóm danh mục gắn dòng chưa có trên bảng mẫu
            out[key] = (sum(qq), qq)
        out["do_dang_dau"] = (self._wip_end_value(year - 1), None)
        out["do_dang_cuoi"] = (self.wip_suggestion(year)[0], None)
        return out

    # ------------------------------------------------------------ bảng lãi lỗ
    def compute(self, year: int, *, live: bool = False) -> dict:
        """Bảng lãi lỗ của năm. Năm đã chốt trả về số chốt (live=True để xem số đang có trong dữ liệu)."""
        info = self.year_info(year)
        if info["DaChot"] and info["SoLieuChot"] and not live:
            res = json.loads(info["SoLieuChot"])
            res["da_chot"] = True
            return res
        auto = self.auto_values(year, info["CheDoDoanhThu"])
        used, rows, warnings = set(), [], []
        stt = 0
        for r in self.rows(year):
            d = dict(r)
            a_val, a_q = auto.get(r["Nguon"], (None, None)) if r["Nguon"] != "nhap" else (None, None)
            used.add(r["Nguon"])
            if r["GiaTriNhap"] is not None:
                val, q, ghi_de = int(r["GiaTriNhap"]), None, r["Nguon"] != "nhap"
            else:
                val, q, ghi_de = int(a_val or 0), a_q, False
            hidden = bool(r["An"]) or (r["Nguon"] in HIDE_WHEN_ZERO and val == 0)
            if r["An"] and val:
                warnings.append(f"Dòng \"{r['TenDong']}\" đang ẩn nên không cộng ({val:,} đ).".replace(",", "."))
            if r["Nhom"] == "Chi phí" and not hidden:
                stt += 1
            d.update(GiaTri=val * r["Dau"], Quy=[x * r["Dau"] for x in q] if q else None, TuDong=a_val,
                     GhiDe=ghi_de, Hien=not hidden, STT=stt if r["Nhom"] == "Chi phí" and not hidden else None,
                     NguonText=SOURCES.get(r["Nguon"], "Tự động"))
            rows.append(d)
        for key, (val, _) in auto.items():
            if key.startswith("dong:") and key not in used and val:
                warnings.append(f"Có {val:,} đ chi phí trong kỳ của dòng {key[5:]} chưa có dòng trên bảng "
                                f"(thêm lại dòng tự động hoặc kiểm tra).".replace(",", "."))
        live_rows = [r for r in rows if r["Hien"]]

        def total(nhom):
            return sum(r["GiaTri"] for r in live_rows if r["Nhom"] == nhom)

        def total_q(nhoms):
            out = [0, 0, 0, 0]
            for r in live_rows:
                if r["Nhom"] in nhoms and r["Quy"]:
                    out = [a + b for a, b in zip(out, r["Quy"])]
            return out

        tong_dt = total("Doanh thu") + total("Điều chỉnh")
        chi_phi = total("Chi phí")
        loi_nhuan = tong_dt - chi_phi
        ty_le = float(info["TyLeThueTNDN"])
        tndn = (int(info["ThueTNDNNhap"]) if info["ThueTNDNNhap"] is not None
                else _round(Decimal(max(loi_nhuan, 0)) * Decimal(str(ty_le)) / 100))
        tong_thue = tndn + total("Thuế")
        tncn_con = int(info["TNCNConPhaiNop"] or 0)
        pending = self.assets.year_check(year)
        if pending:
            warnings.append(f"Còn {pending} dòng hóa đơn đầu vào năm {year} chờ kế toán phân loại: chi phí, "
                            "khấu hao chưa đủ.")
        return {"nam": year, "rows": rows, "tong_doanh_thu": tong_dt, "chi_phi": chi_phi, "loi_nhuan": loi_nhuan,
                "ty_le": ty_le, "thue_tndn": tndn, "thue_tndn_nhap": info["ThueTNDNNhap"] is not None,
                "tong_thue": tong_thue, "tncn_con": tncn_con, "tong_phai_nop": tong_thue + tncn_con,
                "doanh_thu_quy": total_q(("Doanh thu", "Điều chỉnh")), "chi_phi_quy": total_q(("Chi phí",)),
                "che_do": info["CheDoDoanhThu"], "cho_phan_loai": pending, "canh_bao": warnings, "da_chot": False}

    def close_year(self, year: int) -> None:
        pending = self.assets.year_check(year)
        if pending:
            raise ProfitError(f"Còn {pending} dòng hóa đơn đầu vào năm {year} chưa phân loại. Kế toán phân loại "
                              "xong mới chốt được năm.")
        res = self.compute(year, live=True)
        info = self.year_info(year)
        self.fdb.update("LaiLoNam", info["Ma"], {"DaChot": 1, "SoLieuChot": json.dumps(res, ensure_ascii=False)})
        self.store.save()
        self.db.log(self.store.username, f"Chốt bảng lãi lỗ năm {year}")

    def reopen_year(self, user, year: int) -> None:
        if user["VaiTro"] != "Giám đốc":
            raise ProfitError("Chỉ Giám đốc mở chốt được.")
        info = self.year_info(year)
        self.fdb.update("LaiLoNam", info["Ma"], {"DaChot": 0, "SoLieuChot": None})
        self.store.save()
        self.db.log(self.store.username, f"Mở chốt bảng lãi lỗ năm {year}")

    # ------------------------------------------------------------ xuất Excel
    def export_xlsx(self, path: str | Path, year: int, *, company: str = "", user: str = "",
                    password: str | None = None) -> Path:
        """Bảng lãi lỗ như mẫu của công ty, các dòng tổng là công thức Excel."""
        import io

        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        res = self.compute(year)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Lãi lỗ {year}"
        ws["A1"] = company
        ws["A1"].font = Font(bold=True)
        ws["A2"] = f"BẢNG TỔNG HỢP LÃI LỖ NĂM {year}" + (" (đã chốt)" if res.get("da_chot") else "")
        ws["A2"].font = Font(bold=True, size=14)
        ws["A3"] = f"Người xuất: {user} · {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws["A3"].font = Font(italic=True, color="555555")
        heads = ["STT", "Nội dung", "Số tiền", "Quý 1", "Quý 2", "Quý 3", "Quý 4", "Nguồn số liệu"]
        thin = Side(style="thin", color="999999")
        border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for j, h in enumerate(heads, start=1):
            c = ws.cell(5, j, h)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9E2F3")
            c.border = border
            c.alignment = Alignment(horizontal="center")
        r = 6
        bold_rows = set()

        def put(stt, label, value, quy=None, src="", bold=False):
            nonlocal r
            ws.cell(r, 1, stt)
            ws.cell(r, 2, label)
            ws.cell(r, 3, value)
            for k in range(4):
                if quy:
                    ws.cell(r, 4 + k, quy[k])
            ws.cell(r, 8, src)
            if bold:
                bold_rows.add(r)
            r += 1
            return r - 1

        vis = [x for x in res["rows"] if x["Hien"]]
        cells: dict[str, list[int]] = {g: [] for g in GROUPS}
        for g in ("Doanh thu", "Điều chỉnh"):
            if g == "Điều chỉnh" and any(x["Nhom"] == g for x in vis):
                put("", "Điều chỉnh", None, bold=True)
            for x in (x for x in vis if x["Nhom"] == g):
                cells[g].append(put("", x["TenDong"], x["GiaTri"], x["Quy"], x["NguonText"] +
                                    (" (đã ghi đè)" if x["GhiDe"] else ""), bold=g == "Doanh thu"))
        refs = "+".join(f"C{i}" for i in cells["Doanh thu"] + cells["Điều chỉnh"]) or "0"
        r_dt = put("", "Tổng doanh thu", f"={refs}", src="Công thức", bold=True)
        r_cp = put("", "Chi phí", None, src="Công thức", bold=True)
        for x in (x for x in vis if x["Nhom"] == "Chi phí"):
            cells["Chi phí"].append(put(x["STT"], "- " + x["TenDong"], x["GiaTri"], x["Quy"],
                                        x["NguonText"] + (" (đã ghi đè)" if x["GhiDe"] else "")))
        cp = cells["Chi phí"]
        ws.cell(r_cp, 3, f"=SUM(C{cp[0]}:C{cp[-1]})" if cp else 0)
        r_ln = put("", f"Lợi nhuận {year}", f"=C{r_dt}-C{r_cp}", src="Tổng doanh thu – Chi phí", bold=True)
        if res["thue_tndn_nhap"]:
            r_tndn = put("", "Thuế TNDN", res["thue_tndn"], src="Nhập tay")
        else:
            r_tndn = put("", "Thuế TNDN", f"=ROUND(MAX(C{r_ln},0)*I{r},0)", src="Lợi nhuận × tỷ lệ ở cột I")
            ws.cell(r_tndn, 9, res["ty_le"] / 100).number_format = "0%"
        for x in (x for x in vis if x["Nhom"] == "Thuế"):
            cells["Thuế"].append(put("", x["TenDong"], x["GiaTri"], src=x["NguonText"]))
        r_tt = put("", "Tổng số thuế phải nộp", "=" + "+".join([f"C{r_tndn}"] + [f"C{i}" for i in cells["Thuế"]]),
                   src="Công thức", bold=True)
        r_tncn = put("", "Thuế thu nhập cá nhân còn phải nộp", res["tncn_con"], src="Nhập tay")
        put("", "Tổng chi phí phải nộp", f"=C{r_tt}+C{r_tncn}", src="Công thức", bold=True)
        for i in range(6, r):
            for j in range(1, 9):
                c = ws.cell(i, j)
                c.border = border
                if 3 <= j <= 7:
                    c.number_format = '#,##0;(#,##0);""' if j > 3 else "#,##0;-#,##0;0"
                if i in bold_rows and j <= 3:
                    c.font = Font(bold=True)
        if res["canh_bao"]:
            r += 1
            for w in res["canh_bao"]:
                ws.cell(r, 2, "Lưu ý: " + w).font = Font(color="C00000")
                r += 1
        for col, w in zip("ABCDEFGHI", [6, 44, 18, 15, 15, 15, 15, 56, 8]):
            ws.column_dimensions[col].width = w
        path = Path(path)
        if not password:
            wb.save(path)
            return path
        from msoffcrypto.format.ooxml import OOXMLFile

        plain = io.BytesIO()
        wb.save(plain)
        plain.seek(0)
        with open(path, "wb") as out:
            OOXMLFile(plain).encrypt(password, out)
        return path

    def year_report(self, year: int) -> list:
        """Báo cáo tổng doanh thu, tổng chi phí năm (các bảng Report để xem / xuất Excel)."""
        from .reports import Report

        res = self.compute(year, live=True)
        info = self.year_info(year)
        dt, dq, inv_rows = self.revenue(year, info["CheDoDoanhThu"])
        tong = [{"NoiDung": "Doanh thu (hóa đơn đầu ra)", "Q1": dq[0], "Q2": dq[1], "Q3": dq[2], "Q4": dq[3],
                 "Nam": dt}]
        for x in res["rows"]:
            if x["Hien"] and x["Nhom"] == "Điều chỉnh":
                tong.append({"NoiDung": "Điều chỉnh: " + x["TenDong"], "Nam": x["GiaTri"]})
        tong.append({"NoiDung": "Tổng doanh thu", "Nam": res["tong_doanh_thu"], "_kieu": "cong"})
        for x in res["rows"]:
            if x["Hien"] and x["Nhom"] == "Chi phí":
                q = x["Quy"] or [None] * 4
                tong.append({"NoiDung": f"{x['STT']}. {x['TenDong']}", "Q1": q[0], "Q2": q[1], "Q3": q[2],
                             "Q4": q[3], "Nam": x["GiaTri"]})
        tong.append({"NoiDung": "Tổng chi phí", "Nam": res["chi_phi"], "_kieu": "cong"})
        tong.append({"NoiDung": "Lợi nhuận", "Nam": res["loi_nhuan"], "_kieu": "cong"})
        qcols = [("NoiDung", "Nội dung", "text"), ("Q1", "Quý 1", "money"), ("Q2", "Quý 2", "money"),
                 ("Q3", "Quý 3", "money"), ("Q4", "Quý 4", "money"), ("Nam", "Cả năm", "money")]
        reps = [Report(f"Tổng hợp doanh thu, chi phí năm {year}", qcols, tong, sheet="Tổng hợp",
                       subtitle="Doanh thu: " + dict(REVENUE_MODES)[info["CheDoDoanhThu"]])]
        dt_cols = [("NgayXuatHD", "Ngày hóa đơn", "date"), ("KyHieuHD", "Ký hiệu", "text"),
                   ("SoHoaDon", "Số", "text"), ("SoHopDong", "Hợp đồng", "text"), ("TenNguoiMua", "Người mua", "text"),
                   ("GiaTriTruocThue", "Trước thuế", "money"), ("VAT", "VAT", "money"),
                   ("GiaTriSauThue", "Sau thuế", "money"), ("DaThu3112", "Đã thu đến 31/12", "money"),
                   ("TinhDoanhThu", "Tính doanh thu", "money")]
        reps.append(Report(f"Doanh thu năm {year}", dt_cols, inv_rows, sheet="Doanh thu",
                           total={"TenNguoiMua": "Cộng", **{k: sum(int(r[k] or 0) for r in inv_rows) for k in
                                  ("GiaTriTruocThue", "VAT", "GiaTriSauThue", "DaThu3112", "TinhDoanhThu")}}))
        lines = self.assets.expense_lines(year)
        cp_cols = [("NgayNhanHD", "Ngày hóa đơn", "date"), ("SoHoaDonGoc", "Số HĐ", "text"),
                   ("TenNhaCungCap", "Người bán", "text"), ("TenHang", "Hàng hóa, dịch vụ", "text"),
                   ("NhomDM", "Nhóm", "text"), ("DongLaiLo", "Dòng lãi lỗ", "text"),
                   ("ThanhTien", "Chưa thuế", "money")]
        reps.append(Report(f"Chi phí trong kỳ từ hóa đơn đầu vào năm {year}", cp_cols, lines, sheet="Chi phí HĐ",
                           total={"TenHang": "Cộng", "ThanhTien": sum(int(r["ThanhTien"] or 0) for r in lines)}))
        kh_rows = []
        for loai in ("TSCĐ", "CCDC"):
            for a in self.assets.register(loai, year, "co_ps"):
                kh_rows.append({"Loai": loai, "MaTS": a["MaTS"], "TenTS": a["TenTS"], "NgayGhiTang": a["NgayGhiTang"],
                                "SoKy": a["SoKy"], "TrongKy": a["TrongKy"], "ConLai": a["ConLai"]})
        kh_cols = [("Loai", "Sổ", "text"), ("MaTS", "Mã", "text"), ("TenTS", "Tên", "text"),
                   ("NgayGhiTang", "Ngày ghi tăng", "date"), ("SoKy", "Số kỳ", "int"),
                   ("TrongKy", "Khấu hao / phân bổ trong năm", "money"), ("ConLai", "Giá trị còn lại", "money")]
        reps.append(Report(f"Khấu hao TSCĐ, phân bổ CCDC năm {year}", kh_cols, kh_rows, sheet="Khấu hao, phân bổ",
                           total={"TenTS": "Cộng", "TrongKy": sum(r["TrongKy"] for r in kh_rows),
                                  "ConLai": sum(r["ConLai"] for r in kh_rows)}))
        return reps


__all__ = ["ProfitService", "ProfitError", "GROUPS", "REVENUE_MODES", "SOURCES", "dm"]
